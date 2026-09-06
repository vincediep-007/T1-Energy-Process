# =================== 1-CLICK MASTER REPORT MERGER ===================
# Standalone utility for Shared Laptop / QC Report Merging
import os
import sys
import tkinter as tk
from tkinter import filedialog, messagebox
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

def merge_exported_to_master(source_export_path: str, master_report_path: str) -> tuple:
    if not os.path.exists(source_export_path):
        return False, f"Source export file not found: {source_export_path}"
    
    try:
        src_wb = openpyxl.load_workbook(source_export_path, data_only=True)
        src_ws = src_wb.active
        
        # Load or create master report workbook
        if os.path.exists(master_report_path):
            dst_wb = openpyxl.load_workbook(master_report_path)
            dst_ws = dst_wb.active
        else:
            parent_dir = os.path.dirname(master_report_path)
            if parent_dir: os.makedirs(parent_dir, exist_ok=True)
            dst_wb = openpyxl.Workbook()
            dst_ws = dst_wb.active
            dst_ws.title = "QC Daily Report"
            headers = [
                "Date", "Order No", "Serial No", "Defect Summary", "Classification", 
                "Result", "Defect Cause", "Defect Location", "Root Cause", 
                "Pre-Layup Photo", "Post-Layup Photo", "Layup Time", "Station", 
                "Eval Shift", "Line", "Responsible Shift"
            ]
            dst_ws.append(headers)
            header_fill = PatternFill(start_color="1A5B82", end_color="1A5B82", fill_type="solid")
            header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
            for col_idx in range(1, len(headers) + 1):
                cell = dst_ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            dst_ws.row_dimensions[1].height = 24

        # Read existing SNs to prevent duplicates
        existing_sns = set()
        for r in range(2, dst_ws.max_row + 1):
            val = dst_ws.cell(row=r, column=3).value
            if val: existing_sns.add(str(val).strip())

        thin_border = Border(left=Side(style='thin', color='E2E8F0'),
                             right=Side(style='thin', color='E2E8F0'),
                             top=Side(style='thin', color='E2E8F0'),
                             bottom=Side(style='thin', color='E2E8F0'))

        merged_count = 0
        for r in range(2, src_ws.max_row + 1):
            sn = str(src_ws.cell(row=r, column=3).value or '').strip()
            if sn and sn in existing_sns and sn != "Pending SN":
                continue

            r_idx = dst_ws.max_row + 1
            for c in range(1, 17):
                val = src_ws.cell(row=r, column=c).value
                if val in ("-", "None", "Pending SN", None):
                    val = ""
                cell = dst_ws.cell(row=r_idx, column=c, value=val)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = Font(name="Segoe UI", size=9)

            # Style Q3 and Scrap
            res_val = str(src_ws.cell(row=r, column=6).value or '').upper()
            if "SCRAP" in res_val:
                dst_ws.cell(row=r_idx, column=6).fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                dst_ws.cell(row=r_idx, column=6).font = Font(name="Segoe UI", size=9, bold=True, color="DC2626")
            elif "Q3" in res_val:
                dst_ws.cell(row=r_idx, column=6).fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                dst_ws.cell(row=r_idx, column=6).font = Font(name="Segoe UI", size=9, bold=True, color="CA8A04")

            # Copy images if present (both Column J Pre-Layup and Column K Post-Layup)
            if hasattr(src_ws, '_images'):
                for img in src_ws._images:
                    anc = getattr(img, 'anchor', None)
                    is_col_j = False
                    is_col_k = False
                    
                    if anc is not None:
                        if hasattr(anc, '_from') and hasattr(anc._from, 'col') and hasattr(anc._from, 'row'):
                            if anc._from.row == r - 1:
                                if anc._from.col == 9: is_col_j = True
                                elif anc._from.col == 10: is_col_k = True
                        elif isinstance(anc, str):
                            if anc.startswith(f"J{r}"): is_col_j = True
                            elif anc.startswith(f"K{r}"): is_col_k = True

                    if is_col_j:
                        try:
                            dst_ws.add_image(img, f"J{r_idx}")
                            dst_ws.row_dimensions[r_idx].height = 80
                        except Exception: pass
                    elif is_col_k:
                        try:
                            dst_ws.add_image(img, f"K{r_idx}")
                            dst_ws.row_dimensions[r_idx].height = 80
                        except Exception: pass

            if sn and sn != "Pending SN":
                existing_sns.add(sn)
            merged_count += 1

        # Auto column widths
        for col in dst_ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            dst_ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        dst_ws.column_dimensions['J'].width = 18
        dst_ws.column_dimensions['K'].width = 18

        dst_wb.save(master_report_path)
        return True, f"Successfully merged {merged_count} new defect inspection records into Master Report!"
    except Exception as e:
        return False, f"Merge Failed: {e}"


def run_gui():
    root = tk.Tk()
    root.title("1-Click Master QC Report Merger")
    root.geometry("540x260")
    root.configure(bg="#0f172a")

    tk.Label(root, text="Master QC Excel Report Merger", font=("Segoe UI", 13, "bold"), bg="#0f172a", fg="#ffffff").pack(pady=(15, 4))
    tk.Label(root, text="Select your exported QC file and master report to merge in 1 click.", font=("Segoe UI", 9), bg="#0f172a", fg="#94a3b8").pack(pady=(0, 15))

    frame = tk.Frame(root, bg="#1e293b", padx=12, pady=12)
    frame.pack(fill=tk.BOTH, expand=True, padx=15, pady=(0, 15))

    tk.Label(frame, text="1. Exported QC File (.xlsx):", font=("Segoe UI", 9, "bold"), bg="#1e293b", fg="#cbd5e1").grid(row=0, column=0, sticky="w", pady=4)
    ent_src = tk.Entry(frame, width=36, font=("Segoe UI", 9))
    ent_src.grid(row=0, column=1, padx=6, pady=4)
    tk.Button(frame, text="Browse", command=lambda: (ent_src.delete(0, tk.END), ent_src.insert(0, filedialog.askopenfilename(filetypes=[("Excel Files", "*.xlsx")]))), bg="#3b82f6", fg="#ffffff", relief=tk.FLAT).grid(row=0, column=2, padx=4)

    tk.Label(frame, text="2. Master Report (.xlsx):", font=("Segoe UI", 9, "bold"), bg="#1e293b", fg="#cbd5e1").grid(row=1, column=0, sticky="w", pady=4)
    ent_dst = tk.Entry(frame, width=36, font=("Segoe UI", 9))
    ent_dst.grid(row=1, column=1, padx=6, pady=4)
    tk.Button(frame, text="Browse", command=lambda: (ent_dst.delete(0, tk.END), ent_dst.insert(0, filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel Files", "*.xlsx")]))), bg="#3b82f6", fg="#ffffff", relief=tk.FLAT).grid(row=1, column=2, padx=4)

    def do_merge():
        src = ent_src.get().strip()
        dst = ent_dst.get().strip()
        if not src or not dst:
            messagebox.showwarning("Incomplete", "Please specify both the exported QC file and the master report!")
            return
        ok, msg = merge_exported_to_master(src, dst)
        if ok:
            messagebox.showinfo("Merge Complete", msg)
        else:
            messagebox.showerror("Merge Error", msg)

    tk.Button(frame, text="⚡ MERGE INTO MASTER REPORT", command=do_merge, font=("Segoe UI", 10, "bold"), bg="#06d6a0", fg="#0b132b", padx=12, pady=6, relief=tk.FLAT, cursor="hand2").grid(row=2, column=0, columnspan=3, pady=(12, 0))

    root.mainloop()

if __name__ == "__main__":
    if len(sys.argv) >= 3:
        ok, msg = merge_exported_to_master(sys.argv[1], sys.argv[2])
        print(msg)
        sys.exit(0 if ok else 1)
    else:
        run_gui()
