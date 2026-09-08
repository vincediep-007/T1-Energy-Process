# =================== MAIN APPLICATION - VERSION 18.6 ===================
import sys
import os
import json
import io
import time
import ctypes
import tkinter as tk
from tkinter import ttk, messagebox, scrolledtext, filedialog
import webbrowser
import threading
import queue
from datetime import datetime, timedelta
import calendar
from PIL import Image, ImageTk, ImageGrab
import win32clipboard
import subprocess

try:
    ctypes.windll.shcore.SetProcessDpiAwareness(1)
except Exception:
    try: ctypes.windll.user32.SetProcessDPIAware()
    except Exception: pass

if getattr(sys, 'frozen', False):
    current_dir = os.path.dirname(sys.executable)
else:
    current_dir = os.path.dirname(os.path.abspath(__file__))

if current_dir not in sys.path:
    sys.path.insert(0, current_dir)

MAX_RUNTIME_LOG_BYTES = 10 * 1024 * 1024  # 10 MB limit before rotation
RUNTIME_LOG_FILE = os.path.join(current_dir, "app_runtime.log")

class DualTerminalLogger:
    def __init__(self, filepath, original_stream, stream_name="STDOUT"):
        self.filepath = filepath
        self.original_stream = original_stream
        self.stream_name = stream_name
        self.lock = threading.Lock()
        self._at_start_of_line = True

    def write(self, message):
        if not message:
            return
        try:
            if self.original_stream:
                try:
                    self.original_stream.write(message)
                    self.original_stream.flush()
                except UnicodeEncodeError:
                    safe_msg = message.encode(getattr(self.original_stream, 'encoding', 'ascii') or 'ascii', errors='replace').decode('ascii')
                    self.original_stream.write(safe_msg)
                    self.original_stream.flush()
        except Exception:
            pass

        try:
            with self.lock:
                if os.path.exists(self.filepath) and os.path.getsize(self.filepath) > MAX_RUNTIME_LOG_BYTES:
                    backup_path = self.filepath + ".old"
                    try:
                        if os.path.exists(backup_path):
                            os.remove(backup_path)
                        os.rename(self.filepath, backup_path)
                    except Exception:
                        pass

                with open(self.filepath, "a", encoding="utf-8", errors="replace") as f:
                    lines = message.split('\n')
                    for idx, line in enumerate(lines):
                        if idx > 0:
                            f.write('\n')
                            self._at_start_of_line = True
                        if line:
                            if self._at_start_of_line:
                                ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
                                prefix = f"[{ts}] "
                                if self.stream_name == "STDERR":
                                    prefix = f"[{ts}] [STDERR] "
                                f.write(prefix)
                                self._at_start_of_line = False
                            f.write(line)
                    f.flush()
        except Exception:
            pass

    def flush(self):
        try:
            if self.original_stream:
                self.original_stream.flush()
        except Exception:
            pass

    def fileno(self):
        if hasattr(self.original_stream, "fileno"):
            return self.original_stream.fileno()
        raise io.UnsupportedOperation("fileno")

    def isatty(self):
        if hasattr(self.original_stream, "isatty"):
            try:
                return self.original_stream.isatty()
            except Exception:
                pass
        return False

    def reconfigure(self, **kwargs):
        if hasattr(self.original_stream, "reconfigure"):
            try:
                self.original_stream.reconfigure(**kwargs)
            except Exception:
                pass

    @property
    def encoding(self):
        return getattr(self.original_stream, 'encoding', 'utf-8')

# Install dual logger immediately so all prints and startup diagnostics are saved
if not isinstance(sys.stdout, DualTerminalLogger):
    sys.stdout = DualTerminalLogger(RUNTIME_LOG_FILE, sys.stdout, "STDOUT")
if not isinstance(sys.stderr, DualTerminalLogger):
    sys.stderr = DualTerminalLogger(RUNTIME_LOG_FILE, sys.stderr, "STDERR")

print(f"=== QC SUITE RUNTIME LOG SESSION INITIALIZED AT {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} ===")

import config
from search_engine import ImageSearchEngine
from image_processor import ImageProcessor
from data_manager import DataManager, sync_to_master_excel, save_pre_el_search_history, save_module_review_history

from shift_calculator import (
    get_responsible_shift_and_line, get_evaluation_shift, parse_line_from_station,
    get_operational_shift_info, get_pre_el_chinese_shift_folder, get_pinpointed_pre_el_paths,
    get_pre_el_stations_for_layup_station
)
from phone_server import (
    start_phone_server, train_voice_pattern, test_mac_mini_relay_connection,
    get_file_datetime, refresh_save_folders_for_today, update_live_hud_state,
    rename_to_sn_pattern, get_mobile_hud_url, clean_and_validate_sn,
    load_mes_trend_log, save_mes_trend_entry, query_mes_process_log,
    get_mes_trend_analytics, export_mes_trend_csv_string, MES_TREND_FILE,
    get_mes_layup_time_for_sn, parse_mes_datetime, extract_layup_time_from_content,
    MES_CHINESE_DICTIONARY, translate_mes_text
)






def generate_shift_time_slots(interval_str):
    if interval_str == "1 Hour": step = 1
    elif interval_str == "2 Hours": step = 2
    elif interval_str == "4 Hours": step = 4
    else: return ["All Shift (06:00 - 18:00)"]
    slots = ["All Shift (06:00 - 18:00)"]
    start_base = getattr(config, 'SHIFT_START_HOUR', 6)
    for i in range(12 // step):
        sh = start_base + (i * step)
        eh = sh + step
        slots.append(f"{sh:02d}:00 - {eh:02d}:00")
    return slots

def get_current_shift_slot(slots):
    now_h = datetime.now().hour
    for s in slots:
        if "All Shift" in s: continue
        try:
            p1, p2 = s.split(' - ')
            sh, eh = int(p1.split(':')[0]), int(p2.split(':')[0])
            if sh <= now_h < eh: return s
        except Exception: pass
    return slots[0]

def get_zoom_factor():
    try:
        zoom_str = getattr(config, 'CHART_ZOOM_LEVEL', '90%')
        return int(zoom_str.replace("%", "").strip()) / 100.0
    except Exception: return 0.9

class ModernButton(tk.Button):
    def __init__(self, parent, text, command, primary=True, **kwargs):
        self.default_bg = kwargs.pop('bg', config.COLOR_PRIMARY if primary else config.COLOR_SECONDARY_BTN)
        fg = "white" if primary else config.COLOR_SECONDARY_TEXT
        padx = kwargs.pop('padx', 12)
        pady = kwargs.pop('pady', 4)
        super().__init__(parent, text=text, command=command, bg=self.default_bg, fg=fg, font=config.FONT_BODY_BOLD, relief=tk.FLAT, borderwidth=0, padx=padx, pady=pady, cursor="hand2", **kwargs)
        self.primary = primary
        self.bind("<Enter>", lambda e: self.config(bg=config.COLOR_PRIMARY_HOVER if self.primary else "#e2e8f0"))
        self.bind("<Leave>", lambda e: self.config(bg=self.default_bg))

class ModernToggleButton(tk.Button):
    def __init__(self, parent, text, initial_state=True, on_toggle=None, **kwargs):
        self.is_active = initial_state
        self.on_toggle = on_toggle
        self.active_bg = config.COLOR_PRIMARY
        self.active_fg = "white"
        self.inactive_bg = "#e2e8f0"
        self.inactive_fg = config.COLOR_TEXT_SECONDARY
        super().__init__(parent, text=text, font=config.FONT_BODY_BOLD, relief=tk.FLAT, borderwidth=0, padx=kwargs.pop('padx', 8), pady=kwargs.pop('pady', 2), cursor="hand2", command=self._handle_click, **kwargs)
        self._refresh_style()
        self.bind("<Enter>", self._on_hover)
        self.bind("<Leave>", self._on_leave)

    def _handle_click(self):
        self.is_active = not self.is_active
        self._refresh_style()
        if self.on_toggle: self.on_toggle(self.is_active)

    def set_state(self, state):
        self.is_active = state
        self._refresh_style()

    def get_state(self): return self.is_active

    def _refresh_style(self):
        if self.is_active: self.config(bg=self.active_bg, fg=self.active_fg)
        else: self.config(bg=self.inactive_bg, fg=self.inactive_fg)

    def _on_hover(self, e): self.config(bg=config.COLOR_PRIMARY_HOVER if self.is_active else "#cbd5e1")
    def _on_leave(self, e): self._refresh_style()

class SelectableLabel(tk.Entry):
    def __init__(self, parent, text="", bg="#ffffff", fg="#000000", font=None, width=None, justify=tk.LEFT, **kwargs):
        if width is None: width = len(text)
        super().__init__(parent, bg=bg, fg=fg, font=font, bd=0, highlightthickness=0, relief=tk.FLAT, readonlybackground=bg, width=width, justify=justify, **kwargs)
        self.insert(0, text)
        self.config(state="readonly", cursor="xterm")
        self.xview_moveto(0)

class CalendarDialog(tk.Toplevel):
    def __init__(self, parent, current_date=None, callback=None, anchor_widget=None):
        super().__init__(parent)
        self.callback = callback
        self.current_date = current_date or datetime.now()
        self.display_date = self.current_date
        self.title("Select Date")
        self.geometry("300x320")
        self.configure(bg="white")
        self.resizable(False, False)
        self.transient(parent)
        self.grab_set()
        header = tk.Frame(self, bg=config.COLOR_PRIMARY, height=45)
        header.pack(fill=tk.X)
        self.lbl_month = tk.Label(header, text="", font=config.FONT_SUBTITLE, bg=config.COLOR_PRIMARY, fg="white")
        self.lbl_month.pack(side=tk.LEFT, padx=10, pady=8)
        nav = tk.Frame(header, bg=config.COLOR_PRIMARY)
        nav.pack(side=tk.RIGHT, padx=5)
        tk.Button(nav, text="<", command=self.prev_month, bg=config.COLOR_PRIMARY, fg="white", bd=0).pack(side=tk.LEFT)
        tk.Button(nav, text=">", command=self.next_month, bg=config.COLOR_PRIMARY, fg="white", bd=0).pack(side=tk.LEFT)
        self.cal_frame = tk.Frame(self, bg="white")
        self.cal_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        self.render_calendar()

    def render_calendar(self):
        for w in self.cal_frame.winfo_children(): w.destroy()
        days = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
        for i, d in enumerate(days): tk.Label(self.cal_frame, text=d, font=config.FONT_BODY_BOLD, bg="white", fg="#999").grid(row=0, column=i, pady=(0,4))
        month_days = calendar.monthcalendar(self.display_date.year, self.display_date.month)
        for r, week in enumerate(month_days):
            for c, day in enumerate(week):
                if day == 0: continue
                bg = config.COLOR_PRIMARY if (day == self.current_date.day and self.display_date.month == self.current_date.month and self.display_date.year == self.current_date.year) else "white"
                fg = "white" if bg == config.COLOR_PRIMARY else "#333"
                tk.Button(self.cal_frame, text=str(day), bd=0, bg=bg, fg=fg, width=4, command=lambda d=day: self.select_date(d)).grid(row=r+1, column=c, padx=2, pady=2)
        self.lbl_month.config(text=self.display_date.strftime("%B %Y"))

    def prev_month(self):
        self.display_date = (self.display_date.replace(day=1) - timedelta(days=1))
        self.render_calendar()

    def next_month(self):
        self.display_date = (self.display_date.replace(day=calendar.monthrange(self.display_date.year, self.display_date.month)[1]) + timedelta(days=1))
        self.render_calendar()

    def select_date(self, day):
        if self.callback: self.callback(self.display_date.replace(day=day))
        self.destroy()

class ResultCard(tk.Frame):
    def __init__(self, parent, image_info, callback_click):
        status = image_info.get('status', '-')
        is_ng = (str(status).upper() == "NG" or status not in ("OK", "PASS", "Good", "Normal"))
        status_bg = config.COLOR_STATUS_NG if is_ng else config.COLOR_STATUS_OK
        super().__init__(parent, bg=config.COLOR_CARD_BG, highlightbackground=config.COLOR_CARD_BORDER, highlightthickness=1)
        self.image_info = image_info
        self.callback = callback_click
        
        info_frame = tk.Frame(self, bg=config.COLOR_CARD_BG, padx=8, pady=6)
        info_frame.pack(fill=tk.BOTH, expand=True)

        cat_text = str(image_info.get('category', 'UNK')).replace("exteriorPic", "").strip()
        time_str = image_info['datetime'].strftime("%H:%M:%S") if 'datetime' in image_info and hasattr(image_info['datetime'], 'strftime') else ""
        fname = image_info.get('filename', '')

        # Top row: [Front] [NG / OK badge] [Time]
        top_row = tk.Frame(info_frame, bg=config.COLOR_CARD_BG)
        top_row.pack(fill=tk.X, pady=(0, 4))
        
        lbl_cat = tk.Label(top_row, text=cat_text, font=("Segoe UI", 10, "bold"), bg=config.COLOR_CARD_BG, fg=config.COLOR_TEXT_MAIN)
        lbl_cat.pack(side=tk.LEFT)

        lbl_status = tk.Label(top_row, text=f" {status} ", font=("Segoe UI", 8, "bold"), bg=status_bg, fg="white", padx=4, pady=1)
        lbl_status.pack(side=tk.LEFT, padx=(6, 0))

        if time_str:
            lbl_time = tk.Label(top_row, text=time_str, font=("Segoe UI", 8), bg=config.COLOR_CARD_BG, fg=config.COLOR_TEXT_SECONDARY)
            lbl_time.pack(side=tk.RIGHT)

        # Bottom row: [filename] [🔍 View Image button]
        bot_row = tk.Frame(info_frame, bg=config.COLOR_CARD_BG)
        bot_row.pack(fill=tk.X, pady=(2, 0))

        btn_view = tk.Button(bot_row, text="🖼️ View Image", font=("Segoe UI", 8, "bold"), bg=config.COLOR_PRIMARY, fg="white", 
                             activebackground="#0e4b70", activeforeground="white", relief=tk.FLAT, bd=0, padx=8, pady=2, cursor="hand2", command=self._on_click)
        btn_view.pack(side=tk.RIGHT, padx=(4, 0))

        lbl_fname = SelectableLabel(bot_row, text=fname, font=("Segoe UI", 8), bg=config.COLOR_CARD_BG, fg=config.COLOR_TEXT_SECONDARY)
        lbl_fname.pack(side=tk.LEFT, fill=tk.X, expand=True)

    def _on_click(self): self.callback(self.image_info['path'])

class StringBlackCard(tk.Frame):
    def __init__(self, parent, item, callback_click):
        super().__init__(parent, bg=config.COLOR_CARD_BG, highlightbackground=config.COLOR_CARD_BORDER, highlightthickness=1)
        self.item = item
        self.callback = callback_click

        sn = item.get('sn', 'Unknown')
        station = item.get('station', '')
        line = parse_line_from_station(station) or station
        dt = item.get('datetime')
        if isinstance(dt, datetime):
            time_str = dt.strftime("%H:%M:%S")
        else:
            time_str = str(item.get('datetime_str', ''))[11:19]
            
        dark_pct = item.get('dark_pct', 0.0)
        male = item.get('male', False)
        middle = item.get('middle', False)
        female = item.get('female', False)
        fname = item.get('filename', '')
        fpath = item.get('path', '')

        p = tk.Frame(self, bg=config.COLOR_CARD_BG, padx=10, pady=8)
        p.pack(fill=tk.BOTH, expand=True)

        # Row 1: Line Badge, Station, Timestamp
        r1 = tk.Frame(p, bg=config.COLOR_CARD_BG)
        r1.pack(fill=tk.X, pady=(0, 4))
        tk.Label(r1, text=f" {line} ", font=("Segoe UI", 9, "bold"), bg="#e0f2fe", fg="#0369a1", bd=1, relief=tk.SOLID).pack(side=tk.LEFT)
        tk.Label(r1, text=f"  {station}", font=("Segoe UI", 8), bg=config.COLOR_CARD_BG, fg=config.COLOR_TEXT_SECONDARY).pack(side=tk.LEFT)
        tk.Label(r1, text=time_str, font=("Segoe UI", 8, "bold"), bg=config.COLOR_CARD_BG, fg="#475569").pack(side=tk.RIGHT)

        # Row 2: Serial Number & Dark %
        r2 = tk.Frame(p, bg=config.COLOR_CARD_BG)
        r2.pack(fill=tk.X, pady=(2, 4))
        lbl_sn = SelectableLabel(r2, text=sn, font=("Segoe UI", 11, "bold"), bg=config.COLOR_CARD_BG, fg=config.COLOR_PRIMARY)
        lbl_sn.pack(side=tk.LEFT)

        dark_badge_bg = "#fee2e2" if dark_pct >= 50 else "#fef3c7"
        dark_badge_fg = "#b91c1c" if dark_pct >= 50 else "#b45309"
        lbl_dark = tk.Label(r2, text=f" Dark: {dark_pct}% ", font=("Segoe UI", 8, "bold"), bg=dark_badge_bg, fg=dark_badge_fg, bd=1, relief=tk.SOLID)
        lbl_dark.pack(side=tk.RIGHT)

        # Row 3: Zone Badges (Male, Mid, Female)
        r3 = tk.Frame(p, bg=config.COLOR_CARD_BG)
        r3.pack(fill=tk.X, pady=(2, 6))
        if male:
            tk.Label(r3, text=" Top Zone (Male) ", font=("Segoe UI", 8, "bold"), bg="#fecdd3", fg="#be123c", bd=1, relief=tk.SOLID).pack(side=tk.LEFT, padx=(0, 4))
        if middle:
            tk.Label(r3, text=" Mid Zone ", font=("Segoe UI", 8, "bold"), bg="#fed7aa", fg="#c2410c", bd=1, relief=tk.SOLID).pack(side=tk.LEFT, padx=(0, 4))
        if female:
            tk.Label(r3, text=" Bot Zone (Female) ", font=("Segoe UI", 8, "bold"), bg="#e9d5ff", fg="#7e22ce", bd=1, relief=tk.SOLID).pack(side=tk.LEFT)

        # Row 4: Action button & filename
        r4 = tk.Frame(p, bg=config.COLOR_CARD_BG)
        r4.pack(fill=tk.X, pady=(4, 0))
        btn_view = tk.Button(r4, text="🖼️ Open EL Image", font=("Segoe UI", 8, "bold"), bg=config.COLOR_PRIMARY, fg="white",
                             activebackground="#0e4b70", activeforeground="white", relief=tk.FLAT, bd=0, padx=8, pady=3, cursor="hand2",
                             command=lambda: self.callback(fpath))
        btn_view.pack(side=tk.RIGHT)

        lbl_fn = SelectableLabel(r4, text=fname, font=("Segoe UI", 7), bg=config.COLOR_CARD_BG, fg=config.COLOR_TEXT_SECONDARY)
        lbl_fn.pack(side=tk.LEFT, fill=tk.X, expand=True)

# =================== TOP 5 & NON-EQUIPMENT SCRAP TAB ===================

class Top5AnalyticsTab(tk.Frame):
    def __init__(self, parent, app):
        super().__init__(parent, bg=config.COLOR_BG_MAIN, padx=15, pady=10)
        self.app = app
        self.setup_ui()

    def setup_ui(self):
        top_bar = tk.Frame(self, bg=config.COLOR_BG_MAIN)
        top_bar.pack(fill=tk.X, pady=(0, 10))
        tk.Label(top_bar, text="Top 5 Defects & Non-Equipment Scrap Dashboard", font=config.FONT_TITLE, bg=config.COLOR_BG_MAIN, fg=config.COLOR_TEXT_MAIN).pack(side=tk.LEFT)
        ModernButton(top_bar, text="Refresh Charts", command=self.render_charts, primary=True).pack(side=tk.RIGHT)

        self.canvas_area = tk.Canvas(self, bg="white", highlightthickness=0)
        scrollbar = ttk.Scrollbar(self, orient="vertical", command=self.canvas_area.yview)
        self.scroll_frame = tk.Frame(self.canvas_area, bg="white")
        self.scroll_frame.bind("<Configure>", lambda e: self.canvas_area.config(scrollregion=self.canvas_area.bbox("all")))
        self.canvas_area.create_window((0, 0), window=self.scroll_frame, anchor="nw")
        self.canvas_area.configure(yscrollcommand=scrollbar.set)
        self.canvas_area.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.render_charts()

    def reset_for_new_shift(self, shift_info=None):
        self.render_charts()

    def render_charts(self):
        for w in self.scroll_frame.winfo_children(): w.destroy()
        records = getattr(self.app, 'global_review_records', [])
        if not records:
            tk.Label(self.scroll_frame, text="No review records available for charts yet. Add items in QC Review Tab.", font=config.FONT_SUBTITLE, bg="white", fg=config.COLOR_TEXT_SECONDARY).pack(pady=50, padx=50)
            return

        defect_counts = {}
        defect_line_counts = {}
        non_eq_scrap_counts = {}
        eq_scrap_classes = ["Equipment Scrap 设备报废", "Production Scrap 生产报废"]

        for r in records:
            summary = r.get('summary', 'Unknown')
            line = r.get('line', 'Line 1')
            cls_val = r.get('class', '')
            res = r.get('result', '')

            defect_counts[summary] = defect_counts.get(summary, 0) + 1
            if summary not in defect_line_counts: defect_line_counts[summary] = {}
            defect_line_counts[summary][line] = defect_line_counts[summary].get(line, 0) + 1

            if res.lower() == "scrap" and cls_val not in eq_scrap_classes:
                non_eq_scrap_counts[summary] = non_eq_scrap_counts.get(summary, 0) + 1

        sorted_defects = sorted(defect_counts.items(), key=lambda x: x[1], reverse=True)[:5]
        top_5_names = [d[0] for d in sorted_defects]

        w1 = tk.LabelFrame(self.scroll_frame, text=" Top 5 Defects ", font=config.FONT_BODY_BOLD, bg="white", fg=config.COLOR_PRIMARY, padx=10, pady=10)
        w1.pack(fill=tk.X, padx=10, pady=10)
        c1 = tk.Canvas(w1, width=900, height=240, bg="white", highlightthickness=0)
        c1.pack()
        self._draw_bar_chart(c1, 900, 240, sorted_defects)

        w2 = tk.LabelFrame(self.scroll_frame, text=" Top 5 Defects Detail per Line ", font=config.FONT_BODY_BOLD, bg="white", fg=config.COLOR_PRIMARY, padx=10, pady=10)
        w2.pack(fill=tk.X, padx=10, pady=10)
        c2 = tk.Canvas(w2, width=900, height=280, bg="white", highlightthickness=0)
        c2.pack()
        self._draw_grouped_line_chart(c2, 900, 280, top_5_names, defect_line_counts)

        if non_eq_scrap_counts:
            sorted_non_eq = sorted(non_eq_scrap_counts.items(), key=lambda x: x[1], reverse=True)
            w3 = tk.LabelFrame(self.scroll_frame, text=" Non-Equipment Scrap Breakdown ", font=config.FONT_BODY_BOLD, bg="white", fg=config.COLOR_STATUS_NG, padx=10, pady=10)
            w3.pack(fill=tk.X, padx=10, pady=10)
            c3 = tk.Canvas(w3, width=900, height=220, bg="white", highlightthickness=0)
            c3.pack()
            self._draw_bar_chart(c3, 900, 220, sorted_non_eq, bar_color="#dc2626")

    def _draw_bar_chart(self, canvas, c_w, c_h, data_pairs, bar_color="#1a5b82"):
        pad_l, pad_r, pad_t, pad_b = 40, 30, 20, 50
        plot_w, plot_h = c_w - pad_l - pad_r, c_h - pad_t - pad_b
        if not data_pairs: return
        max_val = max([v for k, v in data_pairs]) if data_pairs else 10
        y_max = max(10, ((max_val // 5) + 1) * 5)
        for i in range(6):
            val = int((y_max / 5) * i)
            y_pos = pad_t + plot_h - (i * (plot_h / 5))
            canvas.create_line(pad_l, y_pos, pad_l + plot_w, y_pos, fill="#f1f5f9", width=1)
            canvas.create_text(pad_l - 8, y_pos, text=str(val), fill="#64748b", font=("Segoe UI", 8), anchor="e")
        canvas.create_line(pad_l, pad_t + plot_h, pad_l + plot_w, pad_t + plot_h, fill="#cbd5e1", width=1.5)
        num_bars = len(data_pairs)
        bar_width = plot_w / max(1, num_bars * 1.8)
        for idx, (name, val) in enumerate(data_pairs):
            bx = pad_l + (idx * (plot_w / num_bars)) + (plot_w / num_bars / 2) - (bar_width / 2)
            bh = (val / y_max) * plot_h
            by = pad_t + plot_h - bh
            canvas.create_rectangle(bx, by, bx + bar_width, pad_t + plot_h, fill=bar_color, outline="")
            canvas.create_text(bx + bar_width/2, by - 8, text=str(val), font=("Segoe UI", 9, "bold"), fill="#334155")
            canvas.create_text(bx + bar_width/2, pad_t + plot_h + 15, text=name, font=("Segoe UI", 9), fill="#334155", angle=15, anchor="ne")

    def _draw_grouped_line_chart(self, canvas, c_w, c_h, top_5_names, defect_line_counts):
        pad_l, pad_r, pad_t, pad_b = 40, 150, 20, 40
        plot_w, plot_h = c_w - pad_l - pad_r, c_h - pad_t - pad_b
        lines = [f"Line {i}" for i in range(1, 8)]
        max_val = max([max(defect_line_counts.get(d, {}).values() or [0]) for d in top_5_names] or [1])
        y_max = max(8, ((max_val // 2) + 1) * 2)
        for i in range(5):
            val = int((y_max / 4) * i)
            y_pos = pad_t + plot_h - (i * (plot_h / 4))
            canvas.create_line(pad_l, y_pos, pad_l + plot_w, y_pos, fill="#f1f5f9", width=1)
            canvas.create_text(pad_l - 8, y_pos, text=str(val), fill="#64748b", font=("Segoe UI", 8), anchor="e")
        canvas.create_line(pad_l, pad_t + plot_h, pad_l + plot_w, pad_t + plot_h, fill="#cbd5e1", width=1.5)
        group_w = plot_w / len(lines)
        colors = ["#1a5b82", "#e8702a", "#64748b", "#ca8a04", "#0284c7"]
        for l_idx, l_name in enumerate(lines):
            g_center = pad_l + (l_idx * group_w) + (group_w / 2)
            bar_w = max(3, group_w / 6.5)
            start_x = g_center - ((len(top_5_names) * bar_w) / 2)
            for d_idx, d_name in enumerate(top_5_names):
                val = defect_line_counts.get(d_name, {}).get(l_name, 0)
                bx = start_x + (d_idx * bar_w)
                if val > 0:
                    bh = (val / y_max) * plot_h
                    by = pad_t + plot_h - bh
                    canvas.create_rectangle(bx, by, bx + bar_w, pad_t + plot_h, fill=colors[d_idx % len(colors)], outline="")
                    canvas.create_text(bx + bar_w/2, by - 6, text=str(val), font=("Segoe UI", 7, "bold"), fill="#334155")
            canvas.create_text(g_center, pad_t + plot_h + 12, text=l_name, font=("Segoe UI", 9, "bold"), fill="#334155")
        lx, ly = pad_l + plot_w + 15, pad_t + 20
        for d_idx, d_name in enumerate(top_5_names):
            col = colors[d_idx % len(colors)]
            canvas.create_rectangle(lx, ly, lx + 10, ly + 10, fill=col, outline="")
            canvas.create_text(lx + 16, ly + 5, text=d_name, font=("Segoe UI", 8), fill="#334155", anchor="w")
            ly += 22


# =================== MES CHINESE FIELD GUIDE & TRANSLATOR DIALOG ===================

class MESChineseFieldGuideDialog(tk.Toplevel):
    """
    Dedicated Desktop Dialog providing comprehensive Chinese <-> English MES dictionary,
    field mappings for FineReport Page 1/2 (组件生产流转), and an interactive
    live text translator tool.
    """
    def __init__(self, parent):
        super().__init__(parent)
        self.title("Factory MES Chinese ↔ English Field Guide & Translator")
        self.geometry("1020x750")
        self.minsize(820, 600)
        self.configure(bg=config.COLOR_BG_MAIN)
        try:
            self.transient(parent.winfo_toplevel() if hasattr(parent, 'winfo_toplevel') else parent)
        except Exception:
            pass

        self.dict_items = list(MES_CHINESE_DICTIONARY)
        self.filtered_items = list(self.dict_items)
        self.current_cat = "All"
        self.sort_col = "cn"
        self.sort_reverse = False

        self.setup_ui()
        self.center_window()
        self.populate_tree()

    def center_window(self):
        self.update_idletasks()
        w = self.winfo_width()
        h = self.winfo_height()
        sw = self.winfo_screenwidth()
        sh = self.winfo_screenheight()
        x = max(0, (sw - w) // 2)
        y = max(0, (sh - h) // 2)
        self.geometry(f"{w}x{h}+{x}+{y}")

    def setup_ui(self):
        # 1. Header Banner
        header = tk.Frame(self, bg=config.COLOR_BG_MAIN, padx=16, pady=10)
        header.pack(fill=tk.X)

        tk.Label(
            header,
            text="🇨🇳/🇺🇸 Factory MES Chinese ↔ English Field Guide & Live Translator",
            font=config.FONT_TITLE,
            bg=config.COLOR_BG_MAIN,
            fg=config.COLOR_TEXT_MAIN
        ).pack(anchor=tk.W)

        tk.Label(
            header,
            text="Comprehensive field mapping & terminology for FineReport Page 1/2 (组件生产流转) and Page 2/2 (Module Product Process Logsheet)",
            font=config.FONT_SMALL,
            bg=config.COLOR_BG_MAIN,
            fg=config.COLOR_TEXT_SECONDARY
        ).pack(anchor=tk.W, pady=(2, 0))

        # 2. Interactive Live Text Translator Box
        trans_frame = tk.Frame(self, bg="white", highlightbackground=config.COLOR_DIVIDER, highlightthickness=1, padx=14, pady=10)
        trans_frame.pack(fill=tk.X, padx=16, pady=(0, 10))

        t_top = tk.Frame(trans_frame, bg="white")
        t_top.pack(fill=tk.X)

        tk.Label(
            t_top,
            text="🌐 Instant Chinese Text Translator",
            font=("Segoe UI", 10, "bold"),
            bg="white",
            fg="#2563eb"
        ).pack(side=tk.LEFT)

        tk.Label(
            t_top,
            text="Paste or type any Chinese text from MES (e.g. M12工序NG自动Hold, 上层3号位, 接线盒打胶 OK, 耐压 合格):",
            font=("Segoe UI", 8),
            bg="white",
            fg="#64748b"
        ).pack(side=tk.LEFT, padx=(10, 0))

        t_row = tk.Frame(trans_frame, bg="white")
        t_row.pack(fill=tk.X, pady=(6, 0))

        self.ent_translate = tk.Entry(t_row, font=("Segoe UI", 10), bg=config.COLOR_INPUT_BG, relief=tk.FLAT, bd=4)
        self.ent_translate.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 6))
        self.ent_translate.bind("<KeyRelease>", lambda e: self.on_live_translate())

        ModernButton(t_row, text="📋 Paste", command=self.paste_to_translator, primary=False, padx=8, pady=2).pack(side=tk.LEFT, padx=(0, 4))
        ModernButton(t_row, text="Translate", command=self.on_live_translate, primary=True, padx=8, pady=2).pack(side=tk.LEFT, padx=(0, 4))
        ModernButton(t_row, text="Clear", command=self.clear_translator, primary=False, padx=6, pady=2).pack(side=tk.LEFT)

        # Translation Result Box
        self.res_frame = tk.Frame(trans_frame, bg="#f0fdf4", highlightbackground="#86efac", highlightthickness=1, padx=10, pady=6)
        self.res_frame.pack(fill=tk.X, pady=(8, 0))

        tk.Label(self.res_frame, text="English Meaning:", font=("Segoe UI", 9, "bold"), bg="#f0fdf4", fg="#166534").pack(side=tk.LEFT, padx=(0, 6))
        self.lbl_trans_result = tk.Label(self.res_frame, text="Type or paste Chinese text above to translate instantly", font=("Segoe UI", 9), bg="#f0fdf4", fg="#15803d", wraplength=700, justify=tk.LEFT)
        self.lbl_trans_result.pack(side=tk.LEFT, fill=tk.X, expand=True)

        ModernButton(self.res_frame, text="📋 Copy", command=self.copy_translation_result, primary=False, padx=6, pady=1).pack(side=tk.RIGHT)

        # 3. Search & Filter Bar for Field Dictionary Table
        ctrl_frame = tk.Frame(self, bg="white", highlightbackground=config.COLOR_DIVIDER, highlightthickness=1, padx=12, pady=6)
        ctrl_frame.pack(fill=tk.X, padx=16, pady=(0, 8))

        tk.Label(ctrl_frame, text="🔍 Search Dictionary:", font=("Segoe UI", 9, "bold"), bg="white", fg="#1e293b").pack(side=tk.LEFT, padx=(0, 5))
        self.ent_search = tk.Entry(ctrl_frame, font=("Segoe UI", 9), width=20, bg=config.COLOR_INPUT_BG, relief=tk.FLAT, bd=3)
        self.ent_search.pack(side=tk.LEFT, padx=(0, 8))
        self.ent_search.bind("<KeyRelease>", lambda e: self.on_search_filter_changed())

        # Category buttons
        cat_box = tk.Frame(ctrl_frame, bg="white")
        cat_box.pack(side=tk.LEFT)

        self.cat_buttons = {}
        cats = [("All", "All"), ("Header", "Headers"), ("Front-End", "Front-End"), ("Back-End", "Back-End"), ("Electrical", "Electrical"), ("Status", "Status")]
        for cat_key, cat_label in cats:
            btn = tk.Button(
                cat_box,
                text=cat_label,
                font=("Segoe UI", 8, "bold" if cat_key == "All" else "normal"),
                bg="#2563eb" if cat_key == "All" else "#f1f5f9",
                fg="white" if cat_key == "All" else "#334155",
                activebackground="#1d4ed8",
                activeforeground="white",
                relief=tk.FLAT,
                bd=1,
                padx=8,
                pady=2,
                cursor="hand2",
                command=lambda k=cat_key: self.set_category_filter(k)
            )
            btn.pack(side=tk.LEFT, padx=2)
            self.cat_buttons[cat_key] = btn

        ModernButton(ctrl_frame, text="Reset", command=self.reset_filter, primary=False, padx=6, pady=2).pack(side=tk.LEFT, padx=(8, 0))

        self.lbl_dict_count = tk.Label(ctrl_frame, text="Showing: 0 / 0", font=("Segoe UI", 9, "bold"), bg="white", fg=config.COLOR_PRIMARY)
        self.lbl_dict_count.pack(side=tk.RIGHT)

        # 4. Dictionary Table (ttk.Treeview)
        tbl_frame = tk.Frame(self, bg="white", highlightbackground=config.COLOR_DIVIDER, highlightthickness=1)
        tbl_frame.pack(fill=tk.BOTH, expand=True, padx=16, pady=(0, 8))

        columns = [
            ("cn", "Chinese Term (中文名称)", 140, "w"),
            ("en", "English Translation (英文翻译)", 220, "w"),
            ("cat", "Category (类别)", 135, "w"),
            ("equipment", "Equipment Prefix (设备编号前缀)", 175, "w"),
            ("page", "FineReport Page", 100, "center"),
            ("desc", "Factory Context & Real Example (说明与示例)", 320, "w")
        ]

        col_ids = [c[0] for c in columns]
        self.tree = ttk.Treeview(tbl_frame, columns=col_ids, show="headings", selectmode="browse")

        v_scroll = ttk.Scrollbar(tbl_frame, orient="vertical", command=self.tree.yview)
        h_scroll = ttk.Scrollbar(tbl_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=v_scroll.set, xscrollcommand=h_scroll.set)

        self.tree.grid(row=0, column=0, sticky="nsew")
        v_scroll.grid(row=0, column=1, sticky="ns")
        h_scroll.grid(row=1, column=0, sticky="ew")
        tbl_frame.grid_rowconfigure(0, weight=1)
        tbl_frame.grid_columnconfigure(0, weight=1)

        for col_id, col_text, col_w, col_anchor in columns:
            self.tree.heading(col_id, text=col_text, command=lambda c=col_id: self.sort_by(c))
            self.tree.column(col_id, width=col_w, anchor=col_anchor, minwidth=60)

        self.tree.tag_configure("even", background="#ffffff")
        self.tree.tag_configure("odd", background="#f8fafc")
        self.tree.bind("<Double-1>", self.on_row_double_clicked)

        # 5. Reference Sample Breakdown Box
        ref_card = tk.Frame(self, bg="#f8fafc", highlightbackground=config.COLOR_DIVIDER, highlightthickness=1, padx=12, pady=6)
        ref_card.pack(fill=tk.X, padx=16, pady=(0, 10))

        tk.Label(
            ref_card,
            text="📌 Real Screenshot Breakdown (Module V01269003041442 on Page 1/2):",
            font=("Segoe UI", 8, "bold"),
            bg="#f8fafc",
            fg="#1e293b"
        ).pack(anchor=tk.W)

        ref_txt = (
            "• 焊接 (Welding): Stringer203 (TUMSOLERING1009) (2026-09-04 20:14) | "
            "• 敷设 (Lay up): TUMLAYUP1004 (2026-09-05 07:27) | "
            "• 层压 (Lamination): Lam18.2 (上层3号位) (2026-09-06 00:07)\n"
            "• 料号 (Part/Lot No): 6A024170 | "
            "• 等级 (Appearance/Final Grade): Q3 | "
            "• 当前工序 (Current Station): M12 (M12工序NG自动Hold ➔ M12 Process NG Auto-Hold)"
        )
        tk.Label(
            ref_card,
            text=ref_txt,
            font=("Consolas", 8),
            bg="#f8fafc",
            fg="#475569",
            justify=tk.LEFT
        ).pack(anchor=tk.W, pady=(2, 0))

        # 6. Bottom Action Buttons
        bot_bar = tk.Frame(self, bg=config.COLOR_BG_MAIN, padx=16)
        bot_bar.pack(fill=tk.X, pady=(0, 10))

        tk.Label(
            bot_bar,
            text="💡 Double-click any row to copy translation to clipboard",
            font=config.FONT_SMALL,
            bg=config.COLOR_BG_MAIN,
            fg=config.COLOR_TEXT_SECONDARY
        ).pack(side=tk.LEFT)

        b_right = tk.Frame(bot_bar, bg=config.COLOR_BG_MAIN)
        b_right.pack(side=tk.RIGHT)

        ModernButton(b_right, text="📋 Copy Selected Term", command=self.copy_selected_term, primary=False).pack(side=tk.LEFT, padx=4)
        ModernButton(b_right, text="📥 Copy Full Table (TSV)", command=self.copy_all_tsv, primary=False).pack(side=tk.LEFT, padx=4)
        ModernButton(b_right, text="Close", command=self.destroy, primary=True).pack(side=tk.LEFT, padx=4)

    def on_live_translate(self, event=None):
        txt = self.ent_translate.get().strip()
        if not txt:
            self.lbl_trans_result.config(text="Type or paste Chinese text above to translate instantly", fg="#64748b")
            return
        translated = translate_mes_text(txt)
        if translated:
            self.lbl_trans_result.config(text=translated, fg="#15803d")
        else:
            self.lbl_trans_result.config(text="No direct translation match found.", fg="#b45309")

    def paste_to_translator(self):
        try:
            txt = self.clipboard_get()
            if txt:
                self.ent_translate.delete(0, tk.END)
                self.ent_translate.insert(0, txt.strip())
                self.on_live_translate()
        except Exception:
            pass

    def clear_translator(self):
        self.ent_translate.delete(0, tk.END)
        self.on_live_translate()

    def copy_translation_result(self):
        res = self.lbl_trans_result.cget("text")
        if res and "Type or paste" not in res:
            self.clipboard_clear()
            self.clipboard_append(res)
            messagebox.showinfo("Copied", f"Copied translation to clipboard:\n\n{res}")

    def set_category_filter(self, cat):
        self.current_cat = cat
        for k, btn in self.cat_buttons.items():
            if k == cat:
                btn.config(bg="#2563eb", fg="white", font=("Segoe UI", 8, "bold"))
            else:
                btn.config(bg="#f1f5f9", fg="#334155", font=("Segoe UI", 8, "normal"))
        self.apply_filter()

    def on_search_filter_changed(self):
        self.apply_filter()

    def reset_filter(self):
        self.ent_search.delete(0, tk.END)
        self.set_category_filter("All")

    def apply_filter(self):
        q = self.ent_search.get().strip().lower()
        items = list(self.dict_items)

        if self.current_cat != "All":
            items = [item for item in items if self.current_cat.lower() in item.get('cat', '').lower()]

        if q:
            items = [
                item for item in items if (
                    q in item.get('cn', '').lower() or
                    q in item.get('en', '').lower() or
                    q in item.get('cat', '').lower() or
                    q in item.get('equipment', '').lower() or
                    q in item.get('desc', '').lower()
                )
            ]

        # Sorting
        items.sort(key=lambda x: str(x.get(self.sort_col, '')).lower(), reverse=self.sort_reverse)
        self.filtered_items = items
        self.populate_tree()

    def sort_by(self, col):
        if self.sort_col == col:
            self.sort_reverse = not self.sort_reverse
        else:
            self.sort_col = col
            self.sort_reverse = False
        self.apply_filter()

    def populate_tree(self):
        for row_id in self.tree.get_children():
            self.tree.delete(row_id)

        for i, item in enumerate(self.filtered_items):
            tag = "even" if i % 2 == 0 else "odd"
            vals = (
                item.get('cn', '-'),
                item.get('en', '-'),
                item.get('cat', '-'),
                item.get('equipment', '-'),
                item.get('page', '-'),
                item.get('desc', '-')
            )
            self.tree.insert("", tk.END, values=vals, tags=(tag,))

        self.lbl_dict_count.config(text=f"Showing: {len(self.filtered_items)} / {len(self.dict_items)}")

    def on_row_double_clicked(self, event):
        sel = self.tree.selection()
        if not sel: return
        vals = self.tree.item(sel[0]).get('values', [])
        if len(vals) >= 2:
            copy_txt = f"{vals[0]} ({vals[1]})"
            self.clipboard_clear()
            self.clipboard_append(copy_txt)
            messagebox.showinfo("Copied", f"Copied to clipboard:\n\n{copy_txt}")

    def copy_selected_term(self):
        self.on_row_double_clicked(None)

    def copy_all_tsv(self):
        lines = ["Chinese\tEnglish Translation\tCategory\tEquipment\tPage\tDescription"]
        for item in self.dict_items:
            lines.append(f"{item.get('cn','')}\t{item.get('en','')}\t{item.get('cat','')}\t{item.get('equipment','')}\t{item.get('page','')}\t{item.get('desc','')}")
        tsv_data = "\n".join(lines)
        self.clipboard_clear()
        self.clipboard_append(tsv_data)
        messagebox.showinfo("Copied All", f"Full MES dictionary ({len(self.dict_items)} terms) copied to clipboard as TSV!")


# =================== MES PROCESS LOG TAB ===================

class MESProcessLogTab(tk.Frame):
    """
    Dedicated Laptop Display Tab for MES Process Log & Machine Trend Tracker.
    Separates factory Electronic Transfer Order activity (TUMSOLDERING, TUMLAYUP, TUMLAMINATION)
    completely from Module Review defect grading.
    """
    def __init__(self, parent, app):
        super().__init__(parent, bg=config.COLOR_BG_MAIN, padx=15, pady=10)
        self.app = app
        self.records = []
        self.filtered_records = []
        self.sort_col = "time"
        self.sort_reverse = True
        self.setup_ui()
        self.load_records()
        self.last_trend_mtime = 0
        self._poll_log_file()

    def setup_ui(self):
        # 1. Top Bar: Title, Subtitle, and Global Action Buttons
        top_bar = tk.Frame(self, bg=config.COLOR_BG_MAIN)
        top_bar.pack(fill=tk.X, pady=(0, 8))

        title_frame = tk.Frame(top_bar, bg=config.COLOR_BG_MAIN)
        title_frame.pack(side=tk.LEFT, fill=tk.Y)
        
        tk.Label(
            title_frame,
            text="📋 Factory MES Process Log (Electronic Transfer Order)",
            font=config.FONT_TITLE,
            bg=config.COLOR_BG_MAIN,
            fg=config.COLOR_TEXT_MAIN
        ).pack(anchor=tk.W)

        self.lbl_server_status = tk.Label(
            title_frame,
            text="Host: 10.200.3.109:8080 (FineReport Decision Platform) | Storage: mes_process_trend_log.json",
            font=config.FONT_SMALL,
            bg=config.COLOR_BG_MAIN,
            fg=config.COLOR_TEXT_SECONDARY
        )
        self.lbl_server_status.pack(anchor=tk.W, pady=(2, 0))

        btn_box = tk.Frame(top_bar, bg=config.COLOR_BG_MAIN)
        btn_box.pack(side=tk.RIGHT)

        ModernButton(btn_box, text="🌐 MES Active Tab", command=self.open_mes_portal, primary=True).pack(side=tk.LEFT, padx=3)
        ModernButton(btn_box, text="🔐 Login (030888)", command=self.open_mes_login, primary=False).pack(side=tk.LEFT, padx=3)
        ModernButton(btn_box, text="🇨🇳 Chinese Guide", command=self.open_chinese_guide, primary=False).pack(side=tk.LEFT, padx=3)
        ModernButton(btn_box, text="📁 MES Photos", command=self.open_mes_photos_folder, primary=False).pack(side=tk.LEFT, padx=3)
        ModernButton(btn_box, text="🔄 Refresh Log", command=self.refresh_log, primary=False).pack(side=tk.LEFT, padx=3)
        ModernButton(btn_box, text="🔄 New Shift Reset", command=lambda: self.reset_for_new_shift(get_operational_shift_info()), primary=False).pack(side=tk.LEFT, padx=3)
        ModernButton(btn_box, text="📦 View Archive", command=self.open_archive_log, primary=False).pack(side=tk.LEFT, padx=3)
        ModernButton(btn_box, text="📥 Export CSV", command=self.export_csv, primary=True).pack(side=tk.LEFT, padx=3)
        ModernButton(btn_box, text="🗑️ Clear Log", command=self.clear_log, primary=False).pack(side=tk.LEFT, padx=3)

        # 2. Top KPI Metric Cards Frame
        kpi_frame = tk.Frame(self, bg=config.COLOR_BG_MAIN)
        kpi_frame.pack(fill=tk.X, pady=(0, 10))

        self.kpi_cards = {}
        cards_def = [
            ("total", "Total Logged Modules", "#2563eb", "0", "Panels recorded"),
            ("soldering", "Welding (Stringer)", "#d97706", "-", "Top stringer line"),
            ("layup", "Lay up (TUMLAYUP)", "#059669", "-", "Top layup machine"),
            ("lamination", "Lamination (Lam)", "#7c3aed", "-", "Top lamination deck")
        ]
        for key, title, col, init_val, sub in cards_def:
            card = tk.Frame(kpi_frame, bg="white", highlightbackground=config.COLOR_DIVIDER, highlightthickness=1, padx=12, pady=8)
            card.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=3)
            
            lbl_title = tk.Label(card, text=title, font=("Segoe UI", 9, "bold"), bg="white", fg=col)
            lbl_title.pack(anchor=tk.W)

            lbl_val = tk.Label(card, text=init_val, font=("Segoe UI", 14, "bold"), bg="white", fg="#0f172a")
            lbl_val.pack(anchor=tk.W, pady=(2, 2))

            lbl_sub = tk.Label(card, text=sub, font=("Segoe UI", 8), bg="white", fg="#64748b")
            lbl_sub.pack(anchor=tk.W)

            self.kpi_cards[key] = (lbl_val, lbl_sub)

        # 3. Control Action Bar: Direct SN Query & Real-Time Filter
        ctrl_bar = tk.Frame(self, bg="white", highlightbackground=config.COLOR_DIVIDER, highlightthickness=1, padx=10, pady=6)
        ctrl_bar.pack(fill=tk.X, pady=(0, 8))

        # Direct Laptop Query
        q_frame = tk.Frame(ctrl_bar, bg="white")
        q_frame.pack(side=tk.LEFT, fill=tk.Y)

        tk.Label(q_frame, text="🔍 Query MES SN:", font=("Segoe UI", 9, "bold"), bg="white", fg="#1e293b").pack(side=tk.LEFT, padx=(0, 5))
        self.ent_query_sn = tk.Entry(q_frame, font=("Segoe UI", 10, "bold"), width=18, bg=config.COLOR_INPUT_BG, relief=tk.FLAT, bd=3)
        self.ent_query_sn.pack(side=tk.LEFT, padx=(0, 6))
        self.ent_query_sn.bind("<Return>", lambda e: self.query_sn())

        ModernButton(q_frame, text="Search & Query", command=self.query_sn, primary=True, padx=8, pady=2).pack(side=tk.LEFT, padx=(0, 8))
        self.lbl_query_status = tk.Label(q_frame, text="Ready", font=("Segoe UI", 9), bg="white", fg="#64748b")
        self.lbl_query_status.pack(side=tk.LEFT)

        # Filter on Right
        f_frame = tk.Frame(ctrl_bar, bg="white")
        f_frame.pack(side=tk.RIGHT, fill=tk.Y)

        tk.Label(f_frame, text="Filter Records:", font=("Segoe UI", 9), bg="white", fg="#1e293b").pack(side=tk.LEFT, padx=(0, 5))
        self.ent_filter = tk.Entry(f_frame, font=("Segoe UI", 9), width=18, bg=config.COLOR_INPUT_BG, relief=tk.FLAT, bd=3)
        self.ent_filter.pack(side=tk.LEFT, padx=(0, 6))
        self.ent_filter.bind("<KeyRelease>", lambda e: self.on_filter_changed())

        ModernButton(f_frame, text="Clear", command=self.clear_filter, primary=False, padx=6, pady=2).pack(side=tk.LEFT, padx=(0, 8))
        self.lbl_count_tag = tk.Label(f_frame, text="Showing: 0 / 0", font=("Segoe UI", 9, "bold"), bg="white", fg=config.COLOR_PRIMARY)
        self.lbl_count_tag.pack(side=tk.LEFT)

        # 4. Table Container with ttk.Treeview
        table_container = tk.Frame(self, bg="white", highlightbackground=config.COLOR_DIVIDER, highlightthickness=1)
        table_container.pack(fill=tk.BOTH, expand=True)

        self.columns = [
            ("time", "Timestamp", 135, "center"),
            ("sn", "Serial Number (V01)", 155, "center"),
            ("soldering", "Welding (Stringer)", 155, "w"),
            ("layup", "Lay up (TUMLAYUP)", 145, "w"),
            ("lamination", "Lamination (Lam)", 165, "w"),
            ("family", "Product Family", 145, "w"),
            ("lot", "Lot No", 95, "center"),
            ("mo", "MO No", 100, "center"),
            ("grade", "Grade", 70, "center"),
            ("defect", "Defect / Note", 130, "w"),
            ("result", "Result", 75, "center")
        ]

        col_ids = [c[0] for c in self.columns]
        self.tree = ttk.Treeview(table_container, columns=col_ids, show="headings", selectmode="browse")

        style = ttk.Style()
        style.configure("MESTree.Treeview", font=("Segoe UI", 9), rowheight=27)
        style.configure("MESTree.Treeview.Heading", font=("Segoe UI", 9, "bold"))
        self.tree.configure(style="MESTree.Treeview")

        v_scroll = ttk.Scrollbar(table_container, orient="vertical", command=self.tree.yview)
        h_scroll = ttk.Scrollbar(table_container, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=v_scroll.set, xscrollcommand=h_scroll.set)

        self.tree.grid(row=0, column=0, sticky="nsew")
        v_scroll.grid(row=0, column=1, sticky="ns")
        h_scroll.grid(row=1, column=0, sticky="ew")
        table_container.grid_rowconfigure(0, weight=1)
        table_container.grid_columnconfigure(0, weight=1)

        for col_id, col_text, col_w, col_anchor in self.columns:
            self.tree.heading(col_id, text=col_text, command=lambda c=col_id: self.sort_by_column(c))
            self.tree.column(col_id, width=col_w, anchor=col_anchor, minwidth=60)

        # Row tag colors
        self.tree.tag_configure("complete", background="#f0fdf4", foreground="#065f46")   # All 3 machines present
        self.tree.tag_configure("partial", background="#fefce8", foreground="#854d0e")    # 1-2 machines present
        self.tree.tag_configure("blank", background="#fef2f2", foreground="#991b1b")      # 0 machines present

        self.tree.bind("<Double-1>", self.on_row_double_clicked)
        self.tree.bind("<Button-3>", self.show_context_menu)

        # Context Menu
        self.context_menu = tk.Menu(self, tearoff=0)
        self.context_menu.add_command(label="📋 Copy Serial Number", command=self.copy_selected_sn)
        self.context_menu.add_command(label="📋 Copy All Record Details", command=self.copy_selected_row)
        self.context_menu.add_separator()
        self.context_menu.add_command(label="🖼️ View Saved MES Photo", command=self.view_selected_mes_photo)
        self.context_menu.add_command(label="📁 Open MES Photos Folder", command=self.open_mes_photos_folder)
        self.context_menu.add_separator()
        self.context_menu.add_command(label="🇨🇳 Chinese MES Field Guide & Translator", command=self.open_chinese_guide)
        self.context_menu.add_command(label="🔍 Re-query MES for this SN", command=self.requery_selected_sn)
        self.context_menu.add_separator()
        self.context_menu.add_command(label="❌ Delete from Log", command=self.delete_selected_record)

        # 5. Footer Bar
        footer = tk.Frame(self, bg=config.COLOR_BG_MAIN)
        footer.pack(fill=tk.X, pady=(6, 0))

        self.lbl_footer = tk.Label(
            footer,
            text="💡 Double-click row to copy SN | Click '🇨🇳 Chinese Guide' for Page 1/2 field glossary | Synchronized in real-time with Mobile Hub",
            font=config.FONT_SMALL,
            bg=config.COLOR_BG_MAIN,
            fg=config.COLOR_TEXT_SECONDARY
        )
        self.lbl_footer.pack(side=tk.LEFT)

    def load_records(self):
        self.records = load_mes_trend_log()
        self.apply_filter()
        self.update_kpi_cards()

    def refresh_log(self):
        self.load_records()
        self.lbl_query_status.config(text="Log refreshed", fg="#059669")

    def _poll_log_file(self):
        try:
            if os.path.exists(MES_TREND_FILE):
                mtime = os.path.getmtime(MES_TREND_FILE)
                if mtime > self.last_trend_mtime:
                    self.last_trend_mtime = mtime
                    self.load_records()
        except Exception:
            pass
        self.after(2000, self._poll_log_file)

    def handle_incoming_mes_record(self, data):
        if not data or not isinstance(data, dict): return
        sn = data.get('sn', '').strip()
        if not sn: return

        # Update in-memory records
        existing_idx = None
        for i, r in enumerate(self.records):
            if r.get('sn') == sn:
                existing_idx = i
                break
        
        if existing_idx is not None:
            self.records[existing_idx].update(data)
        else:
            self.records.insert(0, data)

        self.apply_filter()
        self.update_kpi_cards()
        self.lbl_query_status.config(text=f"✅ Logged {sn}", fg="#059669")

    def on_filter_changed(self, event=None):
        self.apply_filter()

    def clear_filter(self):
        self.ent_filter.delete(0, tk.END)
        self.apply_filter()

    def open_archive_log(self):
        archive_file = os.path.join(config.LOCAL_DATA_DIR, "mes_process_trend_archive.json")
        if not os.path.exists(archive_file):
            messagebox.showinfo("MES Archive", "No archived MES records found yet.")
            return
        try:
            os.startfile(archive_file)
        except Exception:
            try:
                subprocess.Popen(["notepad.exe", archive_file])
            except Exception as e:
                messagebox.showerror("Error", f"Could not open archive file:\n{e}")

    def reset_for_new_shift(self, shift_info=None):
        if hasattr(self, 'ent_query_sn'):
            self.ent_query_sn.delete(0, tk.END)
        if hasattr(self, 'ent_filter'):
            self.ent_filter.delete(0, tk.END)

        # Archive records from previous shift so active workspace starts clean
        if self.records:
            try:
                archive_file = os.path.join(config.LOCAL_DATA_DIR, "mes_process_trend_archive.json")
                archived = []
                if os.path.exists(archive_file):
                    try:
                        with open(archive_file, 'r', encoding='utf-8') as af:
                            archived = json.load(af)
                    except Exception:
                        archived = []
                existing_sns = {r.get('sn') for r in archived if r.get('sn')}
                for r in self.records:
                    if r.get('sn') and r.get('sn') not in existing_sns:
                        archived.append(r)
                with open(archive_file, 'w', encoding='utf-8') as af:
                    json.dump(archived, af, ensure_ascii=False, indent=2, default=str)
            except Exception as e:
                print(f"[MES ARCHIVE ERROR]: {e}")

        # Clear active trend log for the fresh shift
        try:
            with open(MES_TREND_FILE, 'w', encoding='utf-8') as f:
                json.dump([], f, ensure_ascii=False, indent=2)
        except Exception:
            pass

        self.records = []
        self.apply_filter()
        self.update_kpi_cards()
        shift_label = shift_info.get('name', 'New Shift') if shift_info else 'New Shift'
        if hasattr(self, 'lbl_query_status'):
            self.lbl_query_status.config(text=f"Shift Reset: {shift_label} (Log archived & refreshed)", fg="#059669")

    def apply_filter(self):
        query = self.ent_filter.get().strip().lower()
        if not query:
            filtered = list(self.records)
        else:
            filtered = []
            for r in self.records:
                combined = " ".join([
                    str(r.get('sn', '')),
                    str(r.get('tumsoldering', '')),
                    str(r.get('tumlayup', '')),
                    str(r.get('tumlamination', '')),
                    str(r.get('product_family', '')),
                    str(r.get('lot_no', '')),
                    str(r.get('mo_no', '')),
                    str(r.get('appearance_grade', '')),
                    str(r.get('defect', '')),
                    str(r.get('result', ''))
                ]).lower()
                if query in combined:
                    filtered.append(r)

        # Sorting
        def _sort_key(item):
            val = item.get(self.sort_col, '')
            if self.sort_col == 'time': val = item.get('timestamp') or item.get('layup_time', '')
            return str(val).lower()

        filtered.sort(key=_sort_key, reverse=self.sort_reverse)
        self.filtered_records = filtered

        # Populate tree
        for row_id in self.tree.get_children():
            self.tree.delete(row_id)

        for r in self.filtered_records:
            s = str(r.get('tumsoldering', '')).strip()
            ly = str(r.get('tumlayup', '')).strip()
            lm = str(r.get('tumlamination', '')).strip()

            if s and ly and lm:
                tag = "complete"
            elif s or ly or lm:
                tag = "partial"
            else:
                tag = "blank"

            def_val = str(r.get('defect') or r.get('current_step') or '-').strip()
            if def_val and def_val != '-' and any('\u4e00' <= char <= '\u9fff' for char in def_val):
                def_val = translate_mes_text(def_val)

            vals = (
                r.get('layup_time') or r.get('timestamp', '-'),
                r.get('sn', '-'),
                s or '(Blank)',
                ly or '(Blank)',
                lm or '(Blank)',
                r.get('product_family', '-'),
                r.get('lot_no', '-'),
                r.get('mo_no', '-'),
                r.get('appearance_grade', '-'),
                def_val,
                r.get('result', '-')
            )
            self.tree.insert("", tk.END, values=vals, tags=(tag,))

        self.lbl_count_tag.config(text=f"Showing: {len(self.filtered_records)} / {len(self.records)}")

    def sort_by_column(self, col_id):
        if self.sort_col == col_id:
            self.sort_reverse = not self.sort_reverse
        else:
            self.sort_col = col_id
            self.sort_reverse = False
        self.apply_filter()

    def update_kpi_cards(self):
        analytics = get_mes_trend_analytics()
        tot = analytics.get('total_logged', 0)
        
        lbl_tot_val, lbl_tot_sub = self.kpi_cards.get('total', (None, None))
        if lbl_tot_val:
            lbl_tot_val.config(text=str(tot))
            lbl_tot_sub.config(text=f"{len(self.filtered_records)} matching filter" if len(self.filtered_records) != tot else "Panels recorded")

        def _set_kpi(key, top_list, title_name):
            val_el, sub_el = self.kpi_cards.get(key, (None, None))
            if not val_el: return
            if top_list and len(top_list) > 0:
                top_item = top_list[0]
                val_el.config(text=f"{top_item['name']} ({top_item['pct']}%)")
                sub_el.config(text=f"{top_item['count']} panels | {len(top_list)} machine(s) in use")
            else:
                val_el.config(text="-")
                sub_el.config(text=f"No {title_name} data yet")

        _set_kpi('soldering', analytics.get('soldering_top', []), "welding")
        _set_kpi('layup', analytics.get('layup_top', []), "lay up")
        _set_kpi('lamination', analytics.get('lamination_top', []), "lamination")

    def query_sn(self):
        sn = self.ent_query_sn.get().strip().upper()
        if not sn:
            self.lbl_query_status.config(text="Please enter SN", fg="#dc2626")
            return
        if not sn.startswith("V01") or len(sn) < 8:
            messagebox.showwarning("Invalid SN", "Please enter a valid Module Serial Number starting with V01.")
            return

        # Check in-memory records first for instant display
        for r in self.records:
            if r.get('sn') == sn and (r.get('tumsoldering') or r.get('tumlayup') or r.get('tumlamination')):
                lt_info = f" ({r.get('layup_time')})" if r.get('layup_time') else ""
                sol = r.get('tumsoldering', '')
                lay = r.get('tumlayup', '')
                lam = r.get('tumlamination', '')
                mach_str = " | ".join([m for m in [sol, lay, lam] if m])
                self.lbl_query_status.config(text=f"✅ Found (cached): {mach_str}{lt_info}", fg="#059669")
                return

        self.lbl_query_status.config(text=f"Logging in (030888) & Querying MES for {sn}...", fg="#2563eb")

        def _worker():
            try:
                res = query_mes_process_log(sn)
                def _ui():
                    has_mach = bool(res.get('raw_found') or res.get('tumsoldering') or res.get('tumlayup') or res.get('tumlamination'))
                    if has_mach:
                        lt_info = f" ({res.get('layup_time')})" if res.get('layup_time') else ""
                        sol = res.get('tumsoldering', '')
                        lay = res.get('tumlayup', '')
                        lam = res.get('tumlamination', '')
                        mach_str = " | ".join([m for m in [sol, lay, lam] if m])
                        self.lbl_query_status.config(text=f"✅ Found: {mach_str}{lt_info}", fg="#059669")
                        self.load_records()
                    elif res.get('status') == 'offline_pc':
                        self.lbl_query_status.config(text="⚠️ Host PC on office Wi-Fi (10.200.3.109 unreachable). Connect to plant LAN.", fg="#d97706")
                    else:
                        self.lbl_query_status.config(text=f"Complete. No record found for {sn}.", fg="#64748b")
                self.after(0, _ui)
            except Exception as err:
                self.after(0, lambda: self.lbl_query_status.config(text=f"Query error: {err}", fg="#dc2626"))

        threading.Thread(target=_worker, daemon=True).start()

    def export_csv(self):
        csv_content = export_mes_trend_csv_string()
        def_fname = f"MES_Process_Log_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        fpath = filedialog.asksaveasfilename(
            defaultextension=".csv",
            filetypes=[("CSV Files", "*.csv"), ("All Files", "*.*")],
            initialfile=def_fname
        )
        if fpath:
            try:
                with open(fpath, 'w', encoding='utf-8') as f:
                    f.write(csv_content)
                messagebox.showinfo("Export Success", f"MES Process Log exported successfully to:\n{fpath}")
            except Exception as e:
                messagebox.showerror("Export Failed", f"Could not write CSV file:\n{e}")

    def clear_log(self):
        if not self.records:
            messagebox.showinfo("Log Empty", "MES Process Log is already empty.")
            return
        if messagebox.askyesno("Clear MES Log", "Are you sure you want to clear all MES Process Log entries?\nThis will permanently reset the trend log file."):
            try:
                with open(MES_TREND_FILE, 'w', encoding='utf-8') as f:
                    json.dump([], f, ensure_ascii=False, indent=2)
            except Exception as e:
                print(f"Clear log error: {e}")
            self.load_records()
            self.lbl_query_status.config(text="Log cleared", fg="#64748b")

    def open_mes_portal(self):
        base_tab_url = "http://10.200.3.109:8080/webroot/decision/v10/entry/access/416090fb-b706-40e8-9e4d-d698a059f6bf?preview=true"
        sn = self.ent_query_sn.get().strip()
        url = base_tab_url
        if sn:
            enc_sn = urllib.parse.quote(sn)
            url = f"{base_tab_url}&MOUDLEID={enc_sn}&__bypassevent__=true&组件序列号={enc_sn}&ModuleSerialNo={enc_sn}&SN={enc_sn}"
            try:
                self.clipboard_clear()
                self.clipboard_append(sn)
                self.lbl_query_status.config(text=f"📋 Copied {sn} to clipboard. Opening MES Tab...", fg="#059669")
            except Exception:
                pass
        try:
            webbrowser.open(url)
        except Exception as e:
            messagebox.showerror("Cannot Open Browser", f"Error opening URL:\n{e}")

    def open_mes_login(self):
        url = "http://10.200.3.109:8080/webroot/decision/login"
        try:
            self.clipboard_clear()
            self.clipboard_append("030888")
            self.lbl_query_status.config(text="🔐 Copied 030888 to clipboard. Opening Login...", fg="#059669")
        except Exception:
            pass
        try:
            webbrowser.open(url)
        except Exception as e:
            messagebox.showerror("Cannot Open Browser", f"Error opening URL:\n{e}")

    def on_row_double_clicked(self, event):
        sel = self.tree.selection()
        if not sel: return
        item = self.tree.item(sel[0])
        vals = item.get('values', [])
        if len(vals) > 1:
            sn = str(vals[1])
            self.clipboard_clear()
            self.clipboard_append(sn)
            self.lbl_query_status.config(text=f"📋 Copied SN: {sn}", fg="#059669")

    def show_context_menu(self, event):
        item_id = self.tree.identify_row(event.y)
        if item_id:
            self.tree.selection_set(item_id)
            self.context_menu.post(event.x_root, event.y_root)

    def copy_selected_sn(self):
        sel = self.tree.selection()
        if not sel: return
        vals = self.tree.item(sel[0]).get('values', [])
        if len(vals) > 1:
            sn = str(vals[1])
            self.clipboard_clear()
            self.clipboard_append(sn)
            self.lbl_query_status.config(text=f"📋 Copied SN: {sn}", fg="#059669")

    def copy_selected_row(self):
        sel = self.tree.selection()
        if not sel: return
        vals = self.tree.item(sel[0]).get('values', [])
        txt = "\t".join(str(v) for v in vals)
        self.clipboard_clear()
        self.clipboard_append(txt)
        self.lbl_query_status.config(text="📋 Copied row details", fg="#059669")

    def requery_selected_sn(self):
        sel = self.tree.selection()
        if not sel: return
        vals = self.tree.item(sel[0]).get('values', [])
        if len(vals) > 1:
            sn = str(vals[1])
            self.ent_query_sn.delete(0, tk.END)
            self.ent_query_sn.insert(0, sn)
            self.query_sn()

    def delete_selected_record(self):
        sel = self.tree.selection()
        if not sel: return
        vals = self.tree.item(sel[0]).get('values', [])
        if len(vals) > 1:
            sn = str(vals[1])
            if messagebox.askyesno("Delete Record", f"Remove record for SN '{sn}' from MES log?"):
                self.records = [r for r in self.records if r.get('sn') != sn]
                try:
                    with open(MES_TREND_FILE, 'w', encoding='utf-8') as f:
                        json.dump(self.records, f, ensure_ascii=False, indent=2)
                except Exception: pass
                self.apply_filter()
                self.update_kpi_cards()
                self.lbl_query_status.config(text=f"Deleted {sn}", fg="#64748b")

    def open_mes_photos_folder(self):
        mes_dir = getattr(config, 'MES_PHOTOS_DIR', os.path.join(config.LOCAL_DATA_DIR, 'mes_photos'))
        os.makedirs(mes_dir, exist_ok=True)
        try:
            if os.name == 'nt':
                os.startfile(mes_dir)
            else:
                subprocess.Popen(['xdg-open', mes_dir])
        except Exception as e:
            messagebox.showerror("Error", f"Could not open MES photos folder:\n{e}")

    def view_selected_mes_photo(self):
        sel = self.tree.selection()
        if not sel:
            messagebox.showinfo("View Photo", "Please select a record from the table first.")
            return
        vals = self.tree.item(sel[0]).get('values', [])
        if len(vals) < 2: return
        sn = str(vals[1]).strip()

        found_photo = ""
        # 1. Check in-memory records for saved photo_path
        for rec in self.records:
            if rec.get('sn') == sn and rec.get('photo_path') and os.path.exists(rec.get('photo_path')):
                found_photo = rec.get('photo_path')
                break

        # 2. Check MES_PHOTOS_DIR and DailyCache/mes_photos
        if not found_photo:
            mes_dir = getattr(config, 'MES_PHOTOS_DIR', os.path.join(config.LOCAL_DATA_DIR, 'mes_photos'))
            daily_mes_dir = os.path.join(getattr(config, 'DAILY_CACHE_DIR', ''), 'mes_photos')
            for d in [mes_dir, daily_mes_dir]:
                if not os.path.exists(d): continue
                for fname in os.listdir(d):
                    if sn in fname:
                        found_photo = os.path.join(d, fname)
                        break
                if found_photo: break

        if found_photo and os.path.exists(found_photo):
            try:
                if os.name == 'nt':
                    os.startfile(found_photo)
                else:
                    subprocess.Popen(['xdg-open', found_photo])
            except Exception as e:
                messagebox.showerror("Error", f"Could not open photo file:\n{e}")
        else:
            messagebox.showinfo(
                "MES Photo",
                f"No saved MES photo found for Serial Number:\n{sn}\n\n"
                "Photos are captured from Mobile Hub Tab 3 using '1. MES Barcode' or '2. Get SN Pic' and saved into the dedicated 'mes_photos' folder."
            )

    def open_chinese_guide(self):
        MESChineseFieldGuideDialog(self)



# =================== QR CODE GENERATOR & VIEWER ===================

def generate_qr_image(url: str, size: int = 100):
    try:
        import qrcode, re
        m = re.search(r'(https?://[^\s\n\r]+)', str(url))
        clean_url = m.group(1) if m else str(url).strip()
        if not clean_url.startswith("http"):
            clean_url = f"http://{clean_url}"
        qr = qrcode.QRCode(
            version=1,
            error_correction=qrcode.constants.ERROR_CORRECT_M,
            box_size=4,
            border=2,
        )
        qr.add_data(clean_url)
        qr.make(fit=True)
        img = qr.make_image(fill_color="black", back_color="white")
        return img.resize((size, size), Image.Resampling.NEAREST)
    except Exception as e:
        print(f"QR Code generation error: {e}")
        return None


class QRCodeDialog(tk.Toplevel):
    def __init__(self, parent, url):
        super().__init__(parent)
        self.title("Scan to Open Mobile HUD")
        self.geometry("380x440")
        self.resizable(False, False)
        self.configure(bg="#0f172a")
        
        import re
        m = re.search(r'(https?://[^\s\n\r]+)', str(url))
        clean_url = m.group(1) if m else str(url).strip()
        if not clean_url.startswith("http"):
            clean_url = f"http://{clean_url}"
            
        tk.Label(self, text="📱 Mobile HUD Link", font=("Segoe UI", 14, "bold"), bg="#0f172a", fg="#ffffff").pack(pady=(20, 5))
        tk.Label(self, text="Scan with your phone camera to open the live dashboard:", font=("Segoe UI", 9), bg="#0f172a", fg="#94a3b8", wraplength=320).pack(pady=(0, 12))
        
        qr_img = generate_qr_image(clean_url, size=210)
        if qr_img:
            self.tk_qr = ImageTk.PhotoImage(qr_img)
            lbl_qr = tk.Label(self, image=self.tk_qr, bg="white", padx=8, pady=8, relief=tk.SOLID, bd=2)
            lbl_qr.pack(pady=5)
            
        tk.Label(self, text=clean_url, font=("Consolas", 10, "bold"), bg="#1e293b", fg="#38bdf8", padx=12, pady=6).pack(pady=10)
        ModernButton(self, text="Close", command=self.destroy, primary=False, padx=20, pady=4).pack(pady=5)



# =================== DEFECT SNIPPING & INSPECTION TOOL ===================

class ImageSnipperDialog(tk.Toplevel):
    def __init__(self, parent, title, image_path, record=None, on_snip_saved=None, image_category="MR_DEFECT"):
        super().__init__(parent)
        sn_label = f"[{record.get('sn')}]" if record and record.get('sn') and record.get('sn') != 'Pending SN' else ""
        self.title(f"🔍 QC Image Inspection & Snipping Tool - {title} {sn_label}")
        self.geometry("1100x820")
        self.minsize(850, 640)
        self.configure(bg="#0f172a")
        
        self.image_path = image_path
        self.record = record
        self.on_snip_saved = on_snip_saved
        self.image_category = image_category  # "PRE_EL_FRONT", "PRE_EL_EL", "PRE_EL_BACK", "MR_DEFECT"
        
        self.orig_pil_img = None
        self.tk_img = None
        self.orig_w = 0
        self.orig_h = 0
        self.base_fit_scale = 1.0
        self.scale_factor = 1.0
        self.zoom_level = 1.0  # 1.0 = fit to canvas
        self.pan_offset_x = 0
        self.pan_offset_y = 0
        self.img_offset_x = 0
        self.img_offset_y = 0
        self.disp_w = 0
        self.disp_h = 0
        
        # Pan state
        self.pan_start_x = 0
        self.pan_start_y = 0
        self.is_panning = False
        
        # ROI coordinates in original image space
        self.start_x = None
        self.start_y = None
        self.rect_id = None
        self.roi_coords = None  # (orig_x1, orig_y1, orig_x2, orig_y2)
        
        self.bind("<Return>", lambda e: self.save_defect_snip())
        self.bind("<KP_Enter>", lambda e: self.save_defect_snip())
        self.bind("<Escape>", lambda e: self.destroy())
        self.focus_force()

        self._build_ui(title)
        self._load_source_image()

    def _build_ui(self, title):
        # 1. Top Header Bar
        top_bar = tk.Frame(self, bg="#1e293b", padx=14, pady=8)
        top_bar.pack(fill=tk.X)
        
        title_box = tk.Frame(top_bar, bg="#1e293b")
        title_box.pack(side=tk.LEFT, fill=tk.Y)
        
        sn_text = f" • SN: {self.record.get('sn', 'N/A')}" if self.record else ""
        tk.Label(title_box, text=f"🔍 {title}{sn_text}", font=("Segoe UI", 11, "bold"), bg="#1e293b", fg="#f8fafc").pack(anchor="w")
        tk.Label(title_box, text=f"File: {os.path.basename(self.image_path)}", font=("Consolas", 8), bg="#1e293b", fg="#94a3b8").pack(anchor="w")
        
        # Helper instruction badge
        badge = tk.Frame(top_bar, bg="#0f172a", padx=10, pady=4, highlightbackground="#334155", highlightthickness=1)
        badge.pack(side=tk.RIGHT, padx=(10, 0))
        tk.Label(badge, text="🖱️ Wheel: Zoom | Right-Drag: Pan | Left-Drag: Crop ROI | [Enter]: Save Snip", font=("Segoe UI", 9, "bold"), bg="#0f172a", fg="#38bdf8").pack()

        # 2. Main Interactive Canvas Area
        canvas_frame = tk.Frame(self, bg="#090d16")
        canvas_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=(6, 4))
        
        self.canvas = tk.Canvas(canvas_frame, bg="#090d16", highlightthickness=0, cursor="crosshair")
        self.canvas.pack(fill=tk.BOTH, expand=True)
        
        self.canvas.bind("<Configure>", self._on_canvas_resize)
        self.canvas.bind("<ButtonPress-1>", self._on_mouse_down)
        self.canvas.bind("<B1-Motion>", self._on_mouse_drag)
        self.canvas.bind("<ButtonRelease-1>", self._on_mouse_up)
        
        # Mouse Scroll Wheel Zoom (Windows and Linux support)
        self.canvas.bind("<MouseWheel>", self._on_mouse_wheel)
        self.canvas.bind("<Button-4>", lambda e: self._on_mouse_wheel_step(1.22, e.x, e.y))
        self.canvas.bind("<Button-5>", lambda e: self._on_mouse_wheel_step(1.0 / 1.22, e.x, e.y))
        
        # Right-Click & Middle-Click Pan / Drag
        self.canvas.bind("<ButtonPress-3>", self._on_pan_down)
        self.canvas.bind("<B3-Motion>", self._on_pan_drag)
        self.canvas.bind("<ButtonRelease-3>", self._on_pan_up)
        self.canvas.bind("<ButtonPress-2>", self._on_pan_down)
        self.canvas.bind("<B2-Motion>", self._on_pan_drag)
        self.canvas.bind("<ButtonRelease-2>", self._on_pan_up)
        
        # Double-click to reset zoom to Fit Window
        self.canvas.bind("<Double-Button-1>", lambda e: self.reset_zoom())
        
        # 3. Bottom Control & Action Bar
        bottom_bar = tk.Frame(self, bg="#1e293b", padx=14, pady=8)
        bottom_bar.pack(fill=tk.X, padx=10, pady=(0, 8))
        
        status_box = tk.Frame(bottom_bar, bg="#1e293b")
        status_box.pack(side=tk.LEFT, fill=tk.Y)
        
        self.lbl_status = tk.Label(status_box, text="Ready. Scroll wheel to zoom, right-click to pan, drag box to crop.", font=("Segoe UI", 9), bg="#1e293b", fg="#94a3b8")
        self.lbl_status.pack(anchor="w")
        
        btn_box = tk.Frame(bottom_bar, bg="#1e293b")
        btn_box.pack(side=tk.RIGHT)
        
        # Zoom Controls
        zoom_box = tk.Frame(btn_box, bg="#0f172a", padx=4, pady=2, highlightbackground="#334155", highlightthickness=1)
        zoom_box.pack(side=tk.LEFT, padx=(0, 8))
        
        btn_z_out = tk.Button(zoom_box, text="－", font=("Segoe UI", 9, "bold"), bg="#1e293b", fg="#ffffff", relief=tk.FLAT, padx=6, pady=2, cursor="hand2", command=self.zoom_out_button)
        btn_z_out.pack(side=tk.LEFT, padx=1)
        
        self.lbl_zoom_pct = tk.Label(zoom_box, text="🔍 100%", font=("Segoe UI", 9, "bold"), bg="#0f172a", fg="#38bdf8", width=8)
        self.lbl_zoom_pct.pack(side=tk.LEFT, padx=2)
        
        btn_z_in = tk.Button(zoom_box, text="＋", font=("Segoe UI", 9, "bold"), bg="#1e293b", fg="#ffffff", relief=tk.FLAT, padx=6, pady=2, cursor="hand2", command=self.zoom_in_button)
        btn_z_in.pack(side=tk.LEFT, padx=1)
        
        btn_fit = tk.Button(zoom_box, text="↺ Fit", font=("Segoe UI", 8, "bold"), bg="#334155", fg="#e2e8f0", relief=tk.FLAT, padx=6, pady=2, cursor="hand2", command=self.reset_zoom)
        btn_fit.pack(side=tk.LEFT, padx=2)

        btn_100 = tk.Button(zoom_box, text="1:1", font=("Segoe UI", 8, "bold"), bg="#334155", fg="#e2e8f0", relief=tk.FLAT, padx=6, pady=2, cursor="hand2", command=self.set_zoom_100)
        btn_100.pack(side=tk.LEFT, padx=1)
        
        self.btn_save_snip = tk.Button(
            btn_box, text="✂️ Save Defect Snip [Enter]", font=("Segoe UI", 9, "bold"),
            bg="#059669", fg="#ffffff", activebackground="#047857", activeforeground="#ffffff",
            relief=tk.FLAT, padx=12, pady=5, cursor="hand2", command=self.save_defect_snip
        )
        self.btn_save_snip.pack(side=tk.LEFT, padx=4)

        btn_rot_left = tk.Button(
            btn_box, text="⟲ 90°", font=("Segoe UI", 9, "bold"),
            bg="#334155", fg="#38bdf8", activebackground="#475569", activeforeground="#ffffff",
            relief=tk.FLAT, padx=8, pady=5, cursor="hand2", command=lambda: self.rotate_image(90)
        )
        btn_rot_left.pack(side=tk.LEFT, padx=3)
        
        btn_rot_right = tk.Button(
            btn_box, text="⟳ 90°", font=("Segoe UI", 9, "bold"),
            bg="#334155", fg="#38bdf8", activebackground="#475569", activeforeground="#ffffff",
            relief=tk.FLAT, padx=8, pady=5, cursor="hand2", command=lambda: self.rotate_image(270)
        )
        btn_rot_right.pack(side=tk.LEFT, padx=3)
        
        btn_reset = tk.Button(
            btn_box, text="🔄 Reset ROI", font=("Segoe UI", 9),
            bg="#334155", fg="#e2e8f0", activebackground="#475569", activeforeground="#ffffff",
            relief=tk.FLAT, padx=10, pady=5, cursor="hand2", command=self.reset_selection
        )
        btn_reset.pack(side=tk.LEFT, padx=4)
        
        btn_open_full = tk.Button(
            btn_box, text="🔍 Windows Viewer", font=("Segoe UI", 9),
            bg="#2563eb", fg="#ffffff", activebackground="#1d4ed8", activeforeground="#ffffff",
            relief=tk.FLAT, padx=10, pady=5, cursor="hand2", command=self.open_in_windows_photos
        )
        btn_open_full.pack(side=tk.LEFT, padx=4)
        
        btn_close = tk.Button(
            btn_box, text="Close", font=("Segoe UI", 9),
            bg="#475569", fg="#e2e8f0", activebackground="#64748b", activeforeground="#ffffff",
            relief=tk.FLAT, padx=10, pady=5, cursor="hand2", command=self.destroy
        )
        btn_close.pack(side=tk.LEFT, padx=4)

    def _load_source_image(self):
        try:
            if not os.path.exists(self.image_path):
                self.lbl_status.config(text="Image file not found!", fg="#ef4444")
                return
            from PIL import ImageOps
            raw_img = Image.open(self.image_path)
            self.orig_pil_img = ImageOps.exif_transpose(raw_img) or raw_img
            self.orig_w, self.orig_h = self.orig_pil_img.size
            self.lbl_status.config(text=f"Loaded Image: {self.orig_w} × {self.orig_h} px. Scroll wheel to zoom in/out.")
            self.after(60, self._render_canvas_image)
        except Exception as e:
            self.lbl_status.config(text=f"Error loading image: {e}", fg="#ef4444")

    def rotate_image(self, degrees):
        if not self.orig_pil_img: return
        try:
            self.orig_pil_img = self.orig_pil_img.rotate(degrees, expand=True)
            self.orig_w, self.orig_h = self.orig_pil_img.size
            self.reset_selection()
            self.reset_zoom()
            self.lbl_status.config(text=f"🔄 Rotated Image ({self.orig_w} × {self.orig_h} px)", fg="#38bdf8")
            try:
                self.orig_pil_img.save(self.image_path, quality=95)
            except Exception: pass
        except Exception as e:
            self.lbl_status.config(text=f"Rotate error: {e}", fg="#ef4444")

    def _on_canvas_resize(self, event):
        self._render_canvas_image()

    def _on_mouse_wheel(self, event):
        factor = 1.25 if event.delta > 0 else (1.0 / 1.25)
        self._zoom_at_point(factor, event.x, event.y)
        return "break"

    def _on_mouse_wheel_step(self, factor, mouse_x, mouse_y):
        self._zoom_at_point(factor, mouse_x, mouse_y)
        return "break"

    def _zoom_at_point(self, factor, mouse_x, mouse_y):
        if not self.orig_pil_img or self.orig_w <= 0 or self.orig_h <= 0: return

        new_zoom = max(0.3, min(35.0, self.zoom_level * factor))
        if abs(new_zoom - self.zoom_level) < 0.001: return

        cw = max(self.canvas.winfo_width(), 100)
        ch = max(self.canvas.winfo_height(), 100)
        
        avail_w = max(cw - 16, 10)
        avail_h = max(ch - 16, 10)
        self.base_fit_scale = min(avail_w / max(self.orig_w, 1), avail_h / max(self.orig_h, 1))

        old_scale = self.base_fit_scale * self.zoom_level
        new_scale = self.base_fit_scale * new_zoom

        old_img_x = (cw - int(self.orig_w * old_scale)) // 2 + self.pan_offset_x
        old_img_y = (ch - int(self.orig_h * old_scale)) // 2 + self.pan_offset_y

        px = (mouse_x - old_img_x) / max(old_scale, 0.00001)
        py = (mouse_y - old_img_y) / max(old_scale, 0.00001)

        self.zoom_level = new_zoom

        new_img_x = mouse_x - (px * new_scale)
        new_img_y = mouse_y - (py * new_scale)

        base_center_x = (cw - int(self.orig_w * new_scale)) // 2
        base_center_y = (ch - int(self.orig_h * new_scale)) // 2

        self.pan_offset_x = int(new_img_x - base_center_x)
        self.pan_offset_y = int(new_img_y - base_center_y)

        self._render_canvas_image()
        self._update_zoom_label()

    def _on_pan_down(self, event):
        self.pan_start_x = event.x
        self.pan_start_y = event.y
        self.is_panning = True
        self.canvas.config(cursor="fleur")

    def _on_pan_drag(self, event):
        if not self.is_panning: return
        dx = event.x - self.pan_start_x
        dy = event.y - self.pan_start_y
        self.pan_offset_x += dx
        self.pan_offset_y += dy
        self.pan_start_x = event.x
        self.pan_start_y = event.y
        self._render_canvas_image()

    def _on_pan_up(self, event):
        self.is_panning = False
        self.canvas.config(cursor="crosshair")

    def reset_zoom(self):
        self.zoom_level = 1.0
        self.pan_offset_x = 0
        self.pan_offset_y = 0
        self._render_canvas_image()
        self._update_zoom_label()
        self.lbl_status.config(text=f"Reset to Fit Window ({self.orig_w} × {self.orig_h} px)", fg="#94a3b8")

    def set_zoom_100(self):
        if not self.orig_pil_img or self.base_fit_scale <= 0: return
        self.zoom_level = 1.0 / self.base_fit_scale
        self.pan_offset_x = 0
        self.pan_offset_y = 0
        self._render_canvas_image()
        self._update_zoom_label()
        self.lbl_status.config(text=f"100% Native Pixel Scale ({self.orig_w} × {self.orig_h} px)", fg="#38bdf8")

    def zoom_in_button(self):
        cw = max(self.canvas.winfo_width(), 100)
        ch = max(self.canvas.winfo_height(), 100)
        self._zoom_at_point(1.3, cw // 2, ch // 2)

    def zoom_out_button(self):
        cw = max(self.canvas.winfo_width(), 100)
        ch = max(self.canvas.winfo_height(), 100)
        self._zoom_at_point(1.0 / 1.3, cw // 2, ch // 2)

    def _update_zoom_label(self):
        if hasattr(self, 'lbl_zoom_pct'):
            pct = int(round(self.zoom_level * 100))
            self.lbl_zoom_pct.config(text=f"🔍 {pct}%")

    def _render_canvas_image(self):
        if not self.orig_pil_img: return
        cw = max(self.canvas.winfo_width(), 100)
        ch = max(self.canvas.winfo_height(), 100)
        
        avail_w = max(cw - 16, 10)
        avail_h = max(ch - 16, 10)
        
        self.base_fit_scale = min(avail_w / max(self.orig_w, 1), avail_h / max(self.orig_h, 1))
        self.scale_factor = self.base_fit_scale * self.zoom_level
        self.disp_w = max(1, int(self.orig_w * self.scale_factor))
        self.disp_h = max(1, int(self.orig_h * self.scale_factor))
        
        self.img_offset_x = (cw - self.disp_w) // 2 + self.pan_offset_x
        self.img_offset_y = (ch - self.disp_h) // 2 + self.pan_offset_y

        # Fast viewport-aware cropping for instant, ultra-smooth zooming & panning
        vis_x1 = max(0, int((0 - self.img_offset_x) / max(self.scale_factor, 0.00001)))
        vis_y1 = max(0, int((0 - self.img_offset_y) / max(self.scale_factor, 0.00001)))
        vis_x2 = min(self.orig_w, int((cw - self.img_offset_x) / max(self.scale_factor, 0.00001)) + 1)
        vis_y2 = min(self.orig_h, int((ch - self.img_offset_y) / max(self.scale_factor, 0.00001)) + 1)

        self.canvas.delete("all")

        if vis_x2 > vis_x1 and vis_y2 > vis_y1:
            try:
                crop_region = self.orig_pil_img.crop((vis_x1, vis_y1, vis_x2, vis_y2))
                vis_w = max(1, int((vis_x2 - vis_x1) * self.scale_factor))
                vis_h = max(1, int((vis_y2 - vis_y1) * self.scale_factor))
                
                render_img = crop_region.resize((vis_w, vis_h), Image.Resampling.BILINEAR)
                self.tk_img = ImageTk.PhotoImage(render_img)

                draw_x = self.img_offset_x + int(vis_x1 * self.scale_factor)
                draw_y = self.img_offset_y + int(vis_y1 * self.scale_factor)
                self.canvas.create_image(draw_x, draw_y, image=self.tk_img, anchor="nw")
            except Exception as ren_err:
                pass
        
        # Redraw ROI selection rectangle if exists
        if self.roi_coords:
            ox1, oy1, ox2, oy2 = self.roi_coords
            cx1 = self.img_offset_x + int(ox1 * self.scale_factor)
            cy1 = self.img_offset_y + int(oy1 * self.scale_factor)
            cx2 = self.img_offset_x + int(ox2 * self.scale_factor)
            cy2 = self.img_offset_y + int(oy2 * self.scale_factor)
            self.rect_id = self.canvas.create_rectangle(cx1, cy1, cx2, cy2, outline="#00ffcc", width=2, dash=(4, 2))

    def _on_mouse_down(self, event):
        self.start_x = event.x
        self.start_y = event.y
        if self.rect_id:
            self.canvas.delete(self.rect_id)
            self.rect_id = None
        self.rect_id = self.canvas.create_rectangle(self.start_x, self.start_y, self.start_x, self.start_y, outline="#00ffcc", width=2, dash=(4, 2))

    def _on_mouse_drag(self, event):
        if self.start_x is None or self.start_y is None or not self.rect_id: return
        cur_x = event.x
        cur_y = event.y
        self.canvas.coords(self.rect_id, self.start_x, self.start_y, cur_x, cur_y)
        
        rx1, rx2 = sorted([self.start_x, cur_x])
        ry1, ry2 = sorted([self.start_y, cur_y])
        
        ox1 = max(0, min(self.orig_w, int((rx1 - self.img_offset_x) / max(self.scale_factor, 0.00001))))
        oy1 = max(0, min(self.orig_h, int((ry1 - self.img_offset_y) / max(self.scale_factor, 0.00001))))
        ox2 = max(0, min(self.orig_w, int((rx2 - self.img_offset_x) / max(self.scale_factor, 0.00001))))
        oy2 = max(0, min(self.orig_h, int((ry2 - self.img_offset_y) / max(self.scale_factor, 0.00001))))
        
        w_px = max(1, ox2 - ox1)
        h_px = max(1, oy2 - oy1)
        pct = int(round(self.zoom_level * 100))
        self.lbl_status.config(text=f"Selected ROI: {w_px} × {h_px} px | Zoom: {pct}% (Original: {self.orig_w} × {self.orig_h} px)", fg="#38bdf8")

    def _on_mouse_up(self, event):
        if self.start_x is None or self.start_y is None: return
        end_x = event.x
        end_y = event.y
        
        rx1, rx2 = sorted([self.start_x, end_x])
        ry1, ry2 = sorted([self.start_y, end_y])
        
        orig_x1 = max(0, min(self.orig_w, int((rx1 - self.img_offset_x) / max(self.scale_factor, 0.00001))))
        orig_y1 = max(0, min(self.orig_h, int((ry1 - self.img_offset_y) / max(self.scale_factor, 0.00001))))
        orig_x2 = max(0, min(self.orig_w, int((rx2 - self.img_offset_x) / max(self.scale_factor, 0.00001))))
        orig_y2 = max(0, min(self.orig_h, int((ry2 - self.img_offset_y) / max(self.scale_factor, 0.00001))))
        
        if (orig_x2 - orig_x1) >= 15 and (orig_y2 - orig_y1) >= 15:
            self.roi_coords = (orig_x1, orig_y1, orig_x2, orig_y2)
            self.lbl_status.config(text=f"✅ ROI Selected: {orig_x2 - orig_x1} × {orig_y2 - orig_y1} px. Press [Enter] or 'Save Defect Snip' to attach!", fg="#4ade80")
            cx1 = self.img_offset_x + int(orig_x1 * self.scale_factor)
            cy1 = self.img_offset_y + int(orig_y1 * self.scale_factor)
            cx2 = self.img_offset_x + int(orig_x2 * self.scale_factor)
            cy2 = self.img_offset_y + int(orig_y2 * self.scale_factor)
            self.canvas.coords(self.rect_id, cx1, cy1, cx2, cy2)
        else:
            self.reset_selection()
            self.lbl_status.config(text="Selection too small. Drag a larger box over the defect.", fg="#f59e0b")

    def reset_selection(self):
        if self.rect_id:
            self.canvas.delete(self.rect_id)
            self.rect_id = None
        self.roi_coords = None
        self.lbl_status.config(text="Selection cleared. Drag a box over defect area.", fg="#94a3b8")

    def save_defect_snip(self):
        if not self.orig_pil_img:
            messagebox.showwarning("Snipping Tool", "No image loaded to crop!")
            return
            
        crop_box = self.roi_coords
        if not crop_box:
            crop_box = (0, 0, self.orig_w, self.orig_h)
            
        try:
            cropped = self.orig_pil_img.crop(crop_box)
            if cropped.mode in ('RGBA', 'P'):
                cropped = cropped.convert('RGB')
            
            # 1. Downscale if crop dimensions exceed 950px on longest edge
            cw, ch = cropped.size
            max_dim = 950
            if max(cw, ch) > max_dim:
                scale = max_dim / max(cw, ch)
                cropped = cropped.resize((max(1, int(cw * scale)), max(1, int(ch * scale))), Image.Resampling.LANCZOS)
            
            cache_dir = getattr(config, 'DAILY_CACHE_DIR', 'DailyCache')
            os.makedirs(cache_dir, exist_ok=True)
            
            cat = str(getattr(self, 'image_category', 'MR_DEFECT')).upper()
            if "FRONT" in cat:
                cat_tag = "PreEL_Front"
                v_key = "front"
            elif "BACK" in cat:
                cat_tag = "PreEL_Back"
                v_key = "back"
            elif "EL" in cat:
                cat_tag = "PreEL_EL"
                v_key = "el"
            elif "PRE" in cat:
                cat_tag = "PreEL"
                v_key = "front"
            else:
                cat_tag = "Def"
                v_key = "defect"

            sn = self.record.get('sn', 'Defect') if self.record else 'Defect'
            if sn in ("Pending SN", "-", "", None):
                sn = "Defect"
            timestamp = int(time.time())
            snip_filename = f"Snip_{cat_tag}_{sn}_{timestamp}.jpg"
            snip_path = os.path.join(cache_dir, snip_filename)
            
            # 2. Fast single-pass save (< 3ms) targeting ~150-200 KB max for fast Excel loading
            cropped.save(snip_path, format='JPEG', quality=82)
            
            if self.record is not None:
                if "PRE" in cat_tag.upper():
                    # Pre-Layup Photo (Column J in Excel)
                    self.record['pre_el_snip_path'] = snip_path
                    self.record['pre_el_snip_view'] = v_key
                    if 'pre_el_views' in self.record and isinstance(self.record['pre_el_views'], dict):
                        if v_key in self.record['pre_el_views'] and self.record['pre_el_views'][v_key]:
                            self.record['pre_el_views'][v_key]['snip_path'] = snip_path
                else:
                    # Post-Layup Photo (Column K in Excel - MR Defect Pic)
                    self.record['photo_path'] = snip_path
                    self.record['snip_saved'] = True
                
            if self.on_snip_saved:
                self.on_snip_saved(self.record, snip_path, self.image_category)
                
            self.destroy()
        except Exception as e:
            messagebox.showerror("Error Saving Snip", f"Failed to save cropped image:\n{e}")

    def open_in_windows_photos(self):
        try:
            os.startfile(self.image_path)
        except Exception as e:
            messagebox.showerror("Error", f"Could not open in Windows Photos: {e}")

ImageViewerDialog = ImageSnipperDialog


# =================== MODULE REVIEW TAB ===================

class ModuleReviewTab(tk.Frame):
    def __init__(self, parent, app):
        super().__init__(parent, bg=config.COLOR_BG_MAIN, padx=15, pady=10)
        self.app = app
        self.records = []
        self.selected_rec_id = None
        self.active_photo_path = None
        self.last_spoken_raw = ""
        self.last_audio_path = None
        self.processed_photo_aliases = {}  # Maps photo path or basename to record ID
        self.setup_ui()

    def register_photo_alias(self, rec_id, *paths):
        for p in paths:
            if p:
                self.processed_photo_aliases[str(p)] = rec_id
                self.processed_photo_aliases[os.path.basename(str(p))] = rec_id

    def setup_ui(self):
        top_ctrl = tk.Frame(self, bg=config.COLOR_BG_MAIN)
        top_ctrl.pack(fill=tk.X, pady=(0, 10))

        scan_box = tk.LabelFrame(top_ctrl, text=" 1. Barcode / QR Scanner Entry ", font=config.FONT_BODY_BOLD, bg=config.COLOR_BG_MAIN, fg=config.COLOR_PRIMARY, padx=10, pady=8)
        scan_box.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 8))

        tk.Label(scan_box, text="Serial Number:", font=config.FONT_BODY, bg=config.COLOR_BG_MAIN).pack(anchor=tk.W)
        self.ent_sn = tk.Entry(scan_box, font=("Segoe UI", 12, "bold"), bg=config.COLOR_INPUT_BG, relief=tk.FLAT, bd=5)
        self.ent_sn.pack(fill=tk.X, pady=(4, 6))
        self.ent_sn.bind("<Return>", lambda e: self.on_barcode_scanned())
        self.ent_sn.focus_set()

        meta_row = tk.Frame(scan_box, bg=config.COLOR_BG_MAIN)
        meta_row.pack(fill=tk.X, pady=2)

        tk.Label(meta_row, text="Classification:", font=config.FONT_SMALL, bg=config.COLOR_BG_MAIN).grid(row=0, column=0, sticky="w")
        class_options = [""] + list(config.DEFECT_TREE.keys())
        self.cb_class = ttk.Combobox(meta_row, values=class_options, width=24, state="readonly")
        self.cb_class.set("")
        self.cb_class.grid(row=0, column=1, padx=4, pady=2)
        self.cb_class.bind("<<ComboboxSelected>>", lambda e: self.on_class_changed())

        tk.Label(meta_row, text="Defect Summary:", font=config.FONT_SMALL, bg=config.COLOR_BG_MAIN).grid(row=0, column=2, sticky="w", padx=(6, 0))
        self.cb_summary = ttk.Combobox(meta_row, values=[""], width=24)
        self.cb_summary.set("")
        self.cb_summary.grid(row=0, column=3, padx=4, pady=2)

        tk.Label(meta_row, text="Result:", font=config.FONT_SMALL, bg=config.COLOR_BG_MAIN).grid(row=1, column=0, sticky="w")
        self.cb_result = ttk.Combobox(meta_row, values=["", "Q3", "Scrap"], width=10, state="readonly")
        self.cb_result.set("")
        self.cb_result.grid(row=1, column=1, sticky="w", padx=4, pady=2)

        tk.Label(meta_row, text="Line Input:", font=config.FONT_SMALL, bg=config.COLOR_BG_MAIN).grid(row=1, column=2, sticky="w", padx=(6, 0))
        self.cb_line = ttk.Combobox(meta_row, values=["", "Line 1", "Line 2", "Line 3", "Line 4", "Line 5", "Line 6", "Line 7"], width=12, state="readonly")
        self.cb_line.set("")
        self.cb_line.grid(row=1, column=3, sticky="w", padx=4, pady=2)

        tk.Label(meta_row, text="Action Cause:", font=config.FONT_SMALL, bg=config.COLOR_BG_MAIN).grid(row=1, column=4, sticky="w", padx=(6, 0))
        self.cb_cause = ttk.Combobox(meta_row, values=["", "rework", "Missed inspection"], width=16)
        self.cb_cause.set("")
        self.cb_cause.grid(row=1, column=5, padx=4, pady=2)

        phone_box = tk.LabelFrame(top_ctrl, text=" 2. Live Voice & Mobile HUD (Scan QR) ", font=config.FONT_BODY_BOLD, bg=config.COLOR_BG_MAIN, fg=config.COLOR_PRIMARY, padx=8, pady=6)
        phone_box.pack(side=tk.RIGHT, fill=tk.BOTH, padx=(8, 0))

        # Split phone_box into info left and QR code right
        box_inner = tk.Frame(phone_box, bg=config.COLOR_BG_MAIN)
        box_inner.pack(fill=tk.BOTH, expand=True)

        info_frame = tk.Frame(box_inner, bg=config.COLOR_BG_MAIN)
        info_frame.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 6))

        self.lbl_server_url = tk.Label(info_frame, text=f"{self.app.phone_server_url}", font=("Consolas", 8, "bold"), bg=config.COLOR_BG_MAIN, fg=config.COLOR_PRIMARY, justify=tk.LEFT, wraplength=180)
        self.lbl_server_url.pack(anchor=tk.W, pady=(0, 1))
        self.lbl_photo_status = tk.Label(info_frame, text="Monitoring phone for photos & voice...", font=config.FONT_SMALL, bg=config.COLOR_BG_MAIN, fg=config.COLOR_TEXT_SECONDARY)
        self.lbl_photo_status.pack(anchor=tk.W, pady=1)

        self.lbl_hearing = tk.Label(info_frame, text="App Heard: [None]", font=("Segoe UI", 8, "bold"), bg="#f1f5f9", fg=config.COLOR_PRIMARY, padx=6, pady=2, anchor="w", justify=tk.LEFT, wraplength=180)
        self.lbl_hearing.pack(fill=tk.X, pady=1)

        self.btn_learn_voice = ModernButton(info_frame, text="Save Heard Voice", command=self.save_heard_voice_as_sample, primary=False, padx=4, pady=1)
        self.btn_learn_voice.pack_forget()

        # QR Code Display Right Column
        qr_frame = tk.Frame(box_inner, bg=config.COLOR_BG_MAIN)
        qr_frame.pack(side=tk.RIGHT, padx=2)

        self.lbl_qr_thumb = tk.Label(qr_frame, bg="white", cursor="hand2", relief=tk.SOLID, bd=1)
        self.lbl_qr_thumb.pack()
        self.lbl_qr_thumb.bind("<Button-1>", lambda e: self.show_large_qr())

        tk.Label(qr_frame, text="Tap to expand", font=("Segoe UI", 7), bg=config.COLOR_BG_MAIN, fg=config.COLOR_TEXT_SECONDARY).pack(pady=(1, 0))
        self.update_qr_code(self.app.phone_server_url)


        act_row = tk.Frame(self, bg=config.COLOR_BG_MAIN)
        act_row.pack(fill=tk.X, pady=(0, 8))

        ModernButton(act_row, text="+ Add Review Record", command=self.on_barcode_scanned, primary=True).pack(side=tk.LEFT, padx=(0, 6))
        ModernButton(act_row, text="Update Selected Record", command=self.update_selected_record, primary=True).pack(side=tk.LEFT, padx=(0, 6))
        ModernButton(act_row, text="⚡ Sync to Master Report", command=self.sync_to_master_report, primary=True).pack(side=tk.LEFT, padx=(0, 6))
        ModernButton(act_row, text="Export Styled Defect Excel File", command=self.export_styled_excel, primary=False).pack(side=tk.LEFT, padx=(0, 6))
        ModernButton(act_row, text="Open Voice Samples Folder", command=self.open_voice_samples_folder, primary=False).pack(side=tk.LEFT, padx=(0, 6))
        ModernButton(act_row, text="📐 Reset Col Widths", command=self.reset_column_widths, primary=False).pack(side=tk.LEFT, padx=(0, 6))
        ModernButton(act_row, text="🔄 New Shift Reset", command=self.manual_shift_reset, primary=False).pack(side=tk.LEFT, padx=(0, 6))
        ModernButton(act_row, text="📦 View Archive", command=self.open_archive_history, primary=False).pack(side=tk.LEFT, padx=(0, 6))
        ModernButton(act_row, text="Clear Records", command=self.clear_all_records, primary=False).pack(side=tk.LEFT)

        self.lbl_table_count = tk.Label(act_row, text="Total Records Reviewed: 0", font=config.FONT_BODY_BOLD, bg=config.COLOR_BG_MAIN, fg=config.COLOR_PRIMARY)
        self.lbl_table_count.pack(side=tk.RIGHT, padx=10)

        # 10-Column Table Container with interactive column width adjustment and dual-axis scrolling
        table_container = tk.Frame(self, bg="white", highlightbackground=config.COLOR_DIVIDER, highlightthickness=1)
        table_container.pack(fill=tk.BOTH, expand=True)

        self.col_defs = [
            ("date", "Date", 100),
            ("sn", "SN", 175),
            ("summary", "Defect Summary", 160),
            ("grade", "Grade", 80),
            ("pre_el", "PreEL (Front / EL / Back)", 210),
            ("mr_pic", "MR Defect Pic", 125),
            ("layup_time", "PreEL Time", 155),
            ("station", "Station", 110),
            ("shift", "Respon. Shift", 105),
            ("line", "Line", 85)
        ]
        saved_widths = getattr(config, 'REVIEW_TABLE_COL_WIDTHS', {})
        self.col_widths = {k: saved_widths.get(k, def_w) for k, _, def_w in self.col_defs}
        self.header_cells = {}

        total_w = self.get_total_table_width()

        self.table_canvas = tk.Canvas(table_container, bg="#f8fafc", highlightthickness=0)
        v_scroll = ttk.Scrollbar(table_container, orient="vertical", command=self.table_canvas.yview)
        h_scroll = ttk.Scrollbar(table_container, orient="horizontal", command=self.table_canvas.xview)
        
        self.table_inner = tk.Frame(self.table_canvas, bg="#f8fafc", width=total_w)
        self.table_window = self.table_canvas.create_window((0, 0), window=self.table_inner, anchor="nw", width=total_w)
        
        self.table_canvas.configure(xscrollcommand=h_scroll.set, yscrollcommand=v_scroll.set)
        
        self.table_canvas.grid(row=0, column=0, sticky="nsew")
        v_scroll.grid(row=0, column=1, sticky="ns")
        h_scroll.grid(row=1, column=0, sticky="ew")
        table_container.grid_rowconfigure(0, weight=1)
        table_container.grid_columnconfigure(0, weight=1)

        def _on_canvas_configure(event):
            cw = max(event.width, self.get_total_table_width())
            self.table_canvas.itemconfig(self.table_window, width=cw)
            self.table_canvas.configure(scrollregion=self.table_canvas.bbox("all"))

        self.table_canvas.bind("<Configure>", _on_canvas_configure)
        self.table_inner.bind("<Configure>", lambda e: self.table_canvas.configure(scrollregion=self.table_canvas.bbox("all")))

        def _on_table_mousewheel(event):
            try:
                x, y = event.x_root, event.y_root
                tc_x = table_container.winfo_rootx()
                tc_y = table_container.winfo_rooty()
                tc_w = table_container.winfo_width()
                tc_h = table_container.winfo_height()
                if tc_x <= x <= tc_x + tc_w and tc_y <= y <= tc_y + tc_h:
                    self.table_canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")
            except Exception:
                pass

        self.table_canvas.bind_all("<MouseWheel>", _on_table_mousewheel)

        # Header bar with draggable column splitters
        self.header_bar = tk.Frame(self.table_inner, bg="#1e293b", height=36, width=total_w)
        self.header_bar.pack(anchor="nw", fill=tk.X)
        self.header_bar.pack_propagate(False)

        for col_key, title, _ in self.col_defs:
            w = self.col_widths.get(col_key, 100)
            cell = tk.Frame(self.header_bar, width=w, height=36, bg="#1e293b")
            cell.pack(side=tk.LEFT, padx=0)
            cell.pack_propagate(False)
            self.header_cells[col_key] = cell
            lbl = tk.Label(cell, text=title, font=("Segoe UI", 9, "bold"), bg="#1e293b", fg="#ffffff")
            lbl.pack(fill=tk.BOTH, expand=True)

            splitter = tk.Frame(self.header_bar, width=5, height=36, bg="#334155", cursor="sb_h_double_arrow")
            splitter.pack(side=tk.LEFT, fill=tk.Y)
            self._bind_splitter(splitter, col_key)

        # Rows container
        self.rows_frame = tk.Frame(self.table_inner, bg="#f8fafc", width=total_w)
        self.rows_frame.pack(anchor="nw", fill=tk.X)

    def get_total_table_width(self):
        return sum(self.col_widths.get(k, 100) for k, _, _ in self.col_defs) + (len(self.col_defs) * 5) + 20

    def _bind_splitter(self, splitter, col_key):
        def on_enter(e):
            splitter.config(bg="#60a5fa")
        def on_leave(e):
            splitter.config(bg="#334155")
        def on_press(e):
            self._drag_col_key = col_key
            self._drag_start_x = e.x_root
            self._drag_start_w = self.col_widths.get(col_key, 100)
        def on_motion(e):
            if not getattr(self, '_drag_col_key', None) or self._drag_col_key != col_key:
                return
            delta = e.x_root - self._drag_start_x
            new_w = max(45, self._drag_start_w + delta)
            self.update_col_width(col_key, new_w)
        def on_release(e):
            self._drag_col_key = None
            self.save_column_widths()

        splitter.bind("<Enter>", on_enter)
        splitter.bind("<Leave>", on_leave)
        splitter.bind("<ButtonPress-1>", on_press)
        splitter.bind("<B1-Motion>", on_motion)
        splitter.bind("<ButtonRelease-1>", on_release)

    def update_col_width(self, col_key, new_w):
        self.col_widths[col_key] = new_w
        if col_key in self.header_cells:
            self.header_cells[col_key].config(width=new_w)
        
        for row in self.rows_frame.winfo_children():
            if hasattr(row, '_cells') and col_key in row._cells:
                row._cells[col_key].config(width=new_w)
                
        total_w = self.get_total_table_width()
        self.header_bar.config(width=total_w)
        self.rows_frame.config(width=total_w)
        self.table_inner.config(width=total_w)
        if hasattr(self, 'table_canvas') and hasattr(self, 'table_window'):
            cw = max(self.table_canvas.winfo_width(), total_w)
            self.table_canvas.itemconfig(self.table_window, width=cw)
            self.table_canvas.configure(scrollregion=self.table_canvas.bbox("all"))

    def save_column_widths(self):
        config.REVIEW_TABLE_COL_WIDTHS.update(self.col_widths)
        config.save_persistent_settings()

    def reset_column_widths(self):
        for col_key, _, default_w in self.col_defs:
            self.col_widths[col_key] = default_w
        self.save_column_widths()
        for col_key, w in self.col_widths.items():
            if col_key in self.header_cells:
                self.header_cells[col_key].config(width=w)
        self.refresh_treeview()

    def update_qr_code(self, url=None):
        clean_url = get_mobile_hud_url()
        qr_img = generate_qr_image(clean_url, size=85)
        if qr_img:
            self.tk_qr_thumb = ImageTk.PhotoImage(qr_img)
            if hasattr(self, 'lbl_qr_thumb'):
                self.lbl_qr_thumb.config(image=self.tk_qr_thumb)

    def show_large_qr(self):
        QRCodeDialog(self.winfo_toplevel(), get_mobile_hud_url())


    def show_image_viewer(self, title, img_path, record=None, image_category="MR_DEFECT"):
        if not img_path:
            messagebox.showinfo("Image Viewer", "No image available.")
            return

        target = img_path
        if not os.path.exists(target):
            fname = os.path.basename(img_path)
            c1 = os.path.join(getattr(config, 'DAILY_CACHE_DIR', ''), fname)
            c2 = os.path.join(getattr(config, 'LOCAL_DATA_DIR', ''), "PhoneUploads", fname)
            if os.path.exists(c1):
                target = c1
            elif os.path.exists(c2):
                target = c2
            else:
                # Also check FolderArchives subdirectories
                archive_base = os.path.join(getattr(config, 'LOCAL_DATA_DIR', ''), "FolderArchives")
                found = False
                if os.path.exists(archive_base):
                    try:
                        for sub in os.listdir(archive_base):
                            sub_p = os.path.join(archive_base, sub, fname)
                            if os.path.exists(sub_p):
                                target = sub_p
                                found = True
                                break
                    except Exception: pass
                if not found:
                    messagebox.showinfo("Image Viewer", f"Image file not found:\n{img_path}")
                    return

        try:
            ImageSnipperDialog(self.winfo_toplevel(), title, target, record=record, on_snip_saved=self.on_snip_saved, image_category=image_category)
        except Exception as e:
            try:
                os.startfile(target)
            except Exception:
                messagebox.showerror("Error", f"Could not open photo:\n{e}")

    def on_snip_saved(self, rec, snip_path, image_category="MR_DEFECT"):
        if rec:
            rec['worked_on'] = True
            cat = str(image_category).upper()
            if "FRONT" in cat or "BACK" in cat or "PRE" in cat or ("EL" in cat and "DEFECT" not in cat):
                rec['pre_el_snip_path'] = snip_path
                if "FRONT" in cat:
                    v_key = "front"
                elif "BACK" in cat:
                    v_key = "back"
                elif "EL" in cat:
                    v_key = "el"
                else:
                    v_key = "front"
                rec['pre_el_snip_view'] = v_key
                if 'pre_el_views' in rec and isinstance(rec['pre_el_views'], dict):
                    if v_key in rec['pre_el_views'] and rec['pre_el_views'][v_key]:
                        rec['pre_el_views'][v_key]['snip_path'] = snip_path
                self.register_photo_alias(rec.get('id'), snip_path)
            else:
                rec['photo_path'] = snip_path
                rec['snip_saved'] = True
                self.register_photo_alias(rec.get('id'), snip_path)
            self.refresh_treeview()


    def select_record(self, rec):
        self.selected_rec_id = rec.get('id')
        self.ent_sn.delete(0, tk.END)
        sn_val = rec.get('sn', '')
        if sn_val and sn_val != "Pending SN":
            self.ent_sn.insert(0, sn_val)
        
        cls_val = rec.get('class', '')
        if cls_val and cls_val in config.DEFECT_TREE:
            self.cb_class.set(cls_val)
            self.on_class_changed()
        else:
            self.cb_class.set("")
            
        sum_val = rec.get('summary', '')
        self.cb_summary.set(sum_val if sum_val and sum_val != "-" else "")
        self.cb_result.set(rec.get('result', '') if rec.get('result', '') != "-" else "")
        self.cb_cause.set(rec.get('cause', '') if rec.get('cause', '') != "-" else "")
        self.cb_line.set(rec.get('line', '') if rec.get('line') else "")
        
        if rec.get('audio_path'): self.last_audio_path = rec.get('audio_path')
        if rec.get('raw_voice'): self.last_spoken_raw = rec.get('raw_voice')
        self.refresh_treeview()

    def update_selected_record(self):
        target_rec = None
        if self.selected_rec_id:
            for r in self.records:
                if r.get('id') == self.selected_rec_id:
                    target_rec = r
                    break
        if not target_rec and self.records:
            target_rec = self.records[0]

        if not target_rec:
            messagebox.showinfo("Update Record", "Please select a record from the table to update.")
            return

        new_sn = self.ent_sn.get().strip()
        new_sum = self.cb_summary.get().strip()
        new_cls = self.cb_class.get().strip()
        new_res = self.cb_result.get().strip()
        new_cause = self.cb_cause.get().strip()
        new_line = self.cb_line.get().strip()

        if new_sn and new_sn != target_rec.get('sn'):
            target_rec['sn'] = new_sn
            target_rec['order'] = new_sn[:9] if len(new_sn) >= 9 else new_sn
            if target_rec.get('photo_path'):
                target_rec['photo_path'] = rename_to_sn_pattern(target_rec['photo_path'], "Def", new_sn)
                self.register_photo_alias(target_rec['id'], target_rec['photo_path'])
            if target_rec.get('sn_photo_path'):
                target_rec['sn_photo_path'] = rename_to_sn_pattern(target_rec['sn_photo_path'], "SN", new_sn)
                self.register_photo_alias(target_rec['id'], target_rec['sn_photo_path'])
            if target_rec.get('pre_el_snip_path'):
                target_rec['pre_el_snip_path'] = rename_to_sn_pattern(target_rec['pre_el_snip_path'], "PreEL", new_sn)
                self.register_photo_alias(target_rec['id'], target_rec['pre_el_snip_path'])
            
            lt, st, ln, sh, views = self.query_pre_el(new_sn, target_rec.get('dt', datetime.now()))
            target_rec['layup_time'] = lt
            target_rec['station'] = st
            target_rec['shift'] = sh
            target_rec['pre_el_views'] = views
            if not new_line and ln != "-": target_rec['line'] = ln

        if new_sum: target_rec['summary'] = new_sum
        if new_cls: target_rec['class'] = new_cls
        if new_res: target_rec['result'] = new_res
        if new_cause: target_rec['cause'] = new_cause
        if new_line: target_rec['line'] = new_line

        rec_audio = target_rec.get('audio_path') or self.last_audio_path
        rec_spoken = target_rec.get('raw_voice') or self.last_spoken_raw
        if new_sum and (rec_audio or rec_spoken):
            train_voice_pattern(rec_spoken or new_sum, new_sum, rec_audio)
            self.lbl_hearing.config(text=f"Auto-Trained! '{rec_spoken or new_sum}' -> {new_sum}", fg=config.COLOR_STATUS_OK)

        target_rec['worked_on'] = True
        self.refresh_treeview()
        messagebox.showinfo("Record Updated", f"Successfully updated record for {target_rec.get('sn')}!")

    def open_voice_samples_folder(self):
        os.makedirs(config.VOICE_SAMPLES_DIR, exist_ok=True)
        try:
            os.startfile(config.VOICE_SAMPLES_DIR)
        except Exception as e:
            messagebox.showinfo("Voice Samples", f"Voice samples folder is located at:\n{config.VOICE_SAMPLES_DIR}")

    def on_class_changed(self):
        sel_class = self.cb_class.get()
        if sel_class in config.DEFECT_TREE:
            summaries = [""] + config.DEFECT_TREE[sel_class]
            self.cb_summary['values'] = summaries
            self.cb_summary.set("")
        else:
            self.cb_summary['values'] = [""]
            self.cb_summary.set("")

    def refresh_defect_dropdowns(self):
        class_options = [""] + list(config.DEFECT_TREE.keys())
        self.cb_class['values'] = class_options
        self.on_class_changed()

    def handle_mobile_defect_logged(self, data):
        def_sn = data.get('sn', '').strip()
        def_class = data.get('class', '')
        def_summary = data.get('summary', '')
        def_result = data.get('result', '')
        def_dt = data.get('dt') or datetime.now()

        if def_class:
            self.cb_class.set(def_class)
            self.on_class_changed()
        if def_summary:
            self.cb_summary.set(def_summary)
        if def_result:
            self.cb_result.set(def_result)

        target_rec = self.find_or_create_panel_session("MOBILE_DEFECT", def_dt, sn=def_sn)
        if (not target_rec.get('sn') or target_rec.get('sn') in ("Pending SN", "-", "")) and def_sn:
            target_rec['sn'] = def_sn
            target_rec['order'] = def_sn[:9] if len(def_sn) >= 9 else def_sn
        target_rec['summary'] = def_summary
        target_rec['class'] = def_class
        target_rec['result'] = def_result
        target_rec['worked_on'] = True
        self.lbl_hearing.config(text=f"📱 Mobile Logged:\n{def_summary} ({def_result})", fg=config.COLOR_STATUS_OK)
        self.refresh_treeview()

    def save_heard_voice_as_sample(self):
        target_summary = self.cb_summary.get().strip()
        if not target_summary:
            messagebox.showwarning("Voice Training", "Please select a Defect Summary in the dropdown first to map this voice sample to!")
            return
        if not self.last_spoken_raw:
            messagebox.showwarning("Voice Training", "No recent voice recording detected to save!")
            return

        train_voice_pattern(self.last_spoken_raw, target_summary, self.last_audio_path)
        self.btn_learn_voice.pack_forget()
        self.lbl_hearing.config(text=f"Learned! '{self.last_spoken_raw}' -> {target_summary}", fg=config.COLOR_STATUS_OK)
        messagebox.showinfo("Trained", f"Successfully saved '{self.last_spoken_raw}' as a training sample for '{target_summary}'!")

    def find_or_create_panel_session(self, item_type: str, item_dt: datetime, sn: str = None, file_path: str = None):
        WINDOW_SECONDS = 120  # 2 minutes window per panel (SN photo, Defect photo, Mobile defect choice)

        # 1. Match explicit SN if provided
        if sn and sn not in ("Pending SN", "-", ""):
            for rec in self.records:
                if rec.get('sn') == sn:
                    return rec

        # 2. Check alias mapping and existing assigned photo paths
        if file_path:
            fname = os.path.basename(file_path)
            rec_id = self.processed_photo_aliases.get(file_path) or self.processed_photo_aliases.get(fname)
            if rec_id:
                for rec in self.records:
                    if rec.get('id') == rec_id:
                        return rec
            for rec in self.records:
                p1 = os.path.basename(rec.get('photo_path', ''))
                p2 = os.path.basename(rec.get('sn_photo_path', ''))
                if fname and (fname == p1 or fname == p2):
                    return rec

        # 3. Look through existing records (newest to oldest) for an open session within 120s
        for rec in self.records:
            rec_dt = rec.get('dt')
            if not rec_dt: continue
            time_diff = abs((item_dt - rec_dt).total_seconds())
            if time_diff <= WINDOW_SECONDS:
                if item_type in ("MOBILE_DEFECT", "VOICE"):
                    if not rec.get('summary') or rec.get('summary') in ("-", "", "Unrecognized"):
                        return rec
                    if rec.get('sn') in ("Pending SN", "-", "") or not rec.get('photo_path'):
                        return rec

                elif item_type == "DEFECT_PHOTO":
                    if not rec.get('photo_path'):
                        return rec

                elif item_type == "SN_PHOTO":
                    if rec.get('sn') in ("Pending SN", "-", ""):
                        return rec

        # 4. Check top record if within 120s and still has open slot
        if self.records:
            top_rec = self.records[0]
            top_dt = top_rec.get('dt')
            if top_dt and abs((item_dt - top_dt).total_seconds()) <= WINDOW_SECONDS:
                if item_type in ("MOBILE_DEFECT", "VOICE") and (not top_rec.get('summary') or top_rec.get('summary') in ("-", "", "Unrecognized")):
                    return top_rec
                if item_type == "DEFECT_PHOTO" and not top_rec.get('photo_path'):
                    return top_rec
                if item_type == "SN_PHOTO" and top_rec.get('sn') in ("Pending SN", "-", ""):
                    return top_rec

        # 5. Exceeds 2 minutes or session already filled -> create a brand new panel record
        eval_shift = get_evaluation_shift(item_dt)
        new_rec = {
            'id': f"rec_{int(time.time()*1000)}_{len(self.records)}",
            'dt': item_dt,
            'date': item_dt.strftime("%m/%d/%Y"),
            'order': sn[:9] if (sn and len(sn) >= 9) else "",
            'sn': sn or "Pending SN",
            'summary': "",
            'class': "",
            'result': "",

            'cause': "",
            'photo_path': "",
            'sn_photo_path': "",
            'pre_el_snip_path': "",
            'pre_el_snip_view': "",
            'layup_time': "",
            'station': "",
            'eval_shift': eval_shift,
            'line': "",
            'shift': "",
            'pre_el_views': {'front': None, 'el': None, 'back': None}
        }
        self.records.insert(0, new_rec)
        return new_rec

    def _create_row_frame(self, rec, idx, total_w):
        row_frame = tk.Frame(self.rows_frame, height=40, width=total_w)
        row_frame.pack(anchor="nw", fill=tk.X, pady=1)
        row_frame.pack_propagate(False)
        row_frame._rec_id = rec.get('id')
        row_frame._cells = {}
        row_frame._widget_refs = {}
        cell_frames = []

        def bind_select(w, r=rec):
            w.bind("<Button-1>", lambda e, rec_item=r: self.select_record(rec_item))

        bind_select(row_frame)

        # 0. Date
        c0 = tk.Frame(row_frame, width=self.col_widths.get("date", 100), height=38)
        c0.pack(side=tk.LEFT, padx=0)
        c0.pack_propagate(False)
        row_frame._cells["date"] = c0
        cell_frames.append(c0)
        l0 = tk.Label(c0, font=("Segoe UI", 9))
        l0.pack(fill=tk.BOTH, expand=True)
        bind_select(l0)
        sp0 = tk.Frame(row_frame, width=5, height=38)
        sp0.pack(side=tk.LEFT, fill=tk.Y)
        cell_frames.append(sp0)

        # 1. SN
        c1 = tk.Frame(row_frame, width=self.col_widths.get("sn", 175), height=38)
        c1.pack(side=tk.LEFT, padx=0)
        c1.pack_propagate(False)
        row_frame._cells["sn"] = c1
        cell_frames.append(c1)
        l1 = tk.Label(c1, font=("Segoe UI", 9, "bold"))
        l1.pack(fill=tk.BOTH, expand=True)
        bind_select(l1)
        sp1 = tk.Frame(row_frame, width=5, height=38)
        sp1.pack(side=tk.LEFT, fill=tk.Y)
        cell_frames.append(sp1)

        # 2. Defect Summary
        c2 = tk.Frame(row_frame, width=self.col_widths.get("summary", 160), height=38)
        c2.pack(side=tk.LEFT, padx=0)
        c2.pack_propagate(False)
        row_frame._cells["summary"] = c2
        cell_frames.append(c2)
        l2 = tk.Label(c2, font=("Segoe UI", 9))
        l2.pack(fill=tk.BOTH, expand=True)
        bind_select(l2)
        sp2 = tk.Frame(row_frame, width=5, height=38)
        sp2.pack(side=tk.LEFT, fill=tk.Y)
        cell_frames.append(sp2)

        # 3. Grade (Q3 / Scrap)
        c3 = tk.Frame(row_frame, width=self.col_widths.get("grade", 80), height=38)
        c3.pack(side=tk.LEFT, padx=0)
        c3.pack_propagate(False)
        row_frame._cells["grade"] = c3
        cell_frames.append(c3)
        l3 = tk.Label(c3, font=("Segoe UI", 9))
        l3.pack(expand=True)
        bind_select(l3)
        sp3 = tk.Frame(row_frame, width=5, height=38)
        sp3.pack(side=tk.LEFT, fill=tk.Y)
        cell_frames.append(sp3)

        # 4. PreEL (3 Buttons: Front, EL, Back)
        c4 = tk.Frame(row_frame, width=self.col_widths.get("pre_el", 210), height=38)
        c4.pack(side=tk.LEFT, padx=0)
        c4.pack_propagate(False)
        row_frame._cells["pre_el"] = c4
        cell_frames.append(c4)
        f_btn_box = tk.Frame(c4)
        f_btn_box.pack(expand=True)
        btn_f = tk.Button(f_btn_box, text="Front", font=("Segoe UI", 8), padx=4, pady=2, relief=tk.FLAT)
        btn_f.pack(side=tk.LEFT, padx=1)
        btn_el = tk.Button(f_btn_box, text="EL", font=("Segoe UI", 8), padx=4, pady=2, relief=tk.FLAT)
        btn_el.pack(side=tk.LEFT, padx=1)
        btn_b = tk.Button(f_btn_box, text="Back", font=("Segoe UI", 8), padx=4, pady=2, relief=tk.FLAT)
        btn_b.pack(side=tk.LEFT, padx=1)
        sp4 = tk.Frame(row_frame, width=5, height=38)
        sp4.pack(side=tk.LEFT, fill=tk.Y)
        cell_frames.append(sp4)

        # 5. MR Defect Pic
        c5 = tk.Frame(row_frame, width=self.col_widths.get("mr_pic", 125), height=38)
        c5.pack(side=tk.LEFT, padx=0)
        c5.pack_propagate(False)
        row_frame._cells["mr_pic"] = c5
        cell_frames.append(c5)
        btn_mr = tk.Button(c5, text="No Pic", font=("Segoe UI", 8), padx=6, pady=2, relief=tk.FLAT)
        btn_mr.pack(expand=True)
        sp5 = tk.Frame(row_frame, width=5, height=38)
        sp5.pack(side=tk.LEFT, fill=tk.Y)
        cell_frames.append(sp5)

        # 6. PreEL Time
        c6 = tk.Frame(row_frame, width=self.col_widths.get("layup_time", 155), height=38)
        c6.pack(side=tk.LEFT, padx=0)
        c6.pack_propagate(False)
        row_frame._cells["layup_time"] = c6
        cell_frames.append(c6)
        l6 = tk.Label(c6, font=("Segoe UI", 9))
        l6.pack(fill=tk.BOTH, expand=True)
        bind_select(l6)
        sp6 = tk.Frame(row_frame, width=5, height=38)
        sp6.pack(side=tk.LEFT, fill=tk.Y)
        cell_frames.append(sp6)

        # 7. Station
        c7 = tk.Frame(row_frame, width=self.col_widths.get("station", 110), height=38)
        c7.pack(side=tk.LEFT, padx=0)
        c7.pack_propagate(False)
        row_frame._cells["station"] = c7
        cell_frames.append(c7)
        l7 = tk.Label(c7, font=("Segoe UI", 9))
        l7.pack(fill=tk.BOTH, expand=True)
        bind_select(l7)
        sp7 = tk.Frame(row_frame, width=5, height=38)
        sp7.pack(side=tk.LEFT, fill=tk.Y)
        cell_frames.append(sp7)

        # 8. Respon. Shift
        c8 = tk.Frame(row_frame, width=self.col_widths.get("shift", 105), height=38)
        c8.pack(side=tk.LEFT, padx=0)
        c8.pack_propagate(False)
        row_frame._cells["shift"] = c8
        cell_frames.append(c8)
        l8 = tk.Label(c8, font=("Segoe UI", 9, "bold"))
        l8.pack(fill=tk.BOTH, expand=True)
        bind_select(l8)
        sp8 = tk.Frame(row_frame, width=5, height=38)
        sp8.pack(side=tk.LEFT, fill=tk.Y)
        cell_frames.append(sp8)

        # 9. Line
        c9 = tk.Frame(row_frame, width=self.col_widths.get("line", 85), height=38)
        c9.pack(side=tk.LEFT, padx=0)
        c9.pack_propagate(False)
        row_frame._cells["line"] = c9
        cell_frames.append(c9)
        l9 = tk.Label(c9, font=("Segoe UI", 9, "bold"))
        l9.pack(fill=tk.BOTH, expand=True)
        bind_select(l9)

        row_frame._widget_refs = {
            'cell_frames': cell_frames,
            'l0': l0, 'l1': l1, 'l2': l2, 'l3': l3,
            'f_btn_box': f_btn_box, 'btn_f': btn_f, 'btn_el': btn_el, 'btn_b': btn_b,
            'btn_mr': btn_mr, 'l6': l6, 'l7': l7, 'l8': l8, 'l9': l9
        }

        self._update_row_widgets(row_frame, rec, idx)
        return row_frame

    def _update_row_widgets(self, row_frame, rec, idx):
        is_selected = (rec.get('id') == self.selected_rec_id)
        
        # Check if this row has been worked on / processed
        def_photo = rec.get('photo_path')
        is_def_snip = bool(def_photo and ("Snip_" in os.path.basename(def_photo) or rec.get('snip_saved')))
        has_pre_snip = bool(rec.get('pre_el_snip_path') and os.path.exists(rec.get('pre_el_snip_path', '')))
        has_summary = bool(rec.get('summary') and str(rec.get('summary')).strip() not in ("-", "None", ""))
        has_grade = bool(rec.get('result') and str(rec.get('result')).strip() not in ("-", "None", ""))
        is_worked = bool(rec.get('worked_on') or is_def_snip or has_pre_snip or has_summary or has_grade)

        if is_selected:
            row_bg = "#dbeafe"  # Active selection highlight (soft sky blue)
            border_col = "#2563eb"
            border_thick = 2
        elif is_worked:
            row_bg = "#f0fdf4" if idx % 2 == 0 else "#ecfdf5"  # Soft mint/emerald highlight for worked rows
            border_col = "#10b981"
            border_thick = 1
        else:
            row_bg = "#ffffff" if idx % 2 == 0 else "#f8fafc"  # Standard alternating for pending rows
            border_col = "#e2e8f0"
            border_thick = 1

        w = getattr(row_frame, '_widget_refs', {})
        row_frame.config(bg=row_bg, highlightbackground=border_col, highlightthickness=border_thick)
        
        for cf in w.get('cell_frames', []):
            cf.config(bg=row_bg)

        # 0. Date
        date_icon = "✓ " if is_worked else ""
        date_fg = "#047857" if (is_worked and not is_selected) else ("#1e3a8a" if is_selected else "#334155")
        w['l0'].config(text=f"{date_icon}{rec.get('date', '-')}", font=("Segoe UI", 9, "bold" if is_worked else "normal"), fg=date_fg, bg=row_bg)

        # 1. SN
        sn_fg = "#065f46" if (is_worked and not is_selected) else ("#1e3a8a" if is_selected else "#1e293b")
        w['l1'].config(text=rec.get('sn', '-'), fg=sn_fg, bg=row_bg)

        # 2. Defect Summary
        w['l2'].config(text=rec.get('summary') or "-", bg=row_bg)

        # 3. Grade (Q3 / Scrap)
        res_val = str(rec.get('result', '') or '').strip()
        if res_val in ("-", "None"): res_val = ""
        if "SCRAP" in res_val.upper():
            w['l3'].config(text=res_val, font=("Segoe UI", 9, "bold"), bg="#fee2e2", fg="#dc2626")
        elif "Q3" in res_val.upper():
            w['l3'].config(text=res_val, font=("Segoe UI", 9, "bold"), bg="#fef3c7", fg="#ca8a04")
        else:
            w['l3'].config(text=res_val, font=("Segoe UI", 9), bg=row_bg, fg="#64748b")

        # 4. PreEL (Front, EL, Back)
        views = rec.get('pre_el_views') or {}
        front_info = views.get('front')
        el_info = views.get('el')
        back_info = views.get('back')
        pre_snip = rec.get('pre_el_snip_path')
        pre_view = str(rec.get('pre_el_snip_view', '')).lower()
        snip_name = os.path.basename(pre_snip).lower() if pre_snip else ""

        w['f_btn_box'].config(bg=row_bg)

        # Front Button
        if front_info and front_info.get('path'):
            st = str(front_info.get('status', 'OK')).upper()
            is_ng = (st == "NG" or st not in ("OK", "PASS", "GOOD", "NORMAL"))
            has_snip = bool(pre_snip and os.path.exists(pre_snip) and (pre_view == 'front' or front_info.get('snip_path') or 'front' in snip_name))
            btn_txt = "✂️ Front" if has_snip else "Front"
            bg_col = "#dc2626" if is_ng else "#059669"
            act_col = "#b91c1c" if is_ng else "#047857"
            w['btn_f'].config(text=btn_txt, font=("Segoe UI", 8, "bold"), bg=bg_col, fg="#ffffff", activebackground=act_col, activeforeground="#ffffff", state="normal", cursor="hand2",
                              command=lambda p=front_info['path'], r=rec: self.show_image_viewer("Pre-EL Front Image", p, record=r, image_category="PRE_EL_FRONT"))
        else:
            w['btn_f'].config(text="Front", font=("Segoe UI", 8), bg="#cbd5e1", fg="#64748b", state="disabled", cursor="")

        # EL Button
        if el_info and el_info.get('path'):
            st = str(el_info.get('status', 'OK')).upper()
            is_ng = (st == "NG" or st not in ("OK", "PASS", "GOOD", "NORMAL"))
            has_snip = bool(pre_snip and os.path.exists(pre_snip) and (pre_view == 'el' or el_info.get('snip_path') or ('_el_' in snip_name or snip_name.endswith('_el.jpg'))))
            btn_txt = "✂️ EL" if has_snip else "EL"
            bg_col = "#dc2626" if is_ng else "#059669"
            act_col = "#b91c1c" if is_ng else "#047857"
            w['btn_el'].config(text=btn_txt, font=("Segoe UI", 8, "bold"), bg=bg_col, fg="#ffffff", activebackground=act_col, activeforeground="#ffffff", state="normal", cursor="hand2",
                               command=lambda p=el_info['path'], r=rec: self.show_image_viewer("Pre-EL EL Image", p, record=r, image_category="PRE_EL_EL"))
        else:
            w['btn_el'].config(text="EL", font=("Segoe UI", 8), bg="#cbd5e1", fg="#64748b", state="disabled", cursor="")

        # Back Button
        if back_info and back_info.get('path'):
            st = str(back_info.get('status', 'OK')).upper()
            is_ng = (st == "NG" or st not in ("OK", "PASS", "GOOD", "NORMAL"))
            has_snip = bool(pre_snip and os.path.exists(pre_snip) and (pre_view == 'back' or back_info.get('snip_path') or 'back' in snip_name))
            btn_txt = "✂️ Back" if has_snip else "Back"
            bg_col = "#dc2626" if is_ng else "#059669"
            act_col = "#b91c1c" if is_ng else "#047857"
            w['btn_b'].config(text=btn_txt, font=("Segoe UI", 8, "bold"), bg=bg_col, fg="#ffffff", activebackground=act_col, activeforeground="#ffffff", state="normal", cursor="hand2",
                              command=lambda p=back_info['path'], r=rec: self.show_image_viewer("Pre-EL Back Image", p, record=r, image_category="PRE_EL_BACK"))
        else:
            w['btn_b'].config(text="Back", font=("Segoe UI", 8), bg="#cbd5e1", fg="#64748b", state="disabled", cursor="")

        # 5. MR Defect Pic
        def_photo = rec.get('photo_path')
        if def_photo and os.path.exists(def_photo):
            is_snip = "Snip_" in os.path.basename(def_photo) or rec.get('snip_saved')
            btn_text = "✂️ Defect Snip" if is_snip else "📷 Defect Pic"
            btn_bg = "#059669" if is_snip else "#2563eb"
            btn_act = "#047857" if is_snip else "#1d4ed8"
            w['btn_mr'].config(text=btn_text, font=("Segoe UI", 8, "bold"), bg=btn_bg, fg="#ffffff", activebackground=btn_act, activeforeground="#ffffff", state="normal", cursor="hand2",
                               command=lambda p=def_photo, r=rec: self.show_image_viewer("MR Defect Photo", p, record=r, image_category="MR_DEFECT"))
        else:
            w['btn_mr'].config(text="No Pic", font=("Segoe UI", 8), bg="#e2e8f0", fg="#94a3b8", state="disabled", cursor="")

        # 6. PreEL Time
        val6 = rec.get('layup_time', '')
        if val6 in ("-", "None", None): val6 = ""
        w['l6'].config(text=val6, bg=row_bg)

        # 7. Station
        val7 = rec.get('station', '')
        if val7 in ("-", "None", None): val7 = ""
        w['l7'].config(text=val7, bg=row_bg)

        # 8. Respon. Shift
        val8 = rec.get('shift', '')
        if val8 in ("-", "None", None): val8 = ""
        w['l8'].config(text=val8, bg=row_bg)

        # 9. Line
        val9 = rec.get('line', '')
        if val9 in ("-", "None", None): val9 = ""
        w['l9'].config(text=val9, bg=row_bg)

    def refresh_treeview(self):
        total_w = self.get_total_table_width()
        self.header_bar.config(width=total_w)
        self.rows_frame.config(width=total_w)
        self.table_inner.config(width=total_w)

        existing_frames = [w for w in self.rows_frame.winfo_children() if hasattr(w, '_rec_id')]
        existing_ids = [getattr(w, '_rec_id', None) for w in existing_frames]
        current_ids = [r.get('id') for r in self.records]

        if existing_ids == current_ids and len(existing_frames) == len(self.records):
            for idx, (rec, row_frame) in enumerate(zip(self.records, existing_frames)):
                self._update_row_widgets(row_frame, rec, idx)
        else:
            for widget in self.rows_frame.winfo_children():
                widget.destroy()

            for idx, rec in enumerate(self.records):
                self._create_row_frame(rec, idx, total_w)

        self.lbl_table_count.config(text=f"Total Records Reviewed: {len(self.records)}")
        self.app.global_review_records = self.records
        try:
            self.table_canvas.configure(scrollregion=self.table_canvas.bbox("all"))
        except Exception:
            pass

        try:
            last_rec = self.records[0] if self.records else None
            update_live_hud_state(last_rec=last_rec, all_records=self.records)
            save_module_review_history(self.records)
        except Exception:
            pass

    def query_pre_el(self, sn, session_dt):
        t_start_all = time.time()
        pre_stations = [f"{config.PRE_EL_STATION_PREFIX}{i}" for i in range(config.PRE_EL_STATION_MIN, config.PRE_EL_STATION_MAX + 1)]
        
        # 1. Before getting time and station, check if MES has layup time for this SN
        # Enables automatic detection of modules produced in prior months (e.g. 2am 8/14 -> window 8/13 to 8/15)
        t0 = time.time()
        mes_dt, mes_station, mes_data = get_mes_layup_time_for_sn(sn)
        t_mes = time.time() - t0

        if mes_dt:
            # Search time frame based on day-1 to day+1 around actual production date
            start_search = mes_dt.replace(hour=0, minute=0, second=0) - timedelta(days=1)
            end_search = mes_dt.replace(hour=23, minute=59, second=59) + timedelta(days=1)
            print(f"[MES SMART SEARCH]: SN '{sn}' MES Layup Time: {mes_dt.strftime('%Y-%m-%d %H:%M:%S')}. Searching Pre-EL window: {start_search.strftime('%Y-%m-%d')} to {end_search.strftime('%Y-%m-%d')}")
        else:
            days_back = getattr(config, 'DEFAULT_PRE_FINAL_DAYS_BACK_START', 4)
            start_search = session_dt - timedelta(days=days_back)
            end_search = session_dt + timedelta(days=1)
            print(f"[PRE-EL DEFAULT SEARCH]: SN '{sn}' searching default {days_back} days back from {session_dt.strftime('%Y-%m-%d')}")

        target_stations = get_pre_el_stations_for_layup_station(mes_station, pre_stations)

        t1 = time.time()
        try:
            results, _ = self.app.search_engine.search_pre_el(sn, target_stations, start_search, end_search, layup_dt=mes_dt, max_results=100)
        except Exception as e:
            print(f"[PRE-EL SEARCH ENGINE ERROR]: {e}")
            results = []
        t_scan = time.time() - t1
        print(f"[PRE-EL TOTAL TIMING]: SN '{sn}' query completed in {time.time() - t_start_all:.3f}s (MES: {t_mes:.3f}s, SMB Scan: {t_scan:.3f}s)")

        pre_el_views = {
            'front': None,
            'el': None,
            'back': None
        }

        if results:
            latest = sorted(results, key=lambda x: x['timestamp'], reverse=True)[0]
            layup_dt = latest['datetime']
            station_found = latest.get('station', '')
            line_str, shift_str = get_responsible_shift_and_line(layup_dt, station_found)

            fronts = [r for r in results if r.get('category') == config.CATEGORY_FRONT]
            els = [r for r in results if r.get('category') == config.CATEGORY_EL]
            backs = [r for r in results if r.get('category') == config.CATEGORY_BACK]

            if fronts:
                f_top = sorted(fronts, key=lambda x: x['timestamp'], reverse=True)[0]
                pre_el_views['front'] = {'path': f_top['path'], 'status': f_top['status'], 'time': f_top['datetime']}
            if els:
                el_top = sorted(els, key=lambda x: x['timestamp'], reverse=True)[0]
                pre_el_views['el'] = {'path': el_top['path'], 'status': el_top['status'], 'time': el_top['datetime']}
            if backs:
                b_top = sorted(backs, key=lambda x: x['timestamp'], reverse=True)[0]
                pre_el_views['back'] = {'path': b_top['path'], 'status': b_top['status'], 'time': b_top['datetime']}

            return layup_dt.strftime("%Y-%m-%d %H:%M:%S"), station_found, line_str, shift_str, pre_el_views
        else:
            # Fallback to MES data if images were not found on disk share
            if mes_dt:
                station_found = mes_station or (mes_data.get('tumlayup') if mes_data else '')
                line_str, shift_str = get_responsible_shift_and_line(mes_dt, station_found)
                print(f"[MES SMART FALLBACK]: Populating layup time {mes_dt} & station '{station_found}' (Line: {line_str}, Shift: {shift_str}) from MES.")
                return mes_dt.strftime("%Y-%m-%d %H:%M:%S"), station_found, line_str, shift_str, pre_el_views
            return "", "", "", "", pre_el_views

    def query_pre_el_async(self, target_rec, sn, file_dt):
        def _fetch():
            try:
                lt, st, ln, sh, views = self.query_pre_el(sn, file_dt)
                target_rec['layup_time'] = lt
                target_rec['station'] = st
                if ln:
                    target_rec['line'] = ln
                elif not target_rec.get('line'):
                    target_rec['line'] = ""
                if sh:
                    target_rec['shift'] = sh
                elif not target_rec.get('shift'):
                    target_rec['shift'] = ""
                target_rec['pre_el_views'] = views

                def _update_ui():
                    try:
                        self.refresh_treeview()
                        rec_id = target_rec.get('id')
                        is_active = (self.selected_rec_id == rec_id) or (not self.selected_rec_id and self.records and self.records[0].get('id') == rec_id)
                        if is_active and ln and hasattr(self, 'cb_line'):
                            self.cb_line.set(ln)
                    except Exception as ex:
                        print(f"[PRE-EL UI UPDATE ERROR]: {ex}")

                self.after(0, _update_ui)
                self.app.incoming_event_queue.put(('REFRESH', {'rec_id': target_rec.get('id'), 'line': ln, 'shift': sh}))
            except Exception as e:
                print(f"[PRE-EL QUERY ERROR]: {e}")

        threading.Thread(target=_fetch, daemon=True).start()

    def handle_incoming_phone_upload(self, file_path, extracted_sn=None, action="NEW_RECORD", v_summary="", v_class="", v_result="", raw_voice="", audio_path="", file_dt=None):
        if file_dt is None and file_path:
            file_dt = get_file_datetime(file_path)
        if file_dt is None:
            file_dt = datetime.now()

        if raw_voice:
            self.last_spoken_raw = raw_voice
            self.last_audio_path = audio_path
            self.lbl_hearing.config(text=f"App Heard: '{raw_voice}'\n-> Defect: {v_summary or 'Unrecognized'} | Result: {v_result or 'Unspecified'}", fg=config.COLOR_PRIMARY)
            if v_summary:
                self.btn_learn_voice.pack(fill=tk.X, pady=4)
            try:
                update_live_hud_state(live_voice={'raw_text': raw_voice, 'defect': v_summary or '-', 'result': v_result or '-'})
            except Exception:
                pass

        if action == "VOICE_DETECTED":
            target_rec = self.find_or_create_panel_session("VOICE", file_dt)
            if v_summary:
                target_rec['summary'] = v_summary
                target_rec['class'] = v_class
            if v_result:
                target_rec['result'] = v_result
            target_rec['audio_path'] = audio_path
            target_rec['raw_voice'] = raw_voice
            self.refresh_treeview()
            return

        if action == "DEFECT_PHOTO":
            target_rec = self.find_or_create_panel_session("DEFECT_PHOTO", file_dt, sn=extracted_sn, file_path=file_path)
            if (not target_rec.get('sn') or target_rec.get('sn') in ("Pending SN", "-", "")) and extracted_sn:
                target_rec['sn'] = extracted_sn
                target_rec['order'] = extracted_sn[:9] if len(extracted_sn) >= 9 else extracted_sn
            if target_rec.get('sn') and target_rec['sn'] not in ("Pending SN", "-", ""):
                renamed_def = rename_to_sn_pattern(file_path, "Def", target_rec['sn'])
            else:
                renamed_def = file_path
            target_rec['photo_path'] = renamed_def
            self.register_photo_alias(target_rec['id'], file_path, renamed_def)
            self.refresh_treeview()
            return

        if action == "SN_PHOTO" or extracted_sn:
            raw_sn = extracted_sn or ""
            sn = clean_and_validate_sn(raw_sn)
            if not sn:
                print(f"[REJECTED PHONE SN]: Discarded non-V01 barcode '{raw_sn}'")
                return

            renamed_sn_path = rename_to_sn_pattern(file_path, "SN", sn)
            target_rec = self.find_or_create_panel_session("SN_PHOTO", file_dt, sn=sn, file_path=file_path)
            eval_shift = get_evaluation_shift(file_dt)

            target_rec['sn'] = sn
            target_rec['order'] = sn[:9] if len(sn) >= 9 else sn
            target_rec['sn_photo_path'] = renamed_sn_path
            self.register_photo_alias(target_rec['id'], file_path, renamed_sn_path)
            if target_rec.get('photo_path'):
                old_p = target_rec['photo_path']
                renamed_def = rename_to_sn_pattern(old_p, "Def", sn)
                target_rec['photo_path'] = renamed_def
                self.register_photo_alias(target_rec['id'], old_p, renamed_def)
            target_rec['date'] = file_dt.strftime("%m/%d/%Y")
            target_rec['eval_shift'] = eval_shift
            self.selected_rec_id = target_rec.get('id')
            self.refresh_treeview()

            # Asynchronously query Pre-EL machine station to prevent GUI lag
            self.query_pre_el_async(target_rec, sn, file_dt)
            return

    def on_barcode_scanned(self):
        raw_sn = self.ent_sn.get().strip()
        if not raw_sn: return

        sn = clean_and_validate_sn(raw_sn)
        if not sn:
            if "NEG" in raw_sn.upper() or "TSM" in raw_sn.upper() or "|" in raw_sn:
                messagebox.showwarning(
                    "Model Barcode Scanned",
                    f"⚠️ You scanned the Model/Specification Barcode:\n'{raw_sn}'\n\n"
                    "Please scan the Module Serial Number Barcode (begins with 'V01') located directly below it."
                )
            else:
                messagebox.showwarning(
                    "Invalid Serial Number",
                    f"⚠️ '{raw_sn}' is not a valid Solar Module Serial Number.\n\n"
                    "Module serial numbers must start with 'V01' (e.g. V01269003050237)."
                )
            self.ent_sn.delete(0, tk.END)
            self.ent_sn.focus_set()
            return

        now_dt = datetime.now()

        target_rec = self.find_or_create_panel_session("SN_PHOTO", now_dt, sn=sn)
        eval_shift = get_evaluation_shift(now_dt)

        target_rec['sn'] = sn
        target_rec['order'] = sn[:9] if len(sn) >= 9 else sn
        if target_rec.get('photo_path'):
            target_rec['photo_path'] = rename_to_sn_pattern(target_rec['photo_path'], "Def", sn)
        if target_rec.get('sn_photo_path'):
            target_rec['sn_photo_path'] = rename_to_sn_pattern(target_rec['sn_photo_path'], "SN", sn)
        if self.cb_summary.get(): target_rec['summary'] = self.cb_summary.get()
        if self.cb_class.get(): target_rec['class'] = self.cb_class.get()
        if self.cb_result.get(): target_rec['result'] = self.cb_result.get()
        if self.cb_cause.get(): target_rec['cause'] = self.cb_cause.get()
        self.selected_rec_id = target_rec.get('id')
        self.refresh_treeview()

        self.query_pre_el_async(target_rec, sn, now_dt)

        self.ent_sn.delete(0, tk.END)
        self.active_photo_path = None
        self.cb_summary.set("")
        self.cb_class.set("")
        self.cb_result.set("")
        self.cb_cause.set("")
        self.ent_sn.focus_set()

    def clear_all_records(self):
        if messagebox.askyesno("Clear", "Clear all records?"):
            self.records.clear()
            self.app.global_review_records = []
            for w in self.rows_frame.winfo_children(): w.destroy()
            self.lbl_table_count.config(text="Total Records Reviewed: 0")

    def manual_shift_reset(self):
        if messagebox.askyesno("New Shift Reset", "Start a new operational shift?\n\nThis will safely archive current review records to history and reset the workspace with a clean table."):
            shift_info = get_operational_shift_info(datetime.now())
            self.reset_for_new_shift(shift_info)
            messagebox.showinfo("New Shift Reset", f"Workspace reset for {shift_info.get('name', 'New Shift')}.\nPrevious records archived safely.")

    def open_archive_history(self):
        archive_file = os.path.join(config.LOCAL_DATA_DIR, "module_review_history_archive.json")
        if not os.path.exists(archive_file):
            messagebox.showinfo("Archive Empty", "No archived Module Review records found yet.")
            return
        try:
            os.startfile(archive_file)
        except Exception:
            try:
                import subprocess
                subprocess.Popen(["notepad.exe", archive_file])
            except Exception as e:
                messagebox.showerror("Error", f"Could not open archive file: {e}")

    def reset_for_new_shift(self, shift_info=None):
        if self.records:
            try:
                archive_file = os.path.join(config.LOCAL_DATA_DIR, "module_review_history_archive.json")
                existing_archive = []
                if os.path.exists(archive_file):
                    try:
                        with open(archive_file, 'r', encoding='utf-8') as f:
                            existing_archive = json.load(f)
                    except Exception:
                        existing_archive = []
                clean_recs = [{k: v for k, v in r.items() if not str(k).startswith('_')} for r in self.records]
                existing_archive.extend(clean_recs)
                with open(archive_file, 'w', encoding='utf-8') as f:
                    json.dump(existing_archive, f, ensure_ascii=False, indent=2, default=str)
                print(f"[MR SHIFT RESET]: Archived {len(clean_recs)} records to {archive_file}")
            except Exception as e:
                print(f"[MR SHIFT ARCHIVE ERROR]: {e}")

        if hasattr(self, 'ent_sn'):
            self.ent_sn.delete(0, tk.END)
        if hasattr(self, 'cb_class'):
            self.cb_class.set("")
        if hasattr(self, 'cb_summary'):
            self.cb_summary.set("")
        if hasattr(self, 'cb_result'):
            self.cb_result.set("")
        if hasattr(self, 'cb_line'):
            self.cb_line.set("")
        if hasattr(self, 'cb_cause'):
            self.cb_cause.set("")
        self.active_photo_path = None
        self.selected_rec_id = None
        self.records.clear()
        self.app.global_review_records = []
        if hasattr(self, 'rows_frame'):
            for w in self.rows_frame.winfo_children():
                w.destroy()
        if hasattr(self, 'lbl_table_count'):
            self.lbl_table_count.config(text="Total Records Reviewed: 0")
        shift_label = shift_info.get('name', 'New Shift') if shift_info else 'New Shift'
        if hasattr(self, 'lbl_hearing'):
            self.lbl_hearing.config(text=f"New Shift ({shift_label}): Workspace refreshed.", fg=config.COLOR_STATUS_OK)
        if hasattr(self, 'lbl_photo_status'):
            self.lbl_photo_status.config(text=f"Shift Rollover: {shift_label}. Ready.")
        try:
            save_module_review_history(self.records)
        except Exception:
            pass

    def export_styled_excel(self):
        if not self.records:
            messagebox.showwarning("Export", "No records!")
            return
        filename = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel Files", "*.xlsx")], title="Export")
        if not filename: return
        try:
            import openpyxl
            from openpyxl.drawing.image import Image as XLImage
            
            wb = openpyxl.Workbook()
            ws = wb.active
            headers = [
                "Date", "Order No", "Serial No", "Defect Classification", "Defect Description", 
                "Result", "Defect Cause", "Defect Location", "Root Cause", 
                "Pre-Layup Photo", "Post-Layup Photo", "Layup Time", "Station", 
                "Eval Shift", "Line", "Responsible Shift"
            ]
            ws.append(headers)
            
            def _clean(v):
                if v is None: return ""
                s = str(v).strip()
                return "" if s in ("-", "None", "Pending SN") else s

            for r_idx, rec in enumerate(self.records, start=2):
                def_summary = _clean(rec.get('summary'))
                def_class = _clean(rec.get('class'))
                if not def_class and def_summary and hasattr(config, 'get_class_for_summary'):
                    def_class = config.get_class_for_summary(def_summary)

                ws.cell(row=r_idx, column=1, value=_clean(rec.get('date')))
                ws.cell(row=r_idx, column=2, value=_clean(rec.get('order')))
                ws.cell(row=r_idx, column=3, value=_clean(rec.get('sn')))
                ws.cell(row=r_idx, column=4, value=def_class)   # Column D: Defect Classification (不良归类)
                ws.cell(row=r_idx, column=5, value=def_summary) # Column E: Defect Description (不良描述)
                ws.cell(row=r_idx, column=6, value=_clean(rec.get('result')))
                ws.cell(row=r_idx, column=7, value=_clean(rec.get('cause')))
                ws.cell(row=r_idx, column=8, value=_clean(rec.get('location')))
                ws.cell(row=r_idx, column=9, value=_clean(rec.get('pass_cause', '')))
                ws.cell(row=r_idx, column=12, value=_clean(rec.get('layup_time')))
                ws.cell(row=r_idx, column=13, value=_clean(rec.get('station')))
                ws.cell(row=r_idx, column=14, value=get_evaluation_shift(rec.get('eval_shift') or rec.get('dt')))
                ws.cell(row=r_idx, column=15, value=_clean(rec.get('line')))
                ws.cell(row=r_idx, column=16, value=_clean(rec.get('shift')))
                ws.row_dimensions[r_idx].height = 40
                
                # 1. Embed Pre-Layup Photo into Column J (PreEL snip from Front/EL/Back)
                pre_photo = rec.get('pre_el_snip_path')
                if pre_photo and os.path.exists(pre_photo):
                    try:
                        from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
                        from openpyxl.utils.units import pixels_to_EMU
                        
                        img_pre = XLImage(pre_photo)
                        col_j = 9  # Column J (0-indexed)
                        row_num = r_idx - 1
                        _from_j = AnchorMarker(col=col_j, colOff=pixels_to_EMU(4), row=row_num, rowOff=pixels_to_EMU(4))
                        _to_j = AnchorMarker(col=col_j + 1, colOff=pixels_to_EMU(-4), row=row_num + 1, rowOff=pixels_to_EMU(-4))
                        img_pre.anchor = TwoCellAnchor(editAs='twoCell', _from=_from_j, to=_to_j)
                        ws.add_image(img_pre)
                    except Exception:
                        try:
                            img_pre = XLImage(pre_photo)
                            img_pre.width = 100
                            img_pre.height = 75
                            ws.add_image(img_pre, f"J{r_idx}")
                        except Exception: pass

                # 2. Embed Post-Layup Photo into Column K (MR Defect Photo / Snip)
                photo_path = rec.get('photo_path')
                if photo_path and os.path.exists(photo_path):
                    try:
                        from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
                        from openpyxl.utils.units import pixels_to_EMU
                        
                        img_post = XLImage(photo_path)
                        col_k = 10  # Column K (0-indexed)
                        row_num = r_idx - 1
                        _from_k = AnchorMarker(col=col_k, colOff=pixels_to_EMU(4), row=row_num, rowOff=pixels_to_EMU(4))
                        _to_k = AnchorMarker(col=col_k + 1, colOff=pixels_to_EMU(-4), row=row_num + 1, rowOff=pixels_to_EMU(-4))
                        img_post.anchor = TwoCellAnchor(editAs='twoCell', _from=_from_k, to=_to_k)
                        ws.add_image(img_post)
                    except Exception:
                        try:
                            img_post = XLImage(photo_path)
                            img_post.width = 100
                            img_post.height = 75
                            ws.add_image(img_post, f"K{r_idx}")
                        except Exception: pass

            ws.column_dimensions['J'].width = 18
            ws.column_dimensions['K'].width = 18

            wb.save(filename)
            messagebox.showinfo("Success", f"Exported styled Excel with images to {filename}")
        except Exception as e: 
            messagebox.showerror("Error", f"{e}")

    def sync_to_master_report(self):
        if not self.records:
            messagebox.showwarning("Sync Master Report", "No defect records available to sync!")
            return
        master_path = getattr(config, 'MASTER_REPORT_PATH', '').strip()
        if not master_path:
            master_path = filedialog.asksaveasfilename(
                title="Select or Create Master Daily QC Excel Report File",
                defaultextension=".xlsx",
                filetypes=[("Excel Workbooks", "*.xlsx"), ("All Files", "*.*")]
            )
            if not master_path:
                return
            config.MASTER_REPORT_PATH = master_path
            if hasattr(config, 'save_persistent_settings'):
                config.save_persistent_settings()

        self.lbl_photo_status.config(text="⚡ Syncing records to Master Report in background...", fg="#38bdf8")
        rec_snapshot = list(self.records)

        def _bg_sync():
            ok, msg = sync_to_master_excel(master_path, rec_snapshot)
            def _done():
                if ok:
                    self.lbl_photo_status.config(text="✅ Master Report Synced Successfully!", fg="#10b981")
                    messagebox.showinfo("Master Report Synced", msg)
                else:
                    self.lbl_photo_status.config(text="❌ Master Report Sync Failed!", fg="#ef4444")
                    messagebox.showerror("Master Sync Error", msg)
            self.root.after(0, _done)

        threading.Thread(target=_bg_sync, daemon=True).start()

    def on_mobile_override(self, data):
        if not self.records: return
        target_rec = self.records[0]
        if 'line' in data and data['line']:
            target_rec['line'] = data['line']
        if 'result' in data and data['result']:
            target_rec['result'] = data['result']
        if 'summary' in data and data['summary']:
            target_rec['summary'] = data['summary']
        self.root.after(0, self.refresh_treeview)

QCDefectReviewTab = ModuleReviewTab


class ConfigPanel(tk.Frame):
    def __init__(self, parent, app):
        super().__init__(parent, bg=config.COLOR_BG_MAIN, padx=25, pady=20)
        self.app = app
        self.setup_ui()

    def setup_ui(self):
        canvas = tk.Canvas(self, bg=config.COLOR_BG_MAIN, highlightthickness=0)
        scrollbar = ttk.Scrollbar(self, orient="vertical", command=canvas.yview)
        scroll_content = tk.Frame(canvas, bg=config.COLOR_BG_MAIN)
        scroll_content.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=scroll_content, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        tk.Label(scroll_content, text="Server, Station & Defect Tree Configuration", font=config.FONT_TITLE, bg=config.COLOR_BG_MAIN, fg=config.COLOR_TEXT_MAIN).pack(anchor=tk.W, pady=(0, 15))

        card_defect = tk.LabelFrame(scroll_content, text=" 1. Defect Category Manager ", bg=config.COLOR_BG_MAIN, font=config.FONT_BODY_BOLD, fg=config.COLOR_PRIMARY, padx=12, pady=10)
        card_defect.pack(fill=tk.X, pady=(0, 12))
        action_bar = tk.Frame(card_defect, bg=config.COLOR_BG_MAIN)
        action_bar.pack(fill=tk.X, pady=(0, 8))
        ModernButton(action_bar, text="Upload Defect Excel File (.xlsx)", command=self.import_defect_excel, primary=True, padx=8, pady=3).pack(side=tk.LEFT, padx=(0, 10))
        ModernButton(action_bar, text="Save Defect Tree as Default", command=self.save_defect_tree_directly, primary=False, padx=8, pady=3).pack(side=tk.LEFT)

        row_c = tk.Frame(card_defect, bg=config.COLOR_BG_MAIN)
        row_c.pack(fill=tk.X, pady=4)
        tk.Label(row_c, text="Select Classification:", font=config.FONT_BODY_BOLD, bg=config.COLOR_BG_MAIN, width=22, anchor="w").pack(side=tk.LEFT)
        self.cfg_cb_class = ttk.Combobox(row_c, values=list(config.DEFECT_TREE.keys()), width=32, state="readonly")
        if config.DEFECT_TREE: self.cfg_cb_class.set(list(config.DEFECT_TREE.keys())[0])
        self.cfg_cb_class.pack(side=tk.LEFT, padx=(0, 10))
        self.cfg_cb_class.bind("<<ComboboxSelected>>", lambda e: self.on_cfg_class_selected())
        self.cfg_ent_class_rename = tk.Entry(row_c, font=config.FONT_BODY, width=26)
        self.cfg_ent_class_rename.pack(side=tk.LEFT, padx=(0, 6))
        ModernButton(row_c, text="Rename Class", command=self.rename_current_classification, primary=False, padx=6, pady=2).pack(side=tk.LEFT)

        row_s = tk.Frame(card_defect, bg=config.COLOR_BG_MAIN)
        row_s.pack(fill=tk.X, pady=4)
        tk.Label(row_s, text="Select Defect Summary:", font=config.FONT_BODY_BOLD, bg=config.COLOR_BG_MAIN, width=22, anchor="w").pack(side=tk.LEFT)
        self.cfg_cb_summary = ttk.Combobox(row_s, values=[], width=32, state="readonly")
        self.cfg_cb_summary.pack(side=tk.LEFT, padx=(0, 10))
        self.cfg_ent_summary_rename = tk.Entry(row_s, font=config.FONT_BODY, width=26)
        self.cfg_ent_summary_rename.pack(side=tk.LEFT, padx=(0, 6))
        ModernButton(row_s, text="Rename Summary", command=self.rename_current_summary, primary=False, padx=6, pady=2).pack(side=tk.LEFT)

        row_add = tk.Frame(card_defect, bg=config.COLOR_BG_MAIN)
        row_add.pack(fill=tk.X, pady=(6, 2))
        tk.Label(row_add, text="Add New Item:", font=config.FONT_BODY_BOLD, bg=config.COLOR_BG_MAIN, width=22, anchor="w").pack(side=tk.LEFT)
        self.cfg_ent_new_item = tk.Entry(row_add, font=config.FONT_BODY, width=32)
        self.cfg_ent_new_item.pack(side=tk.LEFT, padx=(0, 10))
        ModernButton(row_add, text="+ Add as New Summary", command=self.add_new_summary_to_class, primary=False, padx=6, pady=2).pack(side=tk.LEFT, padx=(0, 4))
        ModernButton(row_add, text="+ Add as New Class", command=self.add_new_classification, primary=False, padx=6, pady=2).pack(side=tk.LEFT)
        self.on_cfg_class_selected()

        card_lines = tk.LabelFrame(scroll_content, text=" 2. Active Production Lines (Default Selection) ", bg=config.COLOR_BG_MAIN, font=config.FONT_BODY_BOLD, fg=config.COLOR_PRIMARY, padx=12, pady=10)
        card_lines.pack(fill=tk.X, pady=(0, 12))
        self.line_buttons = {}
        active_lines_dict = getattr(config, 'ACTIVE_LINES', {})
        for i in range(config.FINAL_EL_STATION_MIN, config.FINAL_EL_STATION_MAX + 1):
            st = f"{config.FINAL_EL_STATION_PREFIX}{i}"
            is_active = active_lines_dict.get(st, True)
            line_idx = i - 1000
            btn = ModernToggleButton(card_lines, text=f"Line {line_idx}", initial_state=is_active, padx=10, pady=3)
            btn.pack(side=tk.LEFT, padx=6)
            self.line_buttons[st] = btn

        card_time = tk.LabelFrame(scroll_content, text=" 3. Search Preferences & Display Zoom ", bg=config.COLOR_BG_MAIN, font=config.FONT_BODY_BOLD, fg=config.COLOR_PRIMARY, padx=12, pady=10)
        card_time.pack(fill=tk.X, pady=(0, 12))
        tk.Label(card_time, text="Pre/Final EL Days Back:", font=config.FONT_BODY, bg=config.COLOR_BG_MAIN).grid(row=0, column=0, sticky="w", pady=4)
        self.ent_pre_days = tk.Entry(card_time, width=6, font=config.FONT_BODY, justify=tk.CENTER)
        self.ent_pre_days.insert(0, str(getattr(config, 'DEFAULT_PRE_FINAL_DAYS_BACK_START', 4)))
        self.ent_pre_days.grid(row=0, column=1, padx=6, pady=4)
        tk.Label(card_time, text="String Black Interval:", font=config.FONT_BODY, bg=config.COLOR_BG_MAIN).grid(row=0, column=2, sticky="w", padx=(15, 4), pady=4)
        self.cb_def_interval = ttk.Combobox(card_time, values=getattr(config, 'TIME_INTERVAL_OPTIONS', ["1 Hour", "2 Hours", "4 Hours", "All Shift"]), state="readonly", width=12)
        self.cb_def_interval.set(getattr(config, 'DEFAULT_TIME_INTERVAL', '2 Hours'))
        self.cb_def_interval.grid(row=0, column=3, padx=6, pady=4)
        self.cb_def_interval.bind("<<ComboboxSelected>>", lambda e: self.on_interval_selected())
        tk.Label(card_time, text="Display Zoom:", font=config.FONT_BODY_BOLD, bg=config.COLOR_BG_MAIN, fg=config.COLOR_PRIMARY).grid(row=0, column=4, sticky="w", padx=(15, 4), pady=4)
        zoom_options = getattr(config, 'CHART_ZOOM_OPTIONS', ["70%", "80%", "90%", "100%", "110%", "120%"])
        self.cb_zoom = ttk.Combobox(card_time, values=zoom_options, state="readonly", width=8)
        self.cb_zoom.set(getattr(config, 'CHART_ZOOM_LEVEL', '90%'))
        self.cb_zoom.grid(row=0, column=5, padx=6, pady=4)
        self.cb_zoom.bind("<<ComboboxSelected>>", lambda e: self.on_zoom_selected())

        card_pre = tk.LabelFrame(scroll_content, text=" 4. Pre EL Root Server ", bg=config.COLOR_BG_MAIN, font=config.FONT_BODY_BOLD, fg=config.COLOR_PRIMARY, padx=12, pady=10)
        card_pre.pack(fill=tk.X, pady=(0, 12))
        tk.Label(card_pre, text="Path:", font=config.FONT_BODY, bg=config.COLOR_BG_MAIN).grid(row=0, column=0, sticky="w", pady=4)
        self.ent_pre_root = tk.Entry(card_pre, width=65, font=config.FONT_BODY)
        self.ent_pre_root.insert(0, config.PRE_EL_NETWORK_ROOT)
        self.ent_pre_root.grid(row=0, column=1, padx=8, pady=4)
        self.lbl_pre_status = tk.Label(card_pre, text="Checking...", font=config.FONT_BODY_BOLD, bg=config.COLOR_BG_MAIN, fg=config.COLOR_STATUS_WARNING)
        self.lbl_pre_status.grid(row=0, column=2, padx=8)

        card_fin = tk.LabelFrame(scroll_content, text=" 5. Final EL Station IP Mappings (Line 1 to 7) ", bg=config.COLOR_BG_MAIN, font=config.FONT_BODY_BOLD, fg=config.COLOR_PRIMARY, padx=12, pady=10)
        card_fin.pack(fill=tk.X, pady=(0, 15))
        self.fin_entries = {}
        self.fin_labels = {}
        for row_idx, (st_name, st_path) in enumerate(config.FINAL_EL_MAP.items()):
            tk.Label(card_fin, text=f"{st_name}:", font=config.FONT_BODY_BOLD, bg=config.COLOR_BG_MAIN).grid(row=row_idx, column=0, sticky="w", pady=2)
            ent = tk.Entry(card_fin, width=65, font=config.FONT_BODY)
            ent.insert(0, st_path)
            ent.grid(row=row_idx, column=1, padx=8, pady=2)
            lbl = tk.Label(card_fin, text="Checking...", font=config.FONT_SMALL, bg=config.COLOR_BG_MAIN, fg=config.COLOR_STATUS_WARNING)
            lbl.grid(row=row_idx, column=2, padx=8)
            self.fin_entries[st_name] = ent
            self.fin_labels[st_name] = lbl

        card_phone = tk.LabelFrame(scroll_content, text=" 6. Phone Sync & Mac Mini Relay Setup ", bg=config.COLOR_BG_MAIN, font=config.FONT_BODY_BOLD, fg=config.COLOR_PRIMARY, padx=12, pady=10)
        card_phone.pack(fill=tk.X, pady=(0, 15))

        r_relay = tk.Frame(card_phone, bg=config.COLOR_BG_MAIN)
        r_relay.pack(fill=tk.X, pady=3)
        tk.Label(r_relay, text="Mac Mini Relay URL:", font=config.FONT_BODY_BOLD, bg=config.COLOR_BG_MAIN, width=20, anchor="w").pack(side=tk.LEFT)
        self.ent_relay_url = tk.Entry(r_relay, width=42, font=config.FONT_BODY)
        self.ent_relay_url.insert(0, getattr(config, 'MAC_MINI_RELAY_URL', ''))
        self.ent_relay_url.pack(side=tk.LEFT, padx=(0, 8))
        self.btn_toggle_relay = ModernToggleButton(r_relay, text="Enable Relay", initial_state=getattr(config, 'MAC_MINI_RELAY_ENABLED', False), padx=8, pady=2)
        self.btn_toggle_relay.pack(side=tk.LEFT, padx=(0, 8))
        ModernButton(r_relay, text="Test Relay", command=self.test_relay_connection, primary=False, padx=6, pady=2).pack(side=tk.LEFT, padx=(0, 8))
        self.lbl_relay_status = tk.Label(r_relay, text="Ready", font=config.FONT_SMALL, bg=config.COLOR_BG_MAIN, fg=config.COLOR_TEXT_SECONDARY)
        self.lbl_relay_status.pack(side=tk.LEFT)

        r_cam = tk.Frame(card_phone, bg=config.COLOR_BG_MAIN)
        r_cam.pack(fill=tk.X, pady=3)
        tk.Label(r_cam, text="Custom Camera Dir (Optional):", font=config.FONT_SMALL, bg=config.COLOR_BG_MAIN, width=25, anchor="w").pack(side=tk.LEFT)
        self.ent_phone_cam = tk.Entry(r_cam, width=55, font=config.FONT_SMALL)
        self.ent_phone_cam.insert(0, getattr(config, 'PHONE_CAMERA_DIR', ''))
        self.ent_phone_cam.pack(side=tk.LEFT, padx=(0, 10))

        r_rec = tk.Frame(card_phone, bg=config.COLOR_BG_MAIN)
        r_rec.pack(fill=tk.X, pady=3)
        tk.Label(r_rec, text="Custom Audio Dir (Optional):", font=config.FONT_SMALL, bg=config.COLOR_BG_MAIN, width=25, anchor="w").pack(side=tk.LEFT)
        self.ent_phone_rec = tk.Entry(r_rec, width=55, font=config.FONT_SMALL)
        self.ent_phone_rec.insert(0, getattr(config, 'PHONE_RECORDINGS_DIR', ''))
        self.ent_phone_rec.pack(side=tk.LEFT, padx=(0, 10))

        card_master = tk.LabelFrame(scroll_content, text=" 7. Master Daily Report & Mobile HUD Settings ", bg=config.COLOR_BG_MAIN, font=config.FONT_BODY_BOLD, fg=config.COLOR_PRIMARY, padx=12, pady=10)
        card_master.pack(fill=tk.X, pady=(0, 15))

        r_mpath = tk.Frame(card_master, bg=config.COLOR_BG_MAIN)
        r_mpath.pack(fill=tk.X, pady=3)
        tk.Label(r_mpath, text="Master Report Excel File / Share Path:", font=config.FONT_BODY_BOLD, bg=config.COLOR_BG_MAIN, width=32, anchor="w").pack(side=tk.LEFT)
        self.ent_master_path = tk.Entry(r_mpath, width=48, font=config.FONT_BODY)
        self.ent_master_path.insert(0, getattr(config, 'MASTER_REPORT_PATH', ''))
        self.ent_master_path.pack(side=tk.LEFT, padx=(0, 8))
        ModernButton(r_mpath, text="Browse...", command=self.browse_master_path, primary=False, padx=6, pady=2).pack(side=tk.LEFT)

        r_mport = tk.Frame(card_master, bg=config.COLOR_BG_MAIN)
        r_mport.pack(fill=tk.X, pady=3)
        tk.Label(r_mport, text="Mobile HUD Port:", font=config.FONT_BODY_BOLD, bg=config.COLOR_BG_MAIN, width=32, anchor="w").pack(side=tk.LEFT)
        self.ent_hud_port = tk.Entry(r_mport, width=8, font=config.FONT_BODY, justify=tk.CENTER)
        self.ent_hud_port.insert(0, str(getattr(config, 'MOBILE_HUD_PORT', 8080)))
        self.ent_hud_port.pack(side=tk.LEFT, padx=(0, 10))
        tk.Label(r_mport, text="(Access live dashboard from any phone browser on Wi-Fi)", font=config.FONT_SMALL, bg=config.COLOR_BG_MAIN, fg=config.COLOR_TEXT_SECONDARY).pack(side=tk.LEFT)

        btn_box = tk.Frame(scroll_content, bg=config.COLOR_BG_MAIN)
        btn_box.pack(anchor=tk.W, pady=(0, 20))
        ModernButton(btn_box, text="Test All Connections", command=self.test_connections, primary=False).pack(side=tk.LEFT, padx=(0, 10))
        ModernButton(btn_box, text="Save Settings (Save as Default)", command=self.save_settings, primary=True).pack(side=tk.LEFT)
        self.test_connections()



    def on_cfg_class_selected(self):
        curr_class = self.cfg_cb_class.get()
        self.cfg_ent_class_rename.delete(0, tk.END)
        self.cfg_ent_class_rename.insert(0, curr_class)
        summaries = config.DEFECT_TREE.get(curr_class, [])
        self.cfg_cb_summary['values'] = summaries
        if summaries:
            self.cfg_cb_summary.set(summaries[0])
            self.cfg_ent_summary_rename.delete(0, tk.END)
            self.cfg_ent_summary_rename.insert(0, summaries[0])
        else:
            self.cfg_cb_summary.set("")
            self.cfg_ent_summary_rename.delete(0, tk.END)

    def rename_current_classification(self):
        old_name = self.cfg_cb_class.get()
        new_name = self.cfg_ent_class_rename.get().strip()
        if not old_name or not new_name or old_name == new_name: return
        items = config.DEFECT_TREE.pop(old_name, [])
        config.DEFECT_TREE[new_name] = items
        self.save_defect_tree_directly()
        self.cfg_cb_class['values'] = list(config.DEFECT_TREE.keys())
        self.cfg_cb_class.set(new_name)
        messagebox.showinfo("Renamed", f"Renamed classification:\n'{old_name}' -> '{new_name}'")

    def rename_current_summary(self):
        curr_class = self.cfg_cb_class.get()
        old_sum = self.cfg_cb_summary.get()
        new_sum = self.cfg_ent_summary_rename.get().strip()
        if not curr_class or not old_sum or not new_sum or old_sum == new_sum: return
        summaries = config.DEFECT_TREE.get(curr_class, [])
        if old_sum in summaries:
            idx = summaries.index(old_sum)
            summaries[idx] = new_sum
            config.DEFECT_TREE[curr_class] = summaries
            self.save_defect_tree_directly()
            self.cfg_cb_summary['values'] = summaries
            self.cfg_cb_summary.set(new_sum)
            messagebox.showinfo("Renamed", f"Renamed summary:\n'{old_sum}' -> '{new_sum}'")

    def add_new_summary_to_class(self):
        curr_class = self.cfg_cb_class.get()
        new_item = self.cfg_ent_new_item.get().strip()
        if not curr_class or not new_item: return
        summaries = config.DEFECT_TREE.setdefault(curr_class, [])
        if new_item not in summaries:
            summaries.append(new_item)
            self.save_defect_tree_directly()
            self.cfg_cb_summary['values'] = summaries
            self.cfg_cb_summary.set(new_item)
            self.cfg_ent_new_item.delete(0, tk.END)
            messagebox.showinfo("Added", f"Added '{new_item}' to '{curr_class}'!")

    def add_new_classification(self):
        new_class = self.cfg_ent_new_item.get().strip()
        if not new_class: return
        if new_class not in config.DEFECT_TREE:
            config.DEFECT_TREE[new_class] = []
            self.save_defect_tree_directly()
            self.cfg_cb_class['values'] = list(config.DEFECT_TREE.keys())
            self.cfg_cb_class.set(new_class)
            self.on_cfg_class_selected()
            self.cfg_ent_new_item.delete(0, tk.END)
            messagebox.showinfo("Added", f"Created new classification: '{new_class}'!")

    def import_defect_excel(self):
        file_path = filedialog.askopenfilename(title="Select Defect Table Excel", filetypes=[("Excel Files", "*.xlsx *.xls")])
        if not file_path: return
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            ws = wb.active
            new_tree = {}
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or len(row) < 2: continue
                desc = str(row[0] if row[0] is not None else "").strip()
                cls_val = str(row[1] if row[1] is not None else "").strip()
                if len(row) >= 3 and (not desc or desc.lower() == "none"):
                    desc = str(row[1] if row[1] is not None else "").strip()
                    cls_val = str(row[2] if row[2] is not None else "").strip()
                if not cls_val or not desc or "defect classification" in cls_val.lower(): continue
                if cls_val not in new_tree: new_tree[cls_val] = []
                if desc not in new_tree[cls_val]: new_tree[cls_val].append(desc)
            if new_tree:
                config.DEFECT_TREE = new_tree
                self.save_defect_tree_directly()
                self.cfg_cb_class['values'] = list(config.DEFECT_TREE.keys())
                self.cfg_cb_class.set(list(config.DEFECT_TREE.keys())[0])
                self.on_cfg_class_selected()
                self.app.sync_all_tabs(self.cb_def_interval.get())
                messagebox.showinfo("Import Success", f"Imported classifications from {os.path.basename(file_path)}")
        except Exception as e: messagebox.showerror("Excel Error", f"Failed: {e}")

    def save_defect_tree_directly(self):
        if hasattr(config, 'save_persistent_settings'): config.save_persistent_settings()
        self.app.sync_all_tabs(self.cb_def_interval.get())

    def on_interval_selected(self):
        new_interval = self.cb_def_interval.get()
        config.DEFAULT_TIME_INTERVAL = new_interval
        self.app.sync_all_tabs(new_interval)

    def on_zoom_selected(self):
        config.CHART_ZOOM_LEVEL = self.cb_zoom.get()
        self.app.sync_all_tabs(self.cb_def_interval.get())

    def test_relay_connection(self):
        url = self.ent_relay_url.get().strip()
        if not url:
            self.lbl_relay_status.config(text="URL Empty", fg=config.COLOR_STATUS_WARNING)
            return
        self.lbl_relay_status.config(text="Testing...", fg=config.COLOR_PRIMARY)
        def _test():
            ok, msg = test_mac_mini_relay_connection(url, getattr(config, 'MAC_MINI_RELAY_SECRET', ''))
            def update_ui():
                self.lbl_relay_status.config(text="ONLINE" if ok else "OFFLINE", fg=config.COLOR_STATUS_OK if ok else config.COLOR_STATUS_NG)
                if ok:
                    messagebox.showinfo("Mac Mini Relay", f"Relay is online!\n{msg}")
                else:
                    messagebox.showwarning("Mac Mini Relay", f"Could not reach relay:\n{msg}")
            self.app.root.after(0, update_ui)
        threading.Thread(target=_test, daemon=True).start()

    def test_connections(self):
        pre_path = self.ent_pre_root.get().strip()
        fin_paths = {st_name: ent.get().strip() for st_name, ent in self.fin_entries.items()}
        def _test():
            pre_ok, _ = self.app.search_engine.check_network_access(pre_path)
            fin_results = {st_name: self.app.search_engine.check_network_access(p)[0] for st_name, p in fin_paths.items()}
            def update_ui():
                self.lbl_pre_status.config(text="ONLINE" if pre_ok else "OFFLINE", fg=config.COLOR_STATUS_OK if pre_ok else config.COLOR_STATUS_NG)
                for st_name, ok in fin_results.items():
                    self.fin_labels[st_name].config(text="ONLINE" if ok else "OFFLINE", fg=config.COLOR_STATUS_OK if ok else config.COLOR_STATUS_NG)
            self.app.root.after(0, update_ui)
        threading.Thread(target=_test, daemon=True).start()


    def browse_master_path(self):
        f = filedialog.asksaveasfilename(
            title="Select or Create Master Daily QC Excel Report File",
            defaultextension=".xlsx",
            filetypes=[("Excel Workbooks", "*.xlsx"), ("All Files", "*.*")]
        )
        if f:
            self.ent_master_path.delete(0, tk.END)
            self.ent_master_path.insert(0, f)

    def save_settings(self):
        new_interval = self.cb_def_interval.get()
        config.DEFAULT_TIME_INTERVAL = new_interval
        config.CHART_ZOOM_LEVEL = self.cb_zoom.get()
        try: config.DEFAULT_PRE_FINAL_DAYS_BACK_START = int(self.ent_pre_days.get().strip())
        except ValueError: config.DEFAULT_PRE_FINAL_DAYS_BACK_START = 4
        config.PRE_EL_NETWORK_ROOT = self.ent_pre_root.get().strip()
        for st_name, ent in self.fin_entries.items(): config.FINAL_EL_MAP[st_name] = ent.get().strip()
        for st_name, btn in self.line_buttons.items(): config.ACTIVE_LINES[st_name] = btn.get_state()
        config.MAC_MINI_RELAY_URL = self.ent_relay_url.get().strip()
        config.MAC_MINI_RELAY_ENABLED = self.btn_toggle_relay.get_state()
        config.PHONE_CAMERA_DIR = self.ent_phone_cam.get().strip()
        config.PHONE_RECORDINGS_DIR = self.ent_phone_rec.get().strip()
        config.MASTER_REPORT_PATH = self.ent_master_path.get().strip()
        try: config.MOBILE_HUD_PORT = int(self.ent_hud_port.get().strip())
        except ValueError: config.MOBILE_HUD_PORT = 8080
        if hasattr(config, 'save_persistent_settings'): config.save_persistent_settings()
        self.app.sync_all_tabs(new_interval)
        self.test_connections()
        messagebox.showinfo("Configuration", "Settings saved!")



# =================== STANDARD SEARCH TAB CONTENT ===================

class TabContent(tk.Frame):
    def __init__(self, parent, app, mode="PRE_EL"):
        super().__init__(parent, bg=config.COLOR_BG_MAIN)
        self.app = app
        self.mode = mode
        self.current_results = {}
        self.raw_ng_images = []
        self.sn_list_ordered = []
        self.buttons_map = {}
        self.setup_ui()

    def setup_ui(self):
        filters = tk.Frame(self, bg=config.COLOR_BG_MAIN)
        filters.pack(fill=tk.X, pady=(0, 6))
        header_row = tk.Frame(filters, bg=config.COLOR_BG_MAIN)
        header_row.pack(fill=tk.X, pady=(0, 4))

        if self.mode == "PRE_EL": title_text = "Serial Numbers (Pre EL - PreEL1 to PreEL14)"
        elif self.mode == "FINAL_EL": title_text = "Serial Numbers (Final EL - Line 1 to 7)"
        else: title_text = "String Black Analytics (Multi-Chart View)"

        tk.Label(header_row, text=title_text, font=config.FONT_TITLE, bg=config.COLOR_BG_MAIN, fg=config.COLOR_TEXT_MAIN).pack(side=tk.LEFT)

        if self.mode != "STRING_BLACK":
            self.txt_sn = scrolledtext.ScrolledText(filters, height=3, font=("Segoe UI", 11), relief=tk.FLAT, bd=5, bg=config.COLOR_INPUT_BG, fg=config.COLOR_TEXT_MAIN)
            self.txt_sn.pack(fill=tk.X, pady=(0, 6))

        row2 = tk.Frame(filters, bg=config.COLOR_BG_MAIN)
        row2.pack(fill=tk.X, pady=(0, 4))

        def create_date_picker(parent_widget, label, default_days_back=0, on_change=None):
            container = tk.Frame(parent_widget, bg=config.COLOR_BG_MAIN)
            container.pack(side=tk.LEFT, padx=(0, 10), anchor="n")
            tk.Label(container, text=label, font=config.FONT_SMALL, bg=config.COLOR_BG_MAIN, fg=config.COLOR_TEXT_SECONDARY).pack(anchor=tk.W)
            default_date = datetime.now() - timedelta(days=default_days_back)
            val_holder = {"date": default_date}
            cb = tk.Entry(container, width=11, font=config.FONT_BODY, bg=config.COLOR_INPUT_BG, fg=config.COLOR_TEXT_MAIN, relief=tk.FLAT, bd=5, justify=tk.CENTER, cursor="hand2")
            cb.insert(0, default_date.strftime("%Y-%m-%d"))
            cb.config(state="readonly", readonlybackground=config.COLOR_INPUT_BG)
            cb.pack(pady=(2, 0))
            def _on_pick(d):
                val_holder.update({'date': d})
                cb.config(state='normal')
                cb.delete(0, tk.END)
                cb.insert(0, d.strftime('%Y-%m-%d'))
                cb.config(state='readonly')
                if on_change: on_change()
            cb.bind("<Button-1>", lambda e: CalendarDialog(self.app.root, current_date=val_holder['date'], anchor_widget=cb, callback=_on_pick))
            return val_holder

        days_back_start = getattr(config, 'DEFAULT_PRE_FINAL_DAYS_BACK_START', 4) if self.mode != "STRING_BLACK" else 0
        self.date_from_val = create_date_picker(row2, "From Date", days_back_start, on_change=lambda: self.render_string_black_dashboard() if self.mode == "STRING_BLACK" else None)
        self.date_to_val = create_date_picker(row2, "To Date", 0, on_change=lambda: self.render_string_black_dashboard() if self.mode == "STRING_BLACK" else None)



        if self.mode == "STRING_BLACK":
            f_time = tk.Frame(row2, bg=config.COLOR_BG_MAIN)
            f_time.pack(side=tk.LEFT, padx=(0, 10), anchor="n")
            tk.Label(f_time, text="Filter Window", font=config.FONT_SMALL, bg=config.COLOR_BG_MAIN, fg=config.COLOR_TEXT_SECONDARY).pack(anchor=tk.W)
            slots = generate_shift_time_slots(getattr(config, 'DEFAULT_TIME_INTERVAL', '2 Hours'))
            self.cb_timeslot = ttk.Combobox(f_time, values=slots, width=20, state="readonly")
            self.cb_timeslot.set(get_current_shift_slot(slots))
            self.cb_timeslot.pack(pady=(2, 0))
            self.cb_timeslot.bind("<<ComboboxSelected>>", lambda e: self.render_string_black_dashboard())

            f_line = tk.Frame(row2, bg=config.COLOR_BG_MAIN)
            f_line.pack(side=tk.LEFT, padx=(0, 10), anchor="n")
            tk.Label(f_line, text="Trend Focus Line", font=config.FONT_SMALL, bg=config.COLOR_BG_MAIN, fg=config.COLOR_TEXT_SECONDARY).pack(anchor=tk.W)
            active_lines_dict = getattr(config, 'ACTIVE_LINES', {})
            active_line_names = [f"Line {int(st.replace(config.FINAL_EL_STATION_PREFIX,''))-1000}" for st, a in active_lines_dict.items() if a]
            self.cb_trend_line = ttk.Combobox(f_line, values=active_line_names, width=14, state="readonly")
            if active_line_names: self.cb_trend_line.set(active_line_names[0])
            self.cb_trend_line.pack(pady=(2, 0))
            self.cb_trend_line.bind("<<ComboboxSelected>>", lambda e: self.render_string_black_dashboard())

        summary_container = tk.Frame(row2, bg=config.COLOR_BG_MAIN)
        summary_container.pack(side=tk.LEFT, padx=10, anchor="n")
        tk.Label(summary_container, text=" ", font=config.FONT_SMALL, bg=config.COLOR_BG_MAIN).pack(anchor=tk.W)
        self.lbl_summary = tk.Label(summary_container, text="", font=config.FONT_BODY_BOLD, bg=config.COLOR_BG_MAIN, fg=config.COLOR_PRIMARY)
        self.lbl_summary.pack(side=tk.LEFT, pady=(2, 0))

        self.btn_export = ModernButton(row2, text="Export Excel", command=self.export_excel, primary=False)
        self.btn_export.pack(side=tk.RIGHT, anchor=tk.S, padx=(0, 6))
        self.btn_export.config(state="disabled")

        self.btn_export_imgs = ModernButton(row2, text="Export Images", command=self.export_images_to_folder, primary=False)
        self.btn_export_imgs.pack(side=tk.RIGHT, anchor=tk.S, padx=(0, 10))
        self.btn_export_imgs.config(state="disabled")

        self.btn_go = ModernButton(row2, text="Refresh Cache" if self.mode == "STRING_BLACK" else "Search", command=self.toggle_search, primary=True)
        self.btn_go.pack(side=tk.RIGHT, anchor=tk.S)

        btn_bar = tk.Frame(filters, bg=config.COLOR_INPUT_BG, padx=8, pady=5)
        btn_bar.pack(fill=tk.X, pady=(4, 0))
        active_lines_dict = getattr(config, 'ACTIVE_LINES', {})

        if self.mode == "PRE_EL":
            tk.Label(btn_bar, text="Pre EL Stations:", font=config.FONT_BODY_BOLD, bg=config.COLOR_INPUT_BG).pack(side=tk.LEFT, padx=(0, 8))
            for st_idx in range(1, 15):
                st_code = f"{config.PRE_EL_STATION_PREFIX}{1000 + st_idx}"
                parent_line_idx = (st_idx + 1) // 2
                parent_key = f"{config.FINAL_EL_STATION_PREFIX}{1000 + parent_line_idx}"
                btn = ModernToggleButton(btn_bar, text=f"PreEL{st_idx}", initial_state=active_lines_dict.get(parent_key, True), padx=5, pady=2)
                btn.pack(side=tk.LEFT, padx=2)
                self.buttons_map[st_code] = btn
        else:
            tk.Label(btn_bar, text="Active Lines:", font=config.FONT_BODY_BOLD, bg=config.COLOR_INPUT_BG).pack(side=tk.LEFT, padx=(0, 8))
            for i in range(config.FINAL_EL_STATION_MIN, config.FINAL_EL_STATION_MAX + 1):
                st = f"{config.FINAL_EL_STATION_PREFIX}{i}"
                callback = (lambda state: self.render_string_black_dashboard()) if self.mode == "STRING_BLACK" else None
                btn = ModernToggleButton(btn_bar, text=f"Line {i - 1000}", initial_state=active_lines_dict.get(st, True), on_toggle=callback, padx=8, pady=2)
                btn.pack(side=tk.LEFT, padx=4)
                self.buttons_map[st] = btn

        self.results_area = tk.Frame(self, bg=config.COLOR_BG_MAIN)
        self.results_area.pack(fill=tk.BOTH, expand=True, pady=4)
        self.results_canvas = tk.Canvas(self.results_area, bg=config.COLOR_BG_MAIN, highlightthickness=0)
        scrollbar = ttk.Scrollbar(self.results_area, orient="vertical", command=self.results_canvas.yview)
        self.valid_frame = tk.Frame(self.results_canvas, bg=config.COLOR_BG_MAIN)
        self.valid_frame.bind("<Configure>", lambda e: self.results_canvas.config(scrollregion=self.results_canvas.bbox("all")))
        self.results_canvas.create_window((0, 0), window=self.valid_frame, anchor="nw", tags="inner_frame")
        self.results_canvas.configure(yscrollcommand=scrollbar.set)
        
        def _on_results_canvas_configure(event):
            self.results_canvas.itemconfig("inner_frame", width=event.width)

        self.results_canvas.bind("<Configure>", _on_results_canvas_configure)
        self.results_canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        self.result_container = self.valid_frame

        if self.mode == "STRING_BLACK": self.after(500, self.render_string_black_dashboard)

    def get_selected_stations(self): return [st for st, btn in self.buttons_map.items() if btn.get_state()]
    def sync_line_buttons_from_config(self):
        active_lines_dict = getattr(config, 'ACTIVE_LINES', {})
        for st, btn in self.buttons_map.items(): btn.set_state(active_lines_dict.get(st, True))

    def update_time_slots(self, interval_str):
        self.sync_line_buttons_from_config()
        if hasattr(self, 'cb_timeslot'):
            slots = generate_shift_time_slots(interval_str)
            self.cb_timeslot['values'] = slots
            self.cb_timeslot.set(get_current_shift_slot(slots))
            self.render_string_black_dashboard()

    def reset_for_new_shift(self, shift_info=None):
        if hasattr(self, 'txt_sn'):
            self.txt_sn.delete("1.0", tk.END)
        self.current_results.clear()
        self.raw_ng_images.clear()
        self.sn_list_ordered.clear()
        if hasattr(self, 'lbl_summary'):
            shift_label = shift_info.get('name', 'New Shift') if shift_info else 'New Shift'
            self.lbl_summary.config(text=f"Shift Rollover: {shift_label}")
        if hasattr(self, 'btn_export'):
            self.btn_export.config(state="disabled", text="Export Excel")
        if hasattr(self, 'btn_export_imgs'):
            self.btn_export_imgs.config(state="disabled", text="Export Images")
        if hasattr(self, 'result_container'):
            for w in self.result_container.winfo_children():
                w.destroy()
        if self.mode == "STRING_BLACK" and hasattr(self, 'cb_timeslot'):
            slots = generate_shift_time_slots(getattr(config, 'DEFAULT_TIME_INTERVAL', '2 Hours'))
            self.cb_timeslot['values'] = slots
            self.cb_timeslot.set(get_current_shift_slot(slots))
            self.render_string_black_dashboard()

    def toggle_search(self):
        if self.app.is_searching: self.app.stop_search()
        else: self.app.start_search_for_tab(self)

    def _show_single_sn_results(self, sn):
        self.app.highlight_sidebar_sn(self, sn)
        results = self.current_results.get(sn, [])
        for w in self.result_container.winfo_children(): w.destroy()
        if not results:
            tk.Label(self.result_container, text=f"No data found for {sn}", font=config.FONT_TITLE, bg=config.COLOR_BG_MAIN, fg=config.COLOR_STATUS_NG).pack(pady=40)
            return

        groups = self.app.search_engine.group_by_timestamp(results)
        timestamps = list(groups.keys())
        pass_count = len(timestamps)

        latest_ts = timestamps[0] if timestamps else "N/A"
        latest_imgs = groups.get(latest_ts, [])

        # 1. Top Header Banner with SN, Latest Inspection Time, and Passed Count Badge
        header_card = tk.Frame(self.result_container, bg="#ffffff", bd=1, relief=tk.SOLID, padx=14, pady=10)
        header_card.pack(fill=tk.X, pady=(6, 12), padx=4)

        top_info = tk.Frame(header_card, bg="#ffffff")
        top_info.pack(fill=tk.X)

        tk.Label(top_info, text=f"{sn}", font=("Segoe UI", 13, "bold"), bg="#ffffff", fg=config.COLOR_PRIMARY).pack(side=tk.LEFT)
        tk.Label(top_info, text=f"  •  Last Upload: {latest_ts}", font=("Segoe UI", 11, "bold"), bg="#ffffff", fg="#334155").pack(side=tk.LEFT, padx=(6, 0))

        # Pass Count Badge (e.g. "Passed Station: 7 times" or "Passed Station: 1 time")
        badge_text = f"🎯 Passed Station: {pass_count} {'times' if pass_count != 1 else 'time'}"
        badge_bg = "#fef3c7" if pass_count > 1 else "#dcfce7"
        badge_fg = "#b45309" if pass_count > 1 else "#15803d"

        lbl_badge = tk.Label(top_info, text=f"  {badge_text}  ", font=("Segoe UI", 10, "bold"), bg=badge_bg, fg=badge_fg, bd=1, relief=tk.SOLID)
        lbl_badge.pack(side=tk.RIGHT)

        # 2. Section for Latest Inspection Cards (Front, EL, Back for Pre EL; EL for Final EL)
        latest_section = tk.Frame(self.result_container, bg=config.COLOR_BG_MAIN)
        latest_section.pack(fill=tk.X, pady=(0, 8), padx=2)

        lbl_section = tk.Label(latest_section, text=f"LATEST INSPECTION RUN ({latest_ts})", font=("Segoe UI", 9, "bold"), bg=config.COLOR_BG_MAIN, fg="#64748b")
        lbl_section.pack(anchor="w", padx=6, pady=(0, 4))

        grid = tk.Frame(latest_section, bg=config.COLOR_BG_MAIN)
        grid.pack(fill=tk.X, pady=(0, 6), padx=4)

        num_cols = min(3, max(1, len(latest_imgs)))
        for idx, img in enumerate(latest_imgs):
            row = idx // 3
            col = idx % 3
            card = ResultCard(grid, img, self.app.open_original_image)
            card.grid(row=row, column=col, padx=4, pady=4, sticky="nsew")
        for c in range(num_cols):
            grid.grid_columnconfigure(c, weight=1, uniform="card_col")

        # 3. Optional expandable history container if there are multiple historical runs
        if pass_count > 1:
            history_container = tk.Frame(self.result_container, bg=config.COLOR_BG_MAIN)
            history_container.pack(fill=tk.X, pady=(8, 4), padx=2)

            history_content = tk.Frame(history_container, bg=config.COLOR_BG_MAIN)
            history_shown = [False]

            def toggle_history(btn_toggle=None):
                if history_shown[0]:
                    history_content.pack_forget()
                    history_shown[0] = False
                    if btn_toggle: btn_toggle.config(text=f"🕒 Show All {pass_count} Historical Passes ▼")
                else:
                    for w in history_content.winfo_children(): w.destroy()
                    for ts_idx, ts in enumerate(timestamps[1:], start=2):
                        imgs = groups[ts]
                        f_head = tk.Frame(history_content, bg=config.COLOR_BG_MAIN)
                        f_head.pack(fill=tk.X, pady=(10, 2))
                        tk.Label(f_head, text=f"Pass #{pass_count - ts_idx + 1}  •  {ts}", font=("Segoe UI", 10, "bold"), bg=config.COLOR_BG_MAIN, fg="#475569").pack(side=tk.LEFT, padx=6)
                        
                        h_grid = tk.Frame(history_content, bg=config.COLOR_BG_MAIN)
                        h_grid.pack(fill=tk.X, pady=(0, 6), padx=4)
                        h_cols = min(3, max(1, len(imgs)))
                        for idx, img in enumerate(imgs):
                            r = idx // 3
                            c = idx % 3
                            card = ResultCard(h_grid, img, self.app.open_original_image)
                            card.grid(row=r, column=c, padx=4, pady=4, sticky="nsew")
                        for c in range(h_cols):
                            h_grid.grid_columnconfigure(c, weight=1, uniform="card_col")
                    
                    history_content.pack(fill=tk.X, pady=(4, 10))
                    history_shown[0] = True
                    if btn_toggle: btn_toggle.config(text=f"▲ Hide Previous Passes")
                    self.results_canvas.config(scrollregion=self.results_canvas.bbox("all"))

            btn_hist = ModernButton(history_container, text=f"🕒 Show All {pass_count} Historical Passes ▼", command=lambda: toggle_history(btn_hist), primary=False, padx=10, pady=4)
            btn_hist.pack(anchor="w", padx=6, pady=4)

    def _draw_sb_bar_chart(self, canvas, c_w, c_h, data_pairs, bar_color="#1a5b82", title=""):
        pad_l, pad_r, pad_t, pad_b = 45, 20, 25, 35
        plot_w, plot_h = c_w - pad_l - pad_r, c_h - pad_t - pad_b
        canvas.delete("all")
        if title:
            canvas.create_text(pad_l, 12, text=title, font=("Segoe UI", 9, "bold"), fill="#1e293b", anchor="w")
        if not data_pairs:
            canvas.create_text(c_w/2, c_h/2, text="No data", font=("Segoe UI", 9), fill="#94a3b8")
            return
        max_val = max([v for k, v in data_pairs] or [1])
        y_max = max(5, ((max_val // 5) + 1) * 5)
        for i in range(6):
            val = int((y_max / 5) * i)
            y_pos = pad_t + plot_h - (i * (plot_h / 5))
            canvas.create_line(pad_l, y_pos, pad_l + plot_w, y_pos, fill="#f1f5f9", width=1)
            canvas.create_text(pad_l - 6, y_pos, text=str(val), fill="#64748b", font=("Segoe UI", 8), anchor="e")
        canvas.create_line(pad_l, pad_t + plot_h, pad_l + plot_w, pad_t + plot_h, fill="#cbd5e1", width=1.5)
        num_bars = len(data_pairs)
        bar_width = min(48, plot_w / max(1, num_bars * 1.6))
        for idx, (name, val) in enumerate(data_pairs):
            bx = pad_l + (idx * (plot_w / num_bars)) + (plot_w / num_bars / 2) - (bar_width / 2)
            bh = (val / y_max) * plot_h if y_max > 0 else 0
            by = pad_t + plot_h - bh
            canvas.create_rectangle(bx, by, bx + bar_width, pad_t + plot_h, fill=bar_color, outline="")
            canvas.create_text(bx + bar_width/2, by - 7, text=str(val), font=("Segoe UI", 8, "bold"), fill="#334155")
            canvas.create_text(bx + bar_width/2, pad_t + plot_h + 12, text=name, font=("Segoe UI", 8), fill="#475569", anchor="n")

    def render_string_black_dashboard(self):
        if self.mode != "STRING_BLACK": return
        for w in self.result_container.winfo_children(): w.destroy()

        d_from = self.date_from_val.get('date', datetime.now())
        date_str = d_from.strftime("%Y%m%d")

        # Load cache if not already loaded for this date
        if getattr(self, 'current_loaded_cache_date', None) != date_str or not hasattr(self, 'raw_string_black_data') or self.raw_string_black_data is None:
            cache_file = os.path.join(getattr(config, 'BASE_DIR', '.'), f"StringBlackData_{date_str}.json")
            if os.path.exists(cache_file):
                try:
                    with open(cache_file, 'r', encoding='utf-8') as f:
                        loaded = json.load(f)
                    for item in loaded:
                        if 'datetime' not in item and 'datetime_str' in item:
                            try:
                                item['datetime'] = datetime.strptime(item['datetime_str'], "%Y-%m-%d %H:%M:%S")
                            except Exception: pass
                    self.raw_string_black_data = loaded
                    self.current_loaded_cache_date = date_str
                except Exception as e:
                    print(f"[STRING BLACK LOAD ERROR]: {e}")
                    self.raw_string_black_data = []
            else:
                self.raw_string_black_data = []

        # Parse selected filter window
        slot_text = self.cb_timeslot.get() if hasattr(self, 'cb_timeslot') else "All Shift"
        time_range = None
        if "All Shift" not in slot_text and " - " in slot_text:
            try:
                p1, p2 = slot_text.split(' - ')
                sh, sm = map(int, p1.strip().split(':'))
                eh, em = map(int, p2.strip().split(':'))
                from datetime import time as dt_time
                time_range = (dt_time(sh, sm), dt_time(eh, em))
            except Exception: pass

        # Filter by selected lines & time window
        selected_stations = self.get_selected_stations()
        filtered = []
        for r in (self.raw_string_black_data or []):
            st = r.get('station', '')
            if selected_stations and st not in selected_stations:
                continue
            dt = r.get('datetime')
            if dt and time_range:
                t_val = dt.time()
                t_start, t_end = time_range
                if t_start <= t_end:
                    if not (t_start <= t_val <= t_end): continue
                else:
                    if not (t_val >= t_start or t_val <= t_end): continue
            filtered.append(r)

        self.filtered_string_black_data = filtered
        total_cnt = len(filtered)
        self.lbl_summary.config(text=f"Total String Black NG: {total_cnt}")

        if not filtered:
            card = tk.Frame(self.result_container, bg="white", bd=1, relief=tk.SOLID, padx=25, pady=25)
            card.pack(fill=tk.X, padx=15, pady=25)
            tk.Label(card, text=f"No String Black records for {d_from.strftime('%Y-%m-%d')}", font=("Segoe UI", 12, "bold"), bg="white", fg=config.COLOR_PRIMARY).pack(pady=(0, 4))
            tk.Label(card, text="Filter Window: " + slot_text + f"  •  Active Lines: {len(selected_stations)}", font=("Segoe UI", 9), bg="white", fg=config.COLOR_TEXT_SECONDARY).pack(pady=(0, 12))
            tk.Label(card, text="💡 Click 'Refresh Cache' above to scan factory station folders for this date,\nor choose a date with existing records (e.g. 2026-08-25).", font=("Segoe UI", 9), bg="white", fg="#475569").pack(pady=(0, 15))
            ModernButton(card, text="⚡ Scan Final EL Stations Now", command=self.toggle_search, primary=True).pack()
            self.btn_export.config(state="disabled", text="Export Excel")
            if hasattr(self, 'btn_export_imgs'):
                self.btn_export_imgs.config(state="disabled", text="Export Images")
            return

        # Metrics computation
        male_cnt = sum(1 for r in filtered if r.get('male'))
        mid_cnt = sum(1 for r in filtered if r.get('middle'))
        fem_cnt = sum(1 for r in filtered if r.get('female'))
        
        line_counts = {f"Line {i}": 0 for i in range(1, 8)}
        for r in filtered:
            st = r.get('station', '')
            l_name = parse_line_from_station(st) or st
            if l_name in line_counts:
                line_counts[l_name] += 1

        # 1. Header Card with KPI Pills
        kpi_frame = tk.Frame(self.result_container, bg="white", bd=1, relief=tk.SOLID, padx=14, pady=10)
        kpi_frame.pack(fill=tk.X, padx=6, pady=(6, 10))

        top_info = tk.Frame(kpi_frame, bg="white")
        top_info.pack(fill=tk.X, pady=(0, 8))
        tk.Label(top_info, text=f"String Black Analytics: {d_from.strftime('%Y-%m-%d')}", font=("Segoe UI", 12, "bold"), bg="white", fg=config.COLOR_PRIMARY).pack(side=tk.LEFT)
        tk.Label(top_info, text=f"  •  Window: {slot_text}", font=("Segoe UI", 10, "bold"), bg="white", fg="#64748b").pack(side=tk.LEFT)

        pill_row = tk.Frame(kpi_frame, bg="white")
        pill_row.pack(fill=tk.X)

        def make_pill(parent, label, count, pct, bg_c, fg_c):
            f = tk.Frame(parent, bg=bg_c, bd=1, relief=tk.SOLID, padx=8, pady=4)
            f.pack(side=tk.LEFT, padx=(0, 8))
            tk.Label(f, text=f"{label}: {count} ({pct}%)", font=("Segoe UI", 9, "bold"), bg=bg_c, fg=fg_c).pack()

        make_pill(pill_row, "🚨 Total NG", total_cnt, 100, "#fee2e2", "#991b1b")
        make_pill(pill_row, "Top Zone (Male)", male_cnt, round(male_cnt/total_cnt*100, 1), "#fef3c7", "#92400e")
        make_pill(pill_row, "Middle Zone", mid_cnt, round(mid_cnt/total_cnt*100, 1), "#fed7aa", "#9a3412")
        make_pill(pill_row, "Bot Zone (Female)", fem_cnt, round(fem_cnt/total_cnt*100, 1), "#f3e8ff", "#6b21a8")

        # 2. Charts Section
        chart_section = tk.Frame(self.result_container, bg=config.COLOR_BG_MAIN)
        chart_section.pack(fill=tk.X, padx=4, pady=(0, 10))

        # Chart 1: Zone Distribution
        w_zone = tk.LabelFrame(chart_section, text=" Zone Breakdown (Male / Middle / Female) ", font=config.FONT_BODY_BOLD, bg="white", fg=config.COLOR_PRIMARY, padx=8, pady=8)
        w_zone.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 6))
        c_zone = tk.Canvas(w_zone, width=440, height=200, bg="white", highlightthickness=0)
        c_zone.pack(fill=tk.BOTH, expand=True)
        self._draw_sb_bar_chart(c_zone, 440, 200, [("Male (Top)", male_cnt), ("Middle", mid_cnt), ("Female (Bot)", fem_cnt)], bar_color="#e11d48")

        # Chart 2: Line Breakdown
        w_line = tk.LabelFrame(chart_section, text=" Defect Count by Production Line ", font=config.FONT_BODY_BOLD, bg="white", fg=config.COLOR_PRIMARY, padx=8, pady=8)
        w_line.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(6, 0))
        c_line = tk.Canvas(w_line, width=540, height=200, bg="white", highlightthickness=0)
        c_line.pack(fill=tk.BOTH, expand=True)
        self._draw_sb_bar_chart(c_line, 540, 200, list(line_counts.items()), bar_color="#1a5b82")

        # 3. Defect Cards Gallery
        gallery_header = tk.Frame(self.result_container, bg=config.COLOR_BG_MAIN)
        gallery_header.pack(fill=tk.X, padx=6, pady=(4, 4))
        tk.Label(gallery_header, text=f"DETECTED STRING BLACK DEFECTS ({total_cnt})", font=("Segoe UI", 10, "bold"), bg=config.COLOR_BG_MAIN, fg="#475569").pack(side=tk.LEFT)

        grid = tk.Frame(self.result_container, bg=config.COLOR_BG_MAIN)
        grid.pack(fill=tk.X, padx=4, pady=(0, 10))

        for idx, item in enumerate(filtered):
            r = idx // 3
            c = idx % 3
            card = StringBlackCard(grid, item, self.app.open_original_image)
            card.grid(row=r, column=c, padx=4, pady=4, sticky="nsew")

        for c_idx in range(min(3, max(1, len(filtered)))):
            grid.grid_columnconfigure(c_idx, weight=1, uniform="sb_card_col")

        self.btn_export.config(state="normal", text=f"Export Excel ({total_cnt})")
        if hasattr(self, 'btn_export_imgs'):
            self.btn_export_imgs.config(state="normal", text=f"Export Images ({total_cnt})")
        self.results_canvas.config(scrollregion=self.results_canvas.bbox("all"))

    def export_excel(self):
        if self.mode == "STRING_BLACK":
            records_to_export = getattr(self, 'filtered_string_black_data', [])
            if not records_to_export:
                messagebox.showwarning("Export Excel", "No String Black records available to export!")
                return
            d_val = self.date_from_val.get('date', datetime.now())
            default_name = f"String_Black_Analytics_{d_val.strftime('%Y%m%d')}_{datetime.now().strftime('%H%M%S')}.xlsx"
            filename = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                initialfile=default_name,
                filetypes=[("Excel Files", "*.xlsx")],
                title="Export String Black Analytics to Excel"
            )
            if not filename: return
            try:
                import openpyxl
                from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "String Black Defects"
                headers = ["Date", "Time", "Serial No", "Line", "Station", "Dark %", "Male (Top)", "Middle", "Female (Bot)", "Status", "Image Path"]
                ws.append(headers)
                ws.row_dimensions[1].height = 26

                header_fill = PatternFill(start_color="1A5B82", end_color="1A5B82", fill_type="solid")
                header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
                for col_idx in range(1, len(headers) + 1):
                    c = ws.cell(row=1, column=col_idx)
                    c.fill = header_fill
                    c.font = header_font
                    c.alignment = Alignment(horizontal="center", vertical="center")

                for r_idx, item in enumerate(records_to_export, start=2):
                    dt = item.get('datetime')
                    d_str = dt.strftime("%Y-%m-%d") if isinstance(dt, datetime) else str(item.get('datetime_str', ''))[:10]
                    t_str = dt.strftime("%H:%M:%S") if isinstance(dt, datetime) else str(item.get('datetime_str', ''))[11:]
                    sn_val = item.get('sn', '')
                    st_val = item.get('station', '')
                    line_val = parse_line_from_station(st_val) or st_val
                    dark_pct = item.get('dark_pct', 0.0)
                    m_val = "NG" if item.get('male') else "OK"
                    mid_val = "NG" if item.get('middle') else "OK"
                    f_val = "NG" if item.get('female') else "OK"
                    st_status = item.get('status', 'NG')
                    path_val = item.get('path', '')

                    ws.append([d_str, t_str, sn_val, line_val, st_val, dark_pct, m_val, mid_val, f_val, st_status, path_val])

                for col in ws.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    col_letter = col[0].column_letter
                    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

                wb.save(filename)
                resp = messagebox.askyesno("Export Complete", f"Exported {len(records_to_export)} String Black records to:\n{filename}\n\nOpen Excel file now?")
                if resp: os.startfile(filename)
            except Exception as e:
                messagebox.showerror("Export Error", f"Failed to export Excel:\n{e}")
            return

        if not self.current_results or not any(self.current_results.values()):
            messagebox.showwarning("Export Excel", "No search results available to export!")
            return

        default_name = f"Pre_EL_Last_EL_Time_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx" if self.mode == "PRE_EL" else f"Final_EL_Last_EL_Time_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filename = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            initialfile=default_name,
            filetypes=[("Excel Files", "*.xlsx")],
            title=f"Export {self.mode.replace('_', ' ')} Last EL Time to Excel"
        )
        if not filename:
            return

        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Pre EL Results" if self.mode == "PRE_EL" else "Final EL Results"

            # 4 Standard ABCD Columns: Date, Serial No, Last EL Time, Station
            headers = ["Date", "Serial No", "Last EL Time", "Station"]
            ws.append(headers)
            ws.row_dimensions[1].height = 26

            header_fill = PatternFill(start_color="1A5B82", end_color="1A5B82", fill_type="solid")
            header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
            thin_border = Border(
                left=Side(style='thin', color='E2E8F0'),
                right=Side(style='thin', color='E2E8F0'),
                top=Side(style='thin', color='E2E8F0'),
                bottom=Side(style='thin', color='E2E8F0')
            )

            for col_idx in range(1, 5):
                c = ws.cell(row=1, column=col_idx)
                c.fill = header_fill
                c.font = header_font
                c.alignment = Alignment(horizontal="center", vertical="center")

            ordered_sns = self.sn_list_ordered if self.sn_list_ordered else list(self.current_results.keys())
            r_idx = 2

            for sn in ordered_sns:
                img_list = self.current_results.get(sn, [])
                if not img_list:
                    ws.cell(row=r_idx, column=1, value="-")
                    ws.cell(row=r_idx, column=2, value=sn)
                    c_stat = ws.cell(row=r_idx, column=3, value="No Record")
                    c_stat.font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
                    ws.cell(row=r_idx, column=4, value="-")
                    ws.row_dimensions[r_idx].height = 24
                    for c in range(1, 5):
                        cell_obj = ws.cell(row=r_idx, column=c)
                        cell_obj.border = thin_border
                        cell_obj.alignment = Alignment(horizontal="center", vertical="center")
                    r_idx += 1
                    continue

                groups = self.app.search_engine.group_by_timestamp(img_list)
                latest_ts = list(groups.keys())[0] if groups else ""
                latest_imgs = groups.get(latest_ts, [])

                # Pick EL image or any image from latest inspection run
                el_img = next((img for img in latest_imgs if "el" in str(img.get('category', '')).lower()), None)
                if not el_img and latest_imgs:
                    el_img = latest_imgs[0]

                dt_val = el_img.get('datetime') if el_img else None
                date_str = dt_val.strftime('%Y-%m-%d') if dt_val else (latest_ts.split(' ')[0] if ' ' in latest_ts else "-")
                time_str = dt_val.strftime('%Y-%m-%d %H:%M:%S') if dt_val else latest_ts
                st_val = el_img.get('station', '') if el_img else ""

                ws.cell(row=r_idx, column=1, value=date_str)
                ws.cell(row=r_idx, column=2, value=sn)
                ws.cell(row=r_idx, column=3, value=time_str)
                ws.cell(row=r_idx, column=4, value=st_val)
                ws.row_dimensions[r_idx].height = 24

                for c in range(1, 5):
                    cell_obj = ws.cell(row=r_idx, column=c)
                    cell_obj.border = thin_border
                    cell_obj.alignment = Alignment(horizontal="center", vertical="center")
                    cell_obj.font = Font(name="Segoe UI", size=9)

                r_idx += 1

            ws.column_dimensions['A'].width = 16
            ws.column_dimensions['B'].width = 24
            ws.column_dimensions['C'].width = 24
            ws.column_dimensions['D'].width = 18

            wb.save(filename)
            messagebox.showinfo("Export Excel", f"Successfully exported Last EL Time (Columns A-D) for {len(ordered_sns)} SNs to:\n{filename}")
        except Exception as e:
            messagebox.showerror("Export Error", f"Failed to export Excel report:\n{e}")

    def export_images_to_folder(self):
        if self.mode == "STRING_BLACK":
            records_to_export = getattr(self, 'filtered_string_black_data', [])
            if not records_to_export:
                messagebox.showwarning("Export Images", "No String Black images available to export!")
                return
            dest_folder = filedialog.askdirectory(title="Select Destination Folder for Exported Images")
            if not dest_folder: return
            try:
                import shutil
                copied_count = 0
                for item in records_to_export:
                    src_path = item.get('path')
                    if src_path and os.path.exists(src_path):
                        sn = item.get('sn', 'UNKNOWN')
                        fname = os.path.basename(src_path)
                        dst_name = f"{sn}_{fname}" if sn not in fname else fname
                        dst_path = os.path.join(dest_folder, dst_name)
                        if os.path.exists(dst_path):
                            base, ext = os.path.splitext(dst_name)
                            dst_path = os.path.join(dest_folder, f"{base}_{copied_count+1}{ext}")
                        shutil.copy2(src_path, dst_path)
                        copied_count += 1
                if copied_count > 0:
                    resp = messagebox.askyesno("Export Complete", f"Successfully copied {copied_count} images to:\n{dest_folder}\n\nOpen destination folder now?")
                    if resp: os.startfile(dest_folder)
                else:
                    messagebox.showinfo("Export", "No image files found to copy (check network connection).")
            except Exception as e:
                messagebox.showerror("Export Error", f"Failed to export images:\n{e}")
            return

        if not self.current_results or not any(self.current_results.values()):
            messagebox.showwarning("Export Images", "No search results available to export!")
            return

        dest_folder = filedialog.askdirectory(title="Select Destination Folder for Exported Images")
        if not dest_folder:
            return

        try:
            import shutil
            copied_count = 0
            for sn, img_list in self.current_results.items():
                if not img_list: continue
                for img in img_list:
                    src_path = img.get('path')
                    if src_path and os.path.exists(src_path):
                        fname = os.path.basename(src_path)
                        dst_name = f"{sn}_{fname}" if sn not in fname else fname
                        dst_path = os.path.join(dest_folder, dst_name)
                        
                        if os.path.exists(dst_path):
                            base, ext = os.path.splitext(dst_name)
                            dst_path = os.path.join(dest_folder, f"{base}_{copied_count+1}{ext}")

                        shutil.copy2(src_path, dst_path)
                        copied_count += 1

            if copied_count > 0:
                resp = messagebox.askyesno(
                    "Export Complete",
                    f"Successfully copied {copied_count} images to:\n{dest_folder}\n\nDo you want to open the destination folder now?"
                )
                if resp:
                    os.startfile(dest_folder)
            else:
                messagebox.showinfo("Export", "No image files found to copy.")
        except Exception as e:
            messagebox.showerror("Export Error", f"Failed to export images:\n{e}")

# =================== MAIN APPLICATION ===================

class AOIDashboardApp:
    def __init__(self, root):
        self.root = root
        self.root.title(config.WINDOW_TITLE)
        self.root.geometry(config.WINDOW_SIZE)
        self.root.configure(bg=config.COLOR_BG_MAIN)
        
        self.search_engine = ImageSearchEngine()
        self.image_processor = ImageProcessor()
        self.data_manager = DataManager()
        self.global_review_records = []
        self.phone_server_url = "Phone Link: Initializing..."
        self.all_tabs = []
        self.is_searching = False
        self.search_id = 0
        self.active_tab = None
        self.tab_counters = {"PRE_EL": 0, "FINAL_EL": 0, "STRING_BLACK": 0}
        self.current_app_date = datetime.now().date()
        self.current_shift_info = get_operational_shift_info()
        self.current_shift_key = self.current_shift_info['shift_key']
        self.incoming_event_queue = queue.Queue()
        self.setup_styles()
        self.setup_layout()
        self.phone_server_url = start_phone_server(
            self.on_phone_photo_uploaded,
            defect_callback=self.on_mobile_defect_logged,
            mes_callback=self.on_mes_record_logged
        )

        for tab in self.all_tabs:
            if isinstance(tab, ModuleReviewTab):
                if hasattr(tab, 'lbl_server_url'):
                    tab.lbl_server_url.config(text=self.phone_server_url)
                if hasattr(tab, 'update_qr_code'):
                    tab.update_qr_code(self.phone_server_url)
        self.check_initial_shift_state()
        self.check_shift_and_day_rollover()
        self._poll_incoming_events()

    def on_mes_record_logged(self, data):
        self.incoming_event_queue.put(('MES_LOG', data))

    def on_mobile_defect_logged(self, data):
        self.incoming_event_queue.put(('DEFECT', data))

    def on_phone_photo_uploaded(self, file_path, extracted_sn=None, action="NEW_RECORD", v_summary="", v_class="", v_result="", raw_voice="", audio_path="", file_dt=None):
        self.incoming_event_queue.put(('PHOTO', {
            'file_path': file_path,
            'extracted_sn': extracted_sn,
            'action': action,
            'v_summary': v_summary,
            'v_class': v_class,
            'v_result': v_result,
            'raw_voice': raw_voice,
            'audio_path': audio_path,
            'file_dt': file_dt
        }))

    def _poll_incoming_events(self):
        try:
            while not self.incoming_event_queue.empty():
                evt_type, data = self.incoming_event_queue.get_nowait()
                if evt_type == 'MES_LOG':
                    for tab in self.all_tabs:
                        if isinstance(tab, MESProcessLogTab):
                            tab.handle_incoming_mes_record(data)
                            break
                elif evt_type == 'DEFECT':
                    for tab in self.all_tabs:
                        if isinstance(tab, ModuleReviewTab):
                            tab.handle_mobile_defect_logged(data)
                            break
                elif evt_type == 'PHOTO':
                    for tab in self.all_tabs:
                        if isinstance(tab, ModuleReviewTab):
                            tab.handle_incoming_phone_upload(
                                data.get('file_path'),
                                data.get('extracted_sn'),
                                data.get('action', 'NEW_RECORD'),
                                data.get('v_summary', ''),
                                data.get('v_class', ''),
                                data.get('v_result', ''),
                                data.get('raw_voice', ''),
                                data.get('audio_path', ''),
                                data.get('file_dt')
                            )
                            break
                elif evt_type == 'REFRESH':
                    for tab in self.all_tabs:
                        if isinstance(tab, ModuleReviewTab):
                            tab.refresh_treeview()
                            rec_id = data.get('rec_id') if isinstance(data, dict) else None
                            line_val = data.get('line') if isinstance(data, dict) else None
                            if line_val and hasattr(tab, 'cb_line'):
                                if tab.selected_rec_id == rec_id or not tab.selected_rec_id:
                                    tab.cb_line.set(line_val)
                            break
        except Exception as e:
            print(f"[EVENT POLL ERROR]: {e}")
        finally:
            try:
                self.root.after(50, self._poll_incoming_events)
            except Exception:
                pass

    def check_initial_shift_state(self):
        shift_state_file = os.path.join(config.LOCAL_DATA_DIR, "shift_state.json")
        last_shift_key = None
        if os.path.exists(shift_state_file):
            try:
                with open(shift_state_file, 'r', encoding='utf-8') as f:
                    last_shift_key = json.load(f).get('shift_key')
            except Exception:
                pass

        current_key_str = f"{self.current_shift_info['date']}_{'DAY' if self.current_shift_info['is_day'] else 'NIGHT'}"

        needs_reset = False
        if last_shift_key and last_shift_key != current_key_str:
            print(f"[STARTUP SHIFT DETECTED]: Last recorded shift was '{last_shift_key}'. Current shift is '{current_key_str}'. Performing shift rollover reset and archiving...")
            needs_reset = True
        else:
            # Check if active MES trend log has older records from a previous date
            mes_trend_file = os.path.join(config.LOCAL_DATA_DIR, "mes_process_trend.json")
            if os.path.exists(mes_trend_file):
                try:
                    with open(mes_trend_file, 'r', encoding='utf-8') as f:
                        records = json.load(f)
                    today_str = datetime.now().strftime("%Y-%m-%d")
                    has_old_records = any(str(r.get('timestamp', ''))[:10] < today_str for r in records if r.get('timestamp'))
                    if has_old_records:
                        print(f"[STARTUP CHECK]: Found older shift records in active log. Archiving and resetting for fresh shift.")
                        needs_reset = True
                except Exception:
                    pass

        if needs_reset:
            self.on_operational_shift_changed(self.current_shift_info)

        # Record current shift key
        try:
            with open(shift_state_file, 'w', encoding='utf-8') as f:
                json.dump({'shift_key': current_key_str, 'updated_at': datetime.now().strftime("%Y-%m-%d %H:%M:%S")}, f, indent=2)
        except Exception as e:
            print(f"[SHIFT STATE SAVE ERROR]: {e}")

    def get_current_tab(self):
        try:
            current_id = self.main_notebook.select()
            if current_id:
                return self.main_notebook.nametowidget(current_id)
        except Exception:
            pass
        return None

    def check_shift_and_day_rollover(self):
        now = datetime.now()
        today = now.date()
        shift_info = get_operational_shift_info(now)
        shift_key = shift_info['shift_key']

        # 1. Calendar Day Rollover check (refresh save folders)
        if today != self.current_app_date:
            self.current_app_date = today
            print(f"[APP DATE ROLLOVER]: Day changed to {today}. Refreshing save folders.")
            refresh_save_folders_for_today()

        # 2. Operational Shift Rollover check (6:00 AM & 6:00 PM)
        if shift_key != self.current_shift_key:
            prev_shift = self.current_shift_info.get('name', 'Previous Shift')
            self.current_shift_info = shift_info
            self.current_shift_key = shift_key
            current_key_str = f"{shift_info['date']}_{'DAY' if shift_info['is_day'] else 'NIGHT'}"
            print(f"[APP SHIFT ROLLOVER]: Operational shift changed from {prev_shift} to {shift_info['name']} (Key: {current_key_str}). Resetting scanning & workspaces across all tabs...")
            shift_state_file = os.path.join(config.LOCAL_DATA_DIR, "shift_state.json")
            try:
                with open(shift_state_file, 'w', encoding='utf-8') as f:
                    json.dump({'shift_key': current_key_str, 'updated_at': datetime.now().strftime("%Y-%m-%d %H:%M:%S")}, f, indent=2)
            except Exception:
                pass
            self.on_operational_shift_changed(shift_info)

        # Check every 15 seconds
        self.root.after(15000, self.check_shift_and_day_rollover)

    def on_operational_shift_changed(self, shift_info):
        # 1. Reset all tabs
        for tab in self.all_tabs:
            if hasattr(tab, 'reset_for_new_shift'):
                try:
                    tab.reset_for_new_shift(shift_info)
                except Exception as e:
                    print(f"[SHIFT RESET TAB ERROR]: {e}")

        # 2. Clear sidebar SN navigation list
        try:
            curr_tab = self.get_current_tab()
            if curr_tab:
                self.refresh_sidebar_for_tab(curr_tab)
            else:
                for w in self.sidebar_nav_content.winfo_children():
                    w.destroy()
                self.lbl_sn_nav_title.config(text="Serial Numbers (0)")
        except Exception as e:
            print(f"[SHIFT RESET SIDEBAR ERROR]: {e}")

        # 3. Trigger phone server live HUD reset
        try:
            update_live_hud_state(reset_shift=True)
        except Exception as e:
            print(f"[SHIFT RESET HUD ERROR]: {e}")

    def setup_styles(self):
        style = ttk.Style()
        style.theme_use('clam')
        style.configure("TNotebook.Tab", font=config.FONT_BODY_BOLD, padding=[14, 7])

    def create_sidebar_content(self):
        brand_frame = tk.Frame(self.sidebar, bg=config.COLOR_BG_SIDEBAR)
        brand_frame.pack(fill=tk.X, padx=15, pady=15)
        tk.Label(brand_frame, text="AOI & EL Dashboard", font=config.FONT_SUBTITLE, bg=config.COLOR_BG_SIDEBAR, fg=config.COLOR_TEXT_MAIN).pack(anchor=tk.W)
        tk.Label(brand_frame, text="v18.6 - Complete Suite", font=config.FONT_SMALL, bg=config.COLOR_BG_SIDEBAR, fg=config.COLOR_TEXT_SECONDARY).pack(anchor=tk.W)

        nav_header = tk.Frame(self.sidebar, bg=config.COLOR_BG_SIDEBAR)
        nav_header.pack(fill=tk.X, padx=15, pady=(15, 8))
        self.lbl_sn_nav_title = tk.Label(nav_header, text="Serial Numbers (0)", font=config.FONT_BODY_BOLD, bg=config.COLOR_BG_SIDEBAR, fg=config.COLOR_TEXT_SECONDARY)
        self.lbl_sn_nav_title.pack(side=tk.LEFT)

        self.sidebar_nav_canvas = tk.Canvas(self.sidebar, bg=config.COLOR_BG_SIDEBAR, highlightthickness=0)
        self.sidebar_nav_scroll = ttk.Scrollbar(self.sidebar, orient="vertical", command=self.sidebar_nav_canvas.yview)
        self.sidebar_nav_content = tk.Frame(self.sidebar_nav_canvas, bg=config.COLOR_BG_SIDEBAR)
        self.sidebar_nav_content.bind("<Configure>", lambda e: self.sidebar_nav_canvas.configure(scrollregion=self.sidebar_nav_canvas.bbox("all")))
        self.sidebar_nav_canvas.create_window((0, 0), window=self.sidebar_nav_content, anchor="nw")
        self.sidebar_nav_canvas.configure(yscrollcommand=self.sidebar_nav_scroll.set)
        self.sidebar_nav_canvas.pack(side="left", fill="both", expand=True, padx=(10, 0), pady=(0, 25))
        self.sidebar_nav_scroll.pack(side="right", fill="y", pady=(0, 25))

    def setup_layout(self):

        self.sidebar = tk.Frame(self.root, bg=config.COLOR_BG_SIDEBAR, width=config.LEFT_PANEL_WIDTH)
        self.sidebar.pack(side=tk.LEFT, fill=tk.Y)
        self.sidebar.pack_propagate(False)
        self.create_sidebar_content()

        self.content = tk.Frame(self.root, bg=config.COLOR_BG_MAIN)
        self.content.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=15, pady=10)

        top_bar = tk.Frame(self.content, bg=config.COLOR_BG_MAIN)
        top_bar.pack(fill=tk.X, pady=(0, 5))
        
        ModernButton(top_bar, text="+ Pre EL Tab", command=lambda: self.add_new_tab("PRE_EL"), primary=False, padx=8, pady=3).pack(side=tk.LEFT, padx=(0, 4))
        ModernButton(top_bar, text="+ Final EL Tab", command=lambda: self.add_new_tab("FINAL_EL"), primary=False, padx=8, pady=3).pack(side=tk.LEFT, padx=(0, 4))
        ModernButton(top_bar, text="+ String Black Tab", command=lambda: self.add_new_tab("STRING_BLACK"), primary=False, padx=8, pady=3).pack(side=tk.LEFT, padx=(0, 4))
        ModernButton(top_bar, text="+ Module Review", command=self.add_defect_review_tab, primary=True, padx=8, pady=3).pack(side=tk.LEFT, padx=(0, 4))
        ModernButton(top_bar, text="+ Top 5 Analytics", command=self.add_analytics_tab, primary=False, padx=8, pady=3).pack(side=tk.LEFT, padx=(0, 4))
        ModernButton(top_bar, text="+ MES Process Log", command=self.add_mes_log_tab, primary=False, padx=8, pady=3).pack(side=tk.LEFT, padx=(0, 4))
        ModernButton(top_bar, text="🇨🇳 Chinese MES Guide", command=self.open_chinese_mes_guide, primary=False, padx=8, pady=3).pack(side=tk.LEFT)
        ModernButton(top_bar, text="Close Tab", command=self.close_current_tab, primary=False, padx=10, pady=3).pack(side=tk.RIGHT)
        ModernButton(top_bar, text="📄 View Log", command=self.open_runtime_log, primary=False, padx=10, pady=3).pack(side=tk.RIGHT, padx=(0, 6))

        self.main_notebook = ttk.Notebook(self.content)
        self.main_notebook.pack(fill=tk.BOTH, expand=True)

        self.all_tabs = []
        self.add_new_tab("PRE_EL")
        self.add_new_tab("FINAL_EL")
        self.add_new_tab("STRING_BLACK")
        self.add_defect_review_tab()
        self.add_analytics_tab()
        self.add_mes_log_tab()
        
        self.tab_config = ConfigPanel(self.main_notebook, self)
        self.main_notebook.add(self.tab_config, text="  Config  ")
        self.main_notebook.bind("<<NotebookTabChanged>>", self.on_tab_changed)

    def _insert_tab_before_config(self, tab, label):
        if hasattr(self, 'tab_config') and str(self.tab_config) in self.main_notebook.tabs():
            self.main_notebook.insert(self.tab_config, tab, text=f"  {label}  ")
        else:
            self.main_notebook.add(tab, text=f"  {label}  ")
        self.main_notebook.select(tab)

    def add_new_tab(self, mode="PRE_EL"):
        self.tab_counters[mode] += 1
        tab = TabContent(self.main_notebook, self, mode=mode)
        self.all_tabs.append(tab)
        label = f"Pre EL {self.tab_counters[mode]}" if mode == "PRE_EL" else (f"Final EL {self.tab_counters[mode]}" if mode == "FINAL_EL" else "String Black")
        self._insert_tab_before_config(tab, label)

    def add_defect_review_tab(self):
        tab = ModuleReviewTab(self.main_notebook, self)
        self.all_tabs.append(tab)
        self._insert_tab_before_config(tab, "Module Review")

    def add_analytics_tab(self):
        tab = Top5AnalyticsTab(self.main_notebook, self)
        self.all_tabs.append(tab)
        self._insert_tab_before_config(tab, "Top 5 & Scrap Analytics")

    def add_mes_log_tab(self):
        tab = MESProcessLogTab(self.main_notebook, self)
        self.all_tabs.append(tab)
        self._insert_tab_before_config(tab, "MES Process Log")

    def open_chinese_mes_guide(self):
        MESChineseFieldGuideDialog(self.root)

    def open_runtime_log(self):
        log_path = RUNTIME_LOG_FILE
        if not os.path.exists(log_path):
            try:
                with open(log_path, "w", encoding="utf-8") as f:
                    f.write(f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S.%f')[:-3]}] Runtime log file initialized.\n")
            except Exception:
                pass
        try:
            os.startfile(log_path)
        except Exception:
            try:
                subprocess.Popen(["notepad.exe", log_path])
            except Exception as e:
                messagebox.showerror("View Log Error", f"Could not open log file:\n{log_path}\n\nError: {e}")

    def close_current_tab(self):
        current_id = self.main_notebook.select()
        current_tab = self.main_notebook.nametowidget(current_id)
        if isinstance(current_tab, ConfigPanel): return
        if len(self.main_notebook.tabs()) <= 2: return
        if current_tab in self.all_tabs: self.all_tabs.remove(current_tab)
        self.main_notebook.forget(current_id)


    def sync_all_tabs(self, interval_str):
        for tab in self.all_tabs:
            if hasattr(tab, 'sync_line_buttons_from_config'): tab.sync_line_buttons_from_config()
            if getattr(tab, 'mode', None) == "STRING_BLACK": tab.update_time_slots(interval_str)
            if isinstance(tab, QCDefectReviewTab): tab.refresh_defect_dropdowns()

    def on_tab_changed(self, event):
        current_id = self.main_notebook.select()
        current_tab = self.main_notebook.nametowidget(current_id)
        if isinstance(current_tab, TabContent) and current_tab.mode != "STRING_BLACK":
            self.refresh_sidebar_for_tab(current_tab)
        else:
            self.lbl_sn_nav_title.config(text="Dashboard Mode")

    def highlight_sidebar_sn(self, tab, active_sn):
        tab.active_displayed_sn = active_sn
        if not hasattr(self, 'sidebar_sn_buttons') or not self.sidebar_sn_buttons:
            return

        for sn, item in self.sidebar_sn_buttons.items():
            btn = item.get('btn')
            if not btn or not btn.winfo_exists(): continue
            is_found = item.get('is_found', True)
            base_txt = item.get('txt', sn)

            if sn == active_sn:
                btn.config(
                    text=f"▶{base_txt}",
                    bg=config.COLOR_PRIMARY,
                    fg="#ffffff",
                    activebackground="#154c6d",
                    activeforeground="#ffffff",
                    font=("Segoe UI", 9, "bold"),
                    relief=tk.FLAT
                )
                try:
                    btn_y = btn.winfo_y()
                    total_h = self.sidebar_nav_content.winfo_height()
                    if total_h > 0 and btn_y > 0:
                        frac = max(0.0, min(1.0, (btn_y - 20) / total_h))
                        self.sidebar_nav_canvas.yview_moveto(frac)
                except Exception: pass
            else:
                btn.config(
                    text=base_txt,
                    bg=config.COLOR_INPUT_BG if is_found else "#fff0f0",
                    fg=config.COLOR_TEXT_MAIN if is_found else config.COLOR_STATUS_NG,
                    activebackground="#e2e8f0",
                    activeforeground=config.COLOR_TEXT_MAIN,
                    font=("Segoe UI", 9, "bold"),
                    relief=tk.FLAT
                )

    def refresh_sidebar_for_tab(self, tab):
        for w in self.sidebar_nav_content.winfo_children(): w.destroy()
        self.sidebar_sn_buttons = {}
        ordered_sns = getattr(tab, 'sn_list_ordered', [])
        all_results = getattr(tab, 'current_results', {})
        self.lbl_sn_nav_title.config(text=f"Serial Numbers ({len(ordered_sns)})")
        if not ordered_sns: return
        found_sns = [sn for sn, res in all_results.items() if res]
        active_sn = getattr(tab, 'active_displayed_sn', None)
        if not active_sn and found_sns:
            active_sn = found_sns[0]

        for sn in ordered_sns:
            is_found = sn in found_sns
            if is_found:
                res = all_results.get(sn, [])
                groups = self.search_engine.group_by_timestamp(res)
                p_cnt = len(groups)
                txt = f" {sn} ({p_cnt} {'passes' if p_cnt != 1 else 'pass'})"
            else:
                txt = f" {sn} (0)"
            btn = tk.Button(self.sidebar_nav_content, text=txt, font=("Segoe UI", 9, "bold"), bg=config.COLOR_INPUT_BG if is_found else "#fff0f0", fg=config.COLOR_TEXT_MAIN if is_found else config.COLOR_STATUS_NG, anchor="w", relief=tk.FLAT, bd=0, padx=8, pady=4, cursor="hand2", command=lambda s=sn: tab._show_single_sn_results(s))
            btn.pack(fill=tk.X, pady=2, padx=4)
            self.sidebar_sn_buttons[sn] = {
                'btn': btn,
                'is_found': is_found,
                'txt': txt
            }

        self.highlight_sidebar_sn(tab, active_sn)
        self.root.update_idletasks()
        self.sidebar_nav_canvas.config(scrollregion=self.sidebar_nav_canvas.bbox("all"))

    def start_search_for_tab(self, tab):
        if tab.mode == "STRING_BLACK":
            selected_stations = tab.get_selected_stations()
            if not selected_stations:
                messagebox.showwarning("Selection", "Please select at least one active Line!")
                return
            d_start = tab.date_from_val.get('date', datetime.now()).replace(hour=0, minute=0, second=0)
            d_end = tab.date_to_val.get('date', datetime.now()).replace(hour=23, minute=59, second=59)

            for w in tab.result_container.winfo_children(): w.destroy()
            tab.results_canvas.yview_moveto(0)
            tab.lbl_summary.config(text="Scanning Final EL NG directories...")
            tab.btn_export.config(state="disabled")
            if hasattr(tab, 'btn_export_imgs'):
                tab.btn_export_imgs.config(state="disabled")
            tab.progress_bar = ttk.Progressbar(tab.result_container, mode='indeterminate', length=350)
            tab.progress_bar.pack(pady=(40, 10))
            tab.progress_bar.start(15)

            self.is_searching = True
            self.active_tab = tab
            self.search_id += 1
            current_sid = self.search_id
            tab.btn_go.config(text="Stop", bg=config.COLOR_STATUS_NG)
            self.search_engine.cancel_flag.clear()
            threading.Thread(target=self._run_string_black_search, args=(tab, d_start, d_end, selected_stations, current_sid), daemon=True).start()
            return

        sn_list = [line.strip() for line in tab.txt_sn.get("1.0", tk.END).strip().split('\n') if line.strip()] if hasattr(tab, 'txt_sn') else []
        if not sn_list:
            messagebox.showwarning("Search", "Please enter at least one Serial Number!")
            return
        selected_stations = tab.get_selected_stations()
        if not selected_stations:
            messagebox.showwarning("Selection", "Please select at least one station!")
            return
        d_start = tab.date_from_val.get('date', datetime.now()).replace(hour=0, minute=0, second=0)
        d_end = tab.date_to_val.get('date', datetime.now()).replace(hour=23, minute=59, second=59)

        for w in tab.result_container.winfo_children(): w.destroy()
        tab.results_canvas.yview_moveto(0)
        tab.lbl_summary.config(text="")
        tab.btn_export.config(state="disabled")
        if hasattr(tab, 'btn_export_imgs'):
            tab.btn_export_imgs.config(state="disabled")
        tab.progress_bar = ttk.Progressbar(tab.result_container, mode='determinate', length=350)
        tab.progress_bar.pack(pady=(40, 10))
        
        self.is_searching = True
        self.active_tab = tab
        self.search_id += 1
        current_sid = self.search_id
        tab.btn_go.config(text="Stop", bg=config.COLOR_STATUS_NG)
        self.search_engine.cancel_flag.clear()
        threading.Thread(target=self._run_search, args=(tab, sn_list, d_start, d_end, selected_stations, current_sid), daemon=True).start()

    def stop_search(self):
        if self.is_searching and self.active_tab:
            self.search_engine.cancel_search()
            self.is_searching = False
            self.search_id += 1
            try:
                btn_txt = "Refresh Cache" if self.active_tab.mode == "STRING_BLACK" else "Search"
                self.active_tab.btn_go.config(text=btn_txt, state="normal", bg=config.COLOR_PRIMARY)
                self.active_tab.lbl_summary.config(text="Cancelled")
                for w in self.active_tab.result_container.winfo_children():
                    if isinstance(w, ttk.Progressbar):
                        try:
                            w.stop()
                            w.destroy()
                        except Exception: pass
            except Exception: pass

    def _update_overall_progress(self, tab, curr_sn_idx, total_sns, current_sn):
        try:
            if hasattr(tab, 'progress_bar') and tab.progress_bar.winfo_exists():
                tab.progress_bar['maximum'] = total_sns
                tab.progress_bar['value'] = curr_sn_idx
            if hasattr(tab, 'lbl_summary'):
                tab.lbl_summary.config(text=f"Searching SN {curr_sn_idx}/{total_sns}: {current_sn}")
        except Exception: pass

    def _run_string_black_search(self, tab, start, end, stations_list, search_id):
        if search_id != self.search_id: return
        res, _ = self.search_engine.search_final_el(
            serial_number="",
            stations_list=stations_list,
            start_date=start,
            end_date=end,
            time_filter=None,
            require_string_black=True,
            max_results=3000
        )
        if self.search_engine.cancel_flag.is_set() or search_id != self.search_id:
            return

        date_str = start.strftime("%Y%m%d")
        cache_file = os.path.join(getattr(config, 'BASE_DIR', '.'), f"StringBlackData_{date_str}.json")
        try:
            serializable = []
            for item in res:
                c = dict(item)
                if isinstance(c.get('datetime'), datetime):
                    c['datetime_str'] = c['datetime'].strftime("%Y-%m-%d %H:%M:%S")
                    del c['datetime']
                serializable.append(c)
            with open(cache_file, 'w', encoding='utf-8') as f:
                json.dump(serializable, f, indent=2)
            print(f"[STRING BLACK CACHE]: Saved {len(serializable)} records to {cache_file}")
        except Exception as e:
            print(f"[STRING BLACK CACHE ERROR]: {e}")

        self.root.after(0, lambda: self._display_string_black_results(tab, res, search_id, date_str))

    def _display_string_black_results(self, tab, results, search_id, date_str):
        if search_id != self.search_id: return
        self.is_searching = False
        tab.btn_go.config(state="normal", text="Refresh Cache", bg=config.COLOR_PRIMARY)
        for w in tab.result_container.winfo_children():
            if isinstance(w, ttk.Progressbar):
                try:
                    w.stop()
                    w.destroy()
                except Exception: pass
        tab.raw_string_black_data = results
        tab.current_loaded_cache_date = date_str
        tab.render_string_black_dashboard()

    def _run_search(self, tab, sn_list, start, end, stations_list, search_id):
        if search_id != self.search_id: return
        all_results = {}
        total_sns = len(sn_list)
        for idx, raw_sn in enumerate(sn_list):
            if self.search_engine.cancel_flag.is_set() or search_id != self.search_id: break
            sn = clean_and_validate_sn(raw_sn) or (raw_sn.strip().upper() if raw_sn else "")
            if not sn: continue
            self.root.after(0, lambda i=idx, s=sn: self._update_overall_progress(tab, i + 1, total_sns, s))
            sn_start, sn_end = start, end
            t_sn_start = time.time()
            if tab.mode == "PRE_EL":
                t0 = time.time()
                mes_dt, mes_station, _ = get_mes_layup_time_for_sn(sn)
                t_mes = time.time() - t0
                target_stations = get_pre_el_stations_for_layup_station(mes_station, stations_list)
                if mes_dt and (mes_dt < start or mes_dt > end):
                    sn_start = mes_dt.replace(hour=0, minute=0, second=0) - timedelta(days=1)
                    sn_end = mes_dt.replace(hour=23, minute=59, second=59) + timedelta(days=1)
                t1 = time.time()
                res, _ = self.search_engine.search_pre_el(sn, target_stations, sn_start, sn_end, layup_dt=mes_dt, max_results=getattr(config, 'MAX_IMAGES_PER_SN', 0))
                t_scan = time.time() - t1
                print(f"[BATCH TIMING]: Pre-EL SN '{sn}' completed in {time.time() - t_sn_start:.3f}s (MES: {t_mes:.3f}s, File Scan: {t_scan:.3f}s)")
            else:
                t1 = time.time()
                res, _ = self.search_engine.search_final_el(sn, stations_list, start, end, max_results=getattr(config, 'MAX_IMAGES_PER_SN', 0))
                print(f"[BATCH TIMING]: {tab.mode} SN '{sn}' completed in {time.time() - t1:.3f}s")
            all_results[sn] = res
        self.root.after(0, lambda: self._display_results(tab, all_results, search_id))

    def _display_results(self, tab, all_results, search_id):
        if search_id != self.search_id: return
        self.is_searching = False
        tab.btn_go.config(state="normal", text="Search", bg=config.COLOR_PRIMARY)
        tab.current_results = all_results
        for w in tab.result_container.winfo_children():
            if isinstance(w, ttk.Progressbar): w.destroy()
        if not all_results:
            tk.Label(tab.result_container, text="No results found", font=config.FONT_SUBTITLE, bg=config.COLOR_BG_MAIN, fg=config.COLOR_TEXT_SECONDARY).pack(pady=40)
            return
        found_sns = [sn for sn, res in all_results.items() if res]
        total_imgs = sum(len(res) for res in all_results.values())
        total_passes = sum(len(self.search_engine.group_by_timestamp(res)) for res in all_results.values())
        if found_sns:
            tab.btn_export.config(state="normal", text=f"Export Excel ({len(found_sns)})")
            if hasattr(tab, 'btn_export_imgs'):
                tab.btn_export_imgs.config(state="normal", text=f"Export Images ({total_imgs})")
        else:
            tab.btn_export.config(state="disabled", text="Export Excel")
            if hasattr(tab, 'btn_export_imgs'):
                tab.btn_export_imgs.config(state="disabled", text="Export Images")
        tab.lbl_summary.config(text=f"Found: {len(found_sns)}/{len(all_results)} SNs ({total_passes} passes • {total_imgs} imgs)")
        tab.sn_list_ordered = found_sns + [sn for sn in all_results if sn not in found_sns]
        if tab.mode == "PRE_EL":
            try:
                save_pre_el_search_history(tab.sn_list_ordered, all_results)
            except Exception:
                pass
        self.refresh_sidebar_for_tab(tab)
        if found_sns: tab._show_single_sn_results(found_sns[0])

    def open_original_image(self, path):
        try:
            if path and os.path.exists(path):
                os.startfile(path)
            else:
                messagebox.showerror("Error", f"Image file not found:\n{path}")
        except Exception as e:
            messagebox.showerror("Error", f"Could not open image file:\n{e}")

    def open_image(self, path):
        try:
            title = os.path.basename(path)
            ImageSnipperDialog(self.root, f"Image Inspection - {title}", path)
        except Exception:
            try:
                os.startfile(path)
            except Exception as e:
                messagebox.showerror("Error", f"Could not open image:\n{e}")

if __name__ == "__main__":
    root = tk.Tk()
    app = AOIDashboardApp(root)
    root.protocol("WM_DELETE_WINDOW", lambda: (root.destroy(), sys.exit(0)))
    root.mainloop()